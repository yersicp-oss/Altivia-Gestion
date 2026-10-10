"""ALTIVIA | Gestor integral de proyectos. Ejecutar: streamlit run app.py"""
import os
from html import escape
import io
import sqlite3
import hashlib
import hmac
import secrets
import tempfile
import re
import base64
import requests
import json
from contextlib import contextmanager
from urllib.parse import quote
from datetime import date, timedelta, datetime, timezone
from zoneinfo import ZoneInfo
import pandas as pd
import streamlit as st
import plotly.express as px

st.set_page_config(page_title='ALTIVIA | Gestión de Proyectos', page_icon='🏗️', layout='wide', initial_sidebar_state='expanded')
ROOT=os.path.dirname(os.path.abspath(__file__))
DB=os.environ.get('ALTIVIA_DB',os.path.join(ROOT,'altivia.db'))
TODAY=date.today()

# Mantener el acceso en este dispositivo es opcional. Se guardan únicamente
# tokens aleatorios en el navegador; los tokens se almacenan con hash en SQLite.
REMEMBER_COOKIE = 'gp_altivia_remember_v1'
REMEMBER_DAYS = 7

def today_peru():
    return datetime.now(ZoneInfo('America/Lima')).date()

def peru_timestamp():
    """Marca con zona horaria de Perú para el semáforo de 72 horas."""
    return datetime.now(ZoneInfo('America/Lima')).isoformat(timespec='seconds')

def delivery_signal(status,last_modified,now=None):
    """Verde si el Admin cambió Pendiente, si no amarillo o rojo a las 72 h."""
    if str(status or '').strip()!='Pendiente':
        return 'green',None
    clock=now or datetime.now(ZoneInfo('America/Lima'))
    if last_modified:
        try:
            point=datetime.fromisoformat(str(last_modified).replace('Z','+00:00'))
            if point.tzinfo is None:
                # Registros migrados de SQLite created_at en UTC.
                point=point.replace(tzinfo=timezone.utc)
            diff=clock-point.astimezone(ZoneInfo('America/Lima'))
            if diff>timedelta(hours=72):
                return 'red',int(diff.total_seconds()//86400)
        except (TypeError,ValueError,OverflowError):
            pass
    return 'yellow',None


PROJECT_STATES=['No iniciado','En desarrollo','En revisión','En corrección','Terminado']
TASK_STATES=['No iniciado','En desarrollo','En revisión','Con observaciones','Corregido','Aprobado']
DELIVERY_STATES=['Pendiente','En desarrollo','En revisión interna','Con observaciones','En corrección','Aprobado internamente','Enviado al cliente','Observado por cliente','Corregido','Aprobado','Entregado']
CHANGE_STATES=['Solicitado','Aprobado','Ejecutado']
MEETING_STATES=['Pendiente','En proceso','Completado','Atrasado','Cancelado']
SPECIALTIES=['Arquitectura','Estructuras','Eléctricas','Sanitarias','Comunicaciones','Modelado','Metrados y Presupuesto','BIM','Coordinación','Topografía']
PROJECT_TYPES=['Arquitectura','Estructuras','Instalaciones eléctricas','Instalaciones sanitarias','HVAC','BIM','Expediente técnico','Otro']
PRIORITIES=['Alta','Media','Baja']; RISKS=['Bajo','Medio','Alto','Crítico']
ROLES=['Gerente','Jefe de Proyecto','Coordinador','Arquitecto','Ingeniero','Dibujante','Modelador BIM','Revisor','Asistente','Otro']
CHECKS=['Información general correcta','Nombre del proyecto','Código del plano','Número de versión','Escala','Norte','Ejes','Niveles','Cotas','Nomenclatura','Simbología','Referencias','Detalles','Cuadro de áreas','Capas','Lineweights','Textos','Compatibilidad con otras especialidades','Formato de impresión','Revisión técnica','Revisión gráfica']

# key: (titulo, sql table, [ (column, visible label, type, required?, choices key) ])
SPECS={
 'Proyectos':('projects',[
 ('code','ID Proyecto','text',True,None),('name','Proyecto','text',True,None),('client','Cliente','text',False,None),('type','Tipo','select',False,'types'),('location','Ubicación','text',False,None),('manager','Jefe / Coordinador','person',False,None),('start_date','Fecha inicio','date',False,None),('due_date','Fecha entrega','date',False,None),('status','Estado','select',True,'project_states'),('priority','Prioridad','select',False,'priorities'),('notes','Observaciones','long',False,None)]),
 'Plan de trabajo':('tasks',[
 ('code','ID Tarea','text',True,None),('project_code','ID Proyecto','project',True,None),('activity','Actividad','text',True,None),('delivery_code','Entregable relacionado','delivery',False,None),('specialty','Especialidad','select',False,'specialties'),('owner','Responsable','person',True,None),('reviewer','Revisor','person',False,None),('start_date','Fecha de inicio','date',True,None),('due_date','Fecha término','date',True,None),('progress','Avance (%)','int',True,None),('status','Estado','select',True,'task_states'),('priority','Prioridad','select',False,'priorities'),('updated_at','Fecha actualización','date',False,None),('notes','Observaciones','long',False,None)]),
 'Entregables':('deliverables',[
 ('project_code','ID Proyecto','project',True,None),('code','ID Entregable','text',True,None),('name','Nombre del plano/documento','text',True,None),('specialty','Especialidad','select',False,'specialties'),('owner','Responsable','person',False,None),('reviewer','Revisor','person',False,None),('version','Versión','version',True,None),('due_date','Fecha de Presentación Final','date',True,None),('actual_date','Fecha de entregable','date',False,None),('status','Estado','select',True,'delivery_states'),('review_date','Fecha revisión','date',False,None),('correction_date','Fecha corrección','date',False,None),('approval_date','Fecha aprobación','date',False,None),('notes','Observaciones','long',False,None),('file_path','Enlace del documento (Google Drive / OneDrive / SharePoint)','text',False,None)]),
 'Control de cambios':('changes',[
 ('code','ID Cambio','text',True,None),('project_code','ID Proyecto','project',True,None),('request_date','Fecha solicitud','date',True,None),('requester','Solicitante','text',True,None),('description','Descripción cambio','long',True,None),('reason','Motivo','long',False,None),('specialty','Especialidad afectada','select',False,'specialties'),('affected_drawings','Planos afectados','text',False,None),('owner','Responsable','person',True,None),('schedule_impact','Impacto en plazo','select',False,'risks'),('new_due_date','Nueva fecha entrega','date',False,None),('approved_by','Aprobado por','text',False,None),('approval_date','Fecha aprobación','date',False,None),('status','Estado','select',True,'change_states'),('notes','Observaciones','long',False,None)]),
 'Personal':('people',[
 ('code','Nombre','text',True,None),('role','Cargo','select',False,'roles'),('specialty','Especialidad','select',False,'specialties'),('email','Correo','text',False,None),('phone','Teléfono','text',False,None),('status','Disponibilidad','select',True,'people_states'),('notes','Observaciones','long',False,None)])
}
OPTIONS={'types':PROJECT_TYPES,'project_states':PROJECT_STATES,'task_states':TASK_STATES,'delivery_states':DELIVERY_STATES,'change_states':CHANGE_STATES,'meeting_states':MEETING_STATES,'specialties':SPECIALTIES,'priorities':PRIORITIES,'risks':RISKS,'roles':ROLES,'people_states':['Disponible','Ocupado','No disponible']}
FINISHED_TASK={'Terminado','Aprobado'}
FINISHED_DELIVERY={'Aprobado','Entregado'}
CLOSED_CHANGE={'Ejecutado'}

@contextmanager
def connection():
    con=sqlite3.connect(DB,timeout=20)
    con.row_factory=sqlite3.Row
    con.execute('PRAGMA foreign_keys=ON')
    try:
        yield con
        con.commit()
    except Exception:
        con.rollback()
        raise
    finally:
        con.close()

def initialize():
    with connection() as con:
        for _,(table,fields) in SPECS.items():
            cols=['id INTEGER PRIMARY KEY AUTOINCREMENT','created_at TEXT DEFAULT CURRENT_TIMESTAMP']
            for key,_,kind,_,_ in fields:
                sql_type='REAL' if kind=='float' else 'INTEGER' if kind in ('int','bool') else 'TEXT'
                cols.append(f'"{key}" {sql_type}'+(' UNIQUE' if key=='code' else ''))
            if table=='people':cols.append('name TEXT')  # columna interna para compatibilidad con datos antiguos
            con.execute(f'CREATE TABLE IF NOT EXISTS {table} ({", ".join(cols)})')
        # Catálogo de entregables: alta y modificaciones reservadas al administrador.
        # La tabla 'deliverables' sigue siendo el registro de la entrega y de su versión vigente.
        # Vinculo permanente cuenta Consulta -> ficha en Personal.
        cols_people={r[1] for r in con.execute('PRAGMA table_info(people)')}
        if 'user_id' not in cols_people:
            con.execute('ALTER TABLE people ADD COLUMN user_id INTEGER')
        if 'linked_once' not in cols_people:
            con.execute('ALTER TABLE people ADD COLUMN linked_once INTEGER NOT NULL DEFAULT 0')
        con.execute('UPDATE people SET linked_once=1 WHERE user_id IS NOT NULL')
        con.execute('CREATE UNIQUE INDEX IF NOT EXISTS idx_people_user ON people(user_id)')
        # Multiples responsables y revisores por tarea, sin duplicar personas.
        con.execute('''CREATE TABLE IF NOT EXISTS task_people (
            task_id INTEGER NOT NULL REFERENCES tasks(id) ON DELETE CASCADE,
            person_id INTEGER NOT NULL REFERENCES people(id) ON DELETE RESTRICT,
            assignment_role TEXT NOT NULL CHECK(assignment_role IN ('Responsable','Revisor')),
            PRIMARY KEY(task_id,person_id,assignment_role))''')
        con.execute('CREATE INDEX IF NOT EXISTS idx_task_people_person ON task_people(person_id,task_id)')
        con.execute('''CREATE TABLE IF NOT EXISTS deliverable_catalog (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            project_code TEXT NOT NULL,
            code TEXT NOT NULL UNIQUE,
            name TEXT NOT NULL,
            specialty TEXT,
            final_due_date TEXT NOT NULL,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP)''')
        # Nueva relación: tarea seleccionada en la entrega y snapshot por versión.
        # No rompe entregas históricas que aún no tenían tarea elegida.
        deliverable_cols={r[1] for r in con.execute('PRAGMA table_info(deliverables)')}
        if 'task_id' not in deliverable_cols:
            con.execute('ALTER TABLE deliverables ADD COLUMN task_id INTEGER')
        if 'last_modified_at' not in deliverable_cols:
            con.execute('ALTER TABLE deliverables ADD COLUMN last_modified_at TEXT')
        con.execute('''CREATE TABLE IF NOT EXISTS delivery_version_tasks (
            delivery_id INTEGER NOT NULL REFERENCES deliverables(id) ON DELETE CASCADE,
            version TEXT NOT NULL,
            task_id INTEGER REFERENCES tasks(id) ON DELETE SET NULL,
            task_code TEXT,
            activity TEXT,
            owners TEXT,
            reviewers TEXT,
            PRIMARY KEY(delivery_id,version))''')
        # Backfill conservador: el campo actual_date/última versión representa la
        # última entrega; created_at representa UTC y se usa solo como respaldo.
        con.execute('''UPDATE deliverables SET last_modified_at =
            CASE WHEN actual_date IS NOT NULL AND trim(actual_date)<>''
                 THEN substr(actual_date,1,10)||'T00:00:00-05:00'
                 ELSE replace(created_at,' ','T')||'+00:00' END
            WHERE last_modified_at IS NULL OR trim(last_modified_at)='' ''')
        # Compatibilidad con SQLite existente: preservar vínculos y enlaces viejos, aunque
        # la casilla 'Código del plano' ya no sea visible en el formulario.
        cols={r[1] for r in con.execute('PRAGMA table_info(deliverables)')}
        if 'drawing_code' not in cols:
            con.execute('ALTER TABLE deliverables ADD COLUMN drawing_code TEXT')
        if 'file_path' not in cols:
            con.execute('ALTER TABLE deliverables ADD COLUMN file_path TEXT')
        con.execute('''INSERT OR IGNORE INTO deliverable_catalog(project_code,code,name,specialty,final_due_date)
            SELECT project_code,code,name,specialty,COALESCE(NULLIF(due_date,''),date('now'))
            FROM deliverables WHERE code IS NOT NULL AND project_code IS NOT NULL''')
        con.execute('''CREATE TABLE IF NOT EXISTS meetings (id INTEGER PRIMARY KEY AUTOINCREMENT, code TEXT UNIQUE, project_code TEXT, owner TEXT, due_date TEXT, status TEXT)''')
        con.execute('CREATE TABLE IF NOT EXISTS versions (id INTEGER PRIMARY KEY AUTOINCREMENT, delivery_code TEXT NOT NULL, version TEXT NOT NULL, registered_at TEXT NOT NULL, notes TEXT, UNIQUE(delivery_code,version))')
        con.execute('CREATE TABLE IF NOT EXISTS checklist (id INTEGER PRIMARY KEY AUTOINCREMENT, delivery_code TEXT NOT NULL, criterion TEXT NOT NULL, result TEXT NOT NULL DEFAULT "PENDIENTE", notes TEXT, UNIQUE(delivery_code,criterion))')
        con.execute('CREATE TABLE IF NOT EXISTS audit (id INTEGER PRIMARY KEY AUTOINCREMENT, occurred_at TEXT DEFAULT CURRENT_TIMESTAMP, module TEXT, record_code TEXT, action TEXT)')
        # Compatibilidad: conservar datos de instalaciones anteriores.
        con.execute('''CREATE TABLE IF NOT EXISTS delivery_files (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            delivery_id INTEGER NOT NULL,
            version TEXT NOT NULL,
            kind TEXT NOT NULL CHECK(kind IN ('pdf','editable')),
            filename TEXT NOT NULL,
            content BLOB NOT NULL,
            uploaded_at TEXT DEFAULT CURRENT_TIMESTAMP,
            UNIQUE(delivery_id,version,kind),
            FOREIGN KEY(delivery_id) REFERENCES deliverables(id) ON DELETE CASCADE
        )''')
        con.execute('''CREATE TABLE IF NOT EXISTS drive_files (
            delivery_id INTEGER NOT NULL, version TEXT NOT NULL, kind TEXT NOT NULL,
            file_id TEXT NOT NULL, filename TEXT NOT NULL, uploaded_at TEXT DEFAULT CURRENT_TIMESTAMP,
            PRIMARY KEY(delivery_id,version,kind),
            FOREIGN KEY(delivery_id) REFERENCES deliverables(id) ON DELETE CASCADE)''')
        # Cada versión puede contener múltiples documentos, con PDF y editable propios.
        con.execute('''CREATE TABLE IF NOT EXISTS delivery_documents (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            delivery_id INTEGER NOT NULL,
            version TEXT NOT NULL,
            doc_no TEXT NOT NULL,
            title TEXT NOT NULL,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP,
            UNIQUE(delivery_id,version,doc_no),
            FOREIGN KEY(delivery_id) REFERENCES deliverables(id) ON DELETE CASCADE)''')
        con.execute('''CREATE TABLE IF NOT EXISTS document_files (
            document_id INTEGER NOT NULL,
            kind TEXT NOT NULL CHECK(kind IN ('pdf','editable')),
            file_id TEXT NOT NULL,
            filename TEXT NOT NULL,
            uploaded_at TEXT DEFAULT CURRENT_TIMESTAMP,
            PRIMARY KEY(document_id,kind),
            FOREIGN KEY(document_id) REFERENCES delivery_documents(id) ON DELETE CASCADE)''')
        con.execute('''CREATE TABLE IF NOT EXISTS document_checklist (
            document_id INTEGER NOT NULL,
            criterion TEXT NOT NULL,
            result TEXT NOT NULL DEFAULT 'PENDIENTE',
            notes TEXT,
            PRIMARY KEY(document_id,criterion),
            FOREIGN KEY(document_id) REFERENCES delivery_documents(id) ON DELETE CASCADE)''')
        con.execute('''CREATE TABLE IF NOT EXISTS document_serials (
            delivery_id INTEGER NOT NULL, version TEXT NOT NULL, highest INTEGER NOT NULL,
            PRIMARY KEY(delivery_id,version),
            FOREIGN KEY(delivery_id) REFERENCES deliverables(id) ON DELETE CASCADE)''')
        # Acceso opcional por documento. Si no se configura, hereda responsables y
        # revisores de las tareas relacionadas con el entregable (compatibilidad).
        # Si el Admin configura la lista, solo sus miembros tienen acceso al Checklist.
        con.execute('''CREATE TABLE IF NOT EXISTS document_assignment_mode (
            document_id INTEGER PRIMARY KEY REFERENCES delivery_documents(id) ON DELETE CASCADE)''')
        con.execute('''CREATE TABLE IF NOT EXISTS document_people (
            document_id INTEGER NOT NULL REFERENCES delivery_documents(id) ON DELETE CASCADE,
            person_id INTEGER NOT NULL REFERENCES people(id) ON DELETE CASCADE,
            assignment_role TEXT NOT NULL CHECK(assignment_role IN ('Responsable','Revisor')),
            PRIMARY KEY(document_id,person_id,assignment_role))''')
        con.execute('CREATE INDEX IF NOT EXISTS idx_document_people_person ON document_people(person_id,document_id)')
        # Un marcador evita restaurar D01 si el Administrador lo elimina después.
        con.execute('''CREATE TABLE IF NOT EXISTS legacy_document_imports (
            delivery_id INTEGER NOT NULL, version TEXT NOT NULL,
            PRIMARY KEY(delivery_id,version),
            FOREIGN KEY(delivery_id) REFERENCES deliverables(id) ON DELETE CASCADE)''')
        legacy=con.execute('''SELECT delivery_id,version FROM drive_files
                   UNION SELECT delivery_id,version FROM delivery_files''').fetchall()
        for entry in legacy:
            did,ver=entry['delivery_id'],entry['version']
            flagged=con.execute('SELECT 1 FROM legacy_document_imports WHERE delivery_id=? AND version=?',(did,ver)).fetchone()
            if flagged:continue
            exists=con.execute('SELECT id FROM deliverables WHERE id=?',(did,)).fetchone()
            if not exists:continue
            con.execute('INSERT OR IGNORE INTO delivery_documents(delivery_id,version,doc_no,title) VALUES(?,?,?,?)',
                        (did,ver,'D01','Documento migrado (D01)'))
            doc=con.execute('SELECT id FROM delivery_documents WHERE delivery_id=? AND version=? AND doc_no=?',
                            (did,ver,'D01')).fetchone()
            for f in con.execute('SELECT kind,file_id,filename FROM drive_files WHERE delivery_id=? AND version=?',(did,ver)).fetchall():
                con.execute('INSERT OR IGNORE INTO document_files(document_id,kind,file_id,filename) VALUES(?,?,?,?)',
                            (doc['id'],f['kind'],f['file_id'],f['filename']))
            # El antiguo checklist del entregable se copia a D01 UNA sola vez, sin borrar el original.
            legacy_code=con.execute('SELECT code,version FROM deliverables WHERE id=?',(did,)).fetchone()
            if legacy_code and legacy_code['version']==ver:
                for ch in con.execute('SELECT criterion,result,notes FROM checklist WHERE delivery_code=?',(legacy_code['code'],)).fetchall():
                    con.execute('INSERT OR IGNORE INTO document_checklist(document_id,criterion,result,notes) VALUES(?,?,?,?)',
                                (doc['id'],ch['criterion'],ch['result'],ch['notes']))
            con.execute('INSERT OR IGNORE INTO document_serials(delivery_id,version,highest) VALUES(?,?,1)',(did,ver))
            con.execute('INSERT INTO legacy_document_imports(delivery_id,version) VALUES(?,?)',(did,ver))
        # Sincroniza máximo con documentos existentes sin reducirlo tras eliminaciones.
        for row in con.execute('SELECT delivery_id,version,doc_no FROM delivery_documents').fetchall():
            match=re.fullmatch(r'D(\d+)',row['doc_no'])
            if match:
                con.execute('''INSERT INTO document_serials(delivery_id,version,highest) VALUES(?,?,?)
                    ON CONFLICT(delivery_id,version) DO UPDATE SET highest=MAX(highest,excluded.highest)''',
                    (row['delivery_id'],row['version'],int(match.group(1))))
        # Máximo histórico: aunque se elimine V02, ese número nunca se reutiliza.
        con.execute('''CREATE TABLE IF NOT EXISTS version_counters (
            delivery_code TEXT PRIMARY KEY, highest INTEGER NOT NULL DEFAULT 0)''')
        known=con.execute('SELECT delivery_code,version FROM versions UNION SELECT code,version FROM deliverables').fetchall()
        for r in known:
            m=re.fullmatch(r'V(\d+)',str(r['version'] or '').upper())
            if m:
                con.execute('''INSERT INTO version_counters(delivery_code,highest) VALUES(?,?)
                    ON CONFLICT(delivery_code) DO UPDATE SET highest=MAX(highest,excluded.highest)''',
                    (r['delivery_code'],int(m.group(1))))
        # Importa asignaciones antiguas basadas en texto UNA sola vez. No reimportar
        # al quitar asignaciones desde el editor moderno.
        if not con.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='task_people_migration'").fetchone():
            con.execute('CREATE TABLE task_people_migration (task_id INTEGER PRIMARY KEY)')
        for old in con.execute('SELECT id,owner,reviewer FROM tasks').fetchall():
            if con.execute('SELECT 1 FROM task_people_migration WHERE task_id=?',(old['id'],)).fetchone():
                continue
            for role,field in [('Responsable','owner'),('Revisor','reviewer')]:
                label=str(old[field] or '').strip()
                if not label:continue
                person=con.execute('SELECT id FROM people WHERE name=? COLLATE NOCASE',(label,)).fetchone()
                if person:
                    con.execute('INSERT OR IGNORE INTO task_people(task_id,person_id,assignment_role) VALUES(?,?,?)',
                        (old['id'],person['id'],role))
            con.execute('INSERT OR IGNORE INTO task_people_migration(task_id) VALUES(?)',(old['id'],))
        # Secuencia T001, T002... independiente por proyecto. No se reduce al borrar
        # una tarea. Detecta tareas heredadas que ya utilicen el formato nuevo.
        con.execute('''CREATE TABLE IF NOT EXISTS task_counters (
            project_code TEXT PRIMARY KEY,
            highest INTEGER NOT NULL DEFAULT 0)''')
        for old_task in con.execute('SELECT project_code,code FROM tasks WHERE project_code IS NOT NULL').fetchall():
            project_code=str(old_task['project_code'] or '').strip()
            match=re.fullmatch(re.escape(project_code)+r'_T(\d+)',str(old_task['code'] or ''),re.IGNORECASE)
            if match:
                con.execute('''INSERT INTO task_counters(project_code,highest) VALUES(?,?)
                    ON CONFLICT(project_code) DO UPDATE SET highest=MAX(highest,excluded.highest)''',
                    (project_code,int(match.group(1))))
        con.execute('CREATE INDEX IF NOT EXISTS idx_tasks_project ON tasks(project_code)')
        con.execute('CREATE INDEX IF NOT EXISTS idx_deliveries_project ON deliverables(project_code)')

def hash_password(password):
    salt=secrets.token_bytes(16)
    digest=hashlib.pbkdf2_hmac('sha256',password.encode('utf-8'),salt,310000)
    return 'pbkdf2_sha256$310000$'+salt.hex()+'$'+digest.hex()

def verify_password(password,stored):
    try:
        algo,iterations,salt,digest=stored.split('$')
        if algo!='pbkdf2_sha256':return False
        actual=hashlib.pbkdf2_hmac('sha256',password.encode('utf-8'),bytes.fromhex(salt),int(iterations))
        return hmac.compare_digest(actual,bytes.fromhex(digest))
    except (ValueError,TypeError):return False

def _unique_person_name(con, full_name, username, exclude_person_id=None):
    """Evita confundir homónimos; conserva nombres existentes siempre que pueda."""
    full_name=str(full_name or '').strip()
    taken=con.execute('SELECT id FROM people WHERE name=? COLLATE NOCASE AND id<>?',
                      (full_name,int(exclude_person_id or -1))).fetchone()
    if not taken:return full_name
    alternative=f'{full_name} ({username})'
    taken=con.execute('SELECT id FROM people WHERE name=? COLLATE NOCASE AND id<>?',
                      (alternative,int(exclude_person_id or -1))).fetchone()
    if taken:raise ValueError('Hay una ficha de Personal con el mismo nombre y usuario; revisa los duplicados.')
    return alternative


def _refresh_task_display(con, task_ids):
    """Actualiza campos legados owner/reviewer para reportes y edición histórica."""
    for task_id in set(map(int,task_ids)):
        labels={}
        for role,field in [('Responsable','owner'),('Revisor','reviewer')]:
            names=[r[0] for r in con.execute("""SELECT p.name FROM task_people tp JOIN people p ON p.id=tp.person_id
                WHERE tp.task_id=? AND tp.assignment_role=? ORDER BY p.name""",(task_id,role))]
            labels[field]=', '.join(names)
        con.execute('UPDATE tasks SET owner=?,reviewer=? WHERE id=?',
                    (labels['owner'],labels['reviewer'],task_id))


def ensure_consulta_person(con,user_id,full_name,username):
    """Crea o vincula Personal a una cuenta Consulta por ID; idempotente."""
    account=con.execute('SELECT role FROM users WHERE id=?',(int(user_id),)).fetchone()
    if not account or account['role']!='Consulta':return None
    existing=con.execute('SELECT id,name FROM people WHERE user_id=?',(int(user_id),)).fetchone()
    if existing:
        new_name=_unique_person_name(con,full_name,username,existing['id'])
        if existing['name']!=new_name:
            con.execute('UPDATE people SET name=?,code=? WHERE id=?',(new_name,new_name,existing['id']))
            _refresh_task_display(con,[r[0] for r in con.execute(
                'SELECT DISTINCT task_id FROM task_people WHERE person_id=?',(existing['id'],))])
        return existing['id']
    # Si había Personal con el mismo nombre y sin vincular, recuperar esa ficha.
    candidate=con.execute("""SELECT id FROM people WHERE name=? COLLATE NOCASE
                             AND user_id IS NULL AND linked_once=0
                             ORDER BY id LIMIT 1""",(str(full_name).strip(),)).fetchone()
    if candidate:
        con.execute('UPDATE people SET user_id=?,linked_once=1 WHERE id=?',(int(user_id),candidate['id']))
        return candidate['id']
    # Una cuenta nueva nunca hereda automáticamente las tareas de una cuenta eliminada,
    # incluso si tienen el mismo nombre completo.
    label=_unique_person_name(con,full_name,username)
    cur=con.execute("""INSERT INTO people(code,name,role,specialty,status,user_id,linked_once)
                       VALUES(?,?,?,'','Disponible',?,1)""",(label,label,'Otro',int(user_id)))
    return cur.lastrowid


def task_assigned_people(task_id):
    with connection() as con:
        rows=con.execute('SELECT person_id,assignment_role FROM task_people WHERE task_id=?',(int(task_id),)).fetchall()
    return ({int(r['person_id']) for r in rows if r['assignment_role']=='Responsable'},
            {int(r['person_id']) for r in rows if r['assignment_role']=='Revisor'})


def visible_consulta_tasks(frame,user_id):
    """Filtro servidor, no depende de listas aportadas por el navegador."""
    with connection() as con:
        ids={r[0] for r in con.execute("""SELECT DISTINCT tp.task_id
            FROM task_people tp JOIN people p ON p.id=tp.person_id
            JOIN users u ON u.id=p.user_id
            WHERE u.id=? AND u.role='Consulta' AND u.active=1""",(int(user_id),))}
    return frame[frame['id'].isin(ids)].copy()


def people_for_assignment():
    with connection() as con:
        rows=con.execute("""SELECT p.id,p.name,p.status,u.role AS account_role,u.active AS account_active
            FROM people p LEFT JOIN users u ON u.id=p.user_id ORDER BY p.name COLLATE NOCASE""").fetchall()
    return [dict(r) for r in rows]


def initialize_auth():
    # Preserva usuarios existentes y elimina la restriccion antigua de roles.
    with sqlite3.connect(DB, timeout=20) as migration:
        row = migration.execute("SELECT sql FROM sqlite_master WHERE type='table' AND name='users'").fetchone()
        if row and 'CHECK(role IN' in (row[0] or ''):
            migration.execute('CREATE TABLE users_migration (id INTEGER PRIMARY KEY AUTOINCREMENT, username TEXT UNIQUE NOT NULL, full_name TEXT NOT NULL, password_hash TEXT NOT NULL, role TEXT NOT NULL, active INTEGER NOT NULL DEFAULT 1, created_at TEXT DEFAULT CURRENT_TIMESTAMP)')
            migration.execute('INSERT INTO users_migration (id,username,full_name,password_hash,role,active,created_at) SELECT id,username,full_name,password_hash,role,active,created_at FROM users')
            migration.execute('DROP TABLE users')
            migration.execute('ALTER TABLE users_migration RENAME TO users')
    with connection() as con:
        con.execute("""CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT, username TEXT UNIQUE NOT NULL,
            full_name TEXT NOT NULL, password_hash TEXT NOT NULL,
            role TEXT NOT NULL,
            active INTEGER NOT NULL DEFAULT 1, created_at TEXT DEFAULT CURRENT_TIMESTAMP)""")
        con.execute('''CREATE TABLE IF NOT EXISTS client_access (user_id INTEGER NOT NULL, delivery_id INTEGER NOT NULL, granted_at TEXT DEFAULT CURRENT_TIMESTAMP, PRIMARY KEY(user_id,delivery_id), FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE CASCADE, FOREIGN KEY(delivery_id) REFERENCES deliverables(id) ON DELETE CASCADE)''')
        con.execute("""CREATE TABLE IF NOT EXISTS remembered_sessions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
            token_hash TEXT NOT NULL UNIQUE,
            password_fingerprint TEXT NOT NULL,
            created_at TEXT NOT NULL,
            expires_at TEXT NOT NULL
        )""")
        con.execute('CREATE INDEX IF NOT EXISTS idx_remembered_user ON remembered_sessions(user_id)')
        # Autorizaciones de VISUALIZACIÓN de proyectos para usuarios Consulta.
        # Se asocian a los ID internos (no al código), resistentes a renombrados.
        con.execute('''CREATE TABLE IF NOT EXISTS consulta_project_access (
            user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
            project_id INTEGER NOT NULL REFERENCES projects(id) ON DELETE CASCADE,
            granted_at TEXT DEFAULT CURRENT_TIMESTAMP,
            PRIMARY KEY(user_id, project_id))''')
        con.execute('CREATE INDEX IF NOT EXISTS idx_consulta_project_access_project ON consulta_project_access(project_id)')
        users=con.execute('SELECT COUNT(*) FROM users').fetchone()[0]
        if users==0:
            password=os.getenv('ALTIVIA_ADMIN_PASSWORD','')
            if not password:
                try:password=st.secrets.get('ALTIVIA_ADMIN_PASSWORD','')
                except Exception:pass
            if len(password)>=6:
                con.execute('INSERT INTO users(username,full_name,password_hash,role) VALUES (?,?,?,?)',
                            ('admin','Administrador principal',hash_password(password),'Administrador'))
            else:return False
        # Incorpora también las cuentas Consulta creadas antes de esta versión.
        for user in con.execute("SELECT id,username,full_name FROM users WHERE role='Consulta'").fetchall():
            ensure_consulta_person(con,user['id'],user['full_name'],user['username'])
        # En instalaciones antiguas podía existir una tarea antes de crear su usuario.
        # Se enlaza solo cuando el nombre coincide exactamente con una ficha y
        # no existe aún una asignación de ese tipo. No se inventan asignaciones.
        for task in con.execute('SELECT id,owner,reviewer FROM tasks').fetchall():
            for role,column in [('Responsable','owner'),('Revisor','reviewer')]:
                if con.execute('SELECT 1 FROM task_people WHERE task_id=? AND assignment_role=?',
                               (task['id'],role)).fetchone():continue
                old_name=str(task[column] or '').strip()
                if not old_name:continue
                person=con.execute('SELECT id FROM people WHERE name=? COLLATE NOCASE',(old_name,)).fetchone()
                if person:
                    con.execute('INSERT OR IGNORE INTO task_people(task_id,person_id,assignment_role) VALUES(?,?,?)',
                                (task['id'],person['id'],role))
    return True

def can_edit():
    return st.session_state.get('role')=='Administrador'

def require_admin():
    if not can_edit():raise PermissionError('Acceso denegado: se requiere rol Administrador.')


def can_create_deliverable():
    return st.session_state.get('role') in ('Administrador','Consulta')


def _password_fingerprint(stored_hash):
    # Invalida sesiones persistentes después de cambiar una contraseña.
    return hashlib.sha256(str(stored_hash).encode('utf-8')).hexdigest()


def issue_remember_token(user_id, stored_hash):
    """Crea una autorización revocable; nunca se guarda el token sin hash."""
    token = secrets.token_urlsafe(48)
    now = datetime.now(timezone.utc)
    expires = now + timedelta(days=REMEMBER_DAYS)
    with connection() as con:
        con.execute('DELETE FROM remembered_sessions WHERE expires_at<=?', (now.isoformat(),))
        con.execute('''INSERT INTO remembered_sessions
            (user_id, token_hash, password_fingerprint, created_at, expires_at)
            VALUES (?,?,?,?,?)''',
            (int(user_id), hashlib.sha256(token.encode()).hexdigest(),
             _password_fingerprint(stored_hash), now.isoformat(), expires.isoformat()))
        # Máximo 10 dispositivos activos por usuario, privilegiando los recientes.
        con.execute('''DELETE FROM remembered_sessions
           WHERE user_id=? AND id NOT IN
           (SELECT id FROM remembered_sessions WHERE user_id=? ORDER BY id DESC LIMIT 10)''',
           (int(user_id),int(user_id)))
    return token


def validate_remember_token(token):
    """Devuelve usuario activo si el token sigue vigente y no cambió su clave."""
    if not isinstance(token,str) or not (30 <= len(token) <= 250):
        return None
    token_hash = hashlib.sha256(token.encode()).hexdigest()
    with connection() as con:
        r=con.execute('''SELECT s.id, s.expires_at, s.password_fingerprint,
                        u.id AS user_id, u.username,u.full_name,u.role,u.active,u.password_hash
                        FROM remembered_sessions s JOIN users u ON u.id=s.user_id
                        WHERE s.token_hash=?''',(token_hash,)).fetchone()
        if not r:return None
        try:expires = datetime.fromisoformat(r['expires_at'])
        except (ValueError,TypeError):expires = datetime.min.replace(tzinfo=timezone.utc)
        if expires.tzinfo is None:expires=expires.replace(tzinfo=timezone.utc)
        if (not r['active'] or expires <= datetime.now(timezone.utc)
            or not hmac.compare_digest(r['password_fingerprint'],_password_fingerprint(r['password_hash']))):
            con.execute('DELETE FROM remembered_sessions WHERE id=?',(r['id'],))
            return None
        return dict(r)


def remembered_cookie_value():
    # st.context.cookies refleja las cookies enviadas por el navegador al cargar
    # esta sesión nueva y permite restaurarla sin incluir tokens en la URL.
    try:return st.context.cookies.get(REMEMBER_COOKIE)
    except (AttributeError,KeyError,TypeError):return None


def clear_remember_token(token):
    if isinstance(token,str) and token:
        with connection() as con:
            con.execute('DELETE FROM remembered_sessions WHERE token_hash=?',
                (hashlib.sha256(token.encode()).hexdigest(),))


def write_remember_cookie(token):
    # CookieController usa JavaScript; la cookie NO es HttpOnly.
    # Solo se incluye un token opaco aleatorio; nunca contraseñas o roles.
    from streamlit_cookies_controller import CookieController
    CookieController().set(REMEMBER_COOKIE, token,
        max_age=REMEMBER_DAYS * 86400, secure=True, same_site='strict', path='/')


def delete_remember_cookie():
    from streamlit_cookies_controller import CookieController
    CookieController().remove(REMEMBER_COOKIE)


def restore_remembered_login():
    if st.session_state.get('user_id') or st.session_state.get('_remember_checked'):
        return
    st.session_state['_remember_checked']=True
    token=remembered_cookie_value()
    account=validate_remember_token(token)
    if account:
        for key in ('user_id','username','full_name','role'):
            st.session_state[key]=account[key]
        st.session_state['_remember_token']=token
        st.session_state['_remembered_login']=True


def end_login_session():
    token=st.session_state.get('_remember_token') or remembered_cookie_value()
    clear_remember_token(token)
    try:delete_remember_cookie()
    except Exception:pass  # La revocación en servidor ya invalida el token.
    for key in ('user_id','role','username','full_name','client_preview',
                '_remember_token','_remembered_login'):
        st.session_state.pop(key,None)
    st.session_state['_remember_checked']=True


def delete_user_account(target_id):
    """Borra solo la cuenta y sus permisos, nunca proyectos ni archivos."""
    require_admin()
    target_id=int(target_id)
    actor=int(st.session_state['user_id'])
    if target_id==actor:
        raise ValueError('No puedes eliminar tu propia cuenta de Administrador.')
    with connection() as con:
        r=con.execute('SELECT username,role,active FROM users WHERE id=?',(target_id,)).fetchone()
        if not r:raise ValueError('El usuario ya no existe.')
        if r['role']=='Administrador' and r['active']:
            remaining=con.execute("SELECT COUNT(*) FROM users WHERE role='Administrador' AND active=1").fetchone()[0]
            if remaining<=1:raise ValueError('Debe quedar al menos un administrador activo.')
        con.execute('DELETE FROM remembered_sessions WHERE user_id=?',(target_id,))
        con.execute('DELETE FROM client_access WHERE user_id=?',(target_id,))
        # Conservar la ficha de Personal y el historial de asignaciones.
        con.execute('UPDATE people SET user_id=NULL WHERE user_id=?',(target_id,))
        con.execute('DELETE FROM users WHERE id=?',(target_id,))
        con.execute('INSERT INTO audit(module,record_code,action) VALUES (?,?,?)',
            ('Usuarios',r['username'],f'Cuenta eliminada por administrador ID {actor}'))
    return r['username']


def login_page():
    st.title('🔐 ALTIVIA | Acceso al sistema')
    st.caption('Ingrese sus credenciales para acceder a la gestión de proyectos.')
    with st.form('login_form'):
        username=st.text_input('Usuario').strip().lower()
        password=st.text_input('Contraseña',type='password')
        remember=st.checkbox('Mantener sesión iniciada en este dispositivo (7 días)',value=True,
            help='No usar en equipos públicos o compartidos. Cerrar sesión revoca el acceso guardado.')
        submitted=st.form_submit_button('Ingresar',type='primary')
    if submitted:
        if st.session_state.get('failed_logins',0)>=8:
            st.error('Demasiados intentos fallidos en esta sesión. Cierre el navegador y contacte al administrador.')
            return
        with connection() as con:
            r=con.execute('SELECT id,username,full_name,role,active,password_hash FROM users WHERE username=?',(username,)).fetchone()
        if r and r['active']==1 and verify_password(password,r['password_hash']):
            st.session_state['user_id']=r['id'];st.session_state['role']=r['role']
            st.session_state['username']=r['username'];st.session_state['full_name']=r['full_name']
            st.session_state['failed_logins']=0
            if remember:
                try:
                    new_token=issue_remember_token(r['id'],r['password_hash'])
                    write_remember_cookie(new_token)
                    st.session_state['_remember_token']=new_token
                except Exception:
                    # No impide usar la app: solo desactiva recordar sesión.
                    if 'new_token' in locals():clear_remember_token(new_token)
                    st.warning('Ingresaste correctamente, pero no se pudo activar «Mantener sesión».')
            else:
                clear_remember_token(remembered_cookie_value())
                try:delete_remember_cookie()
                except Exception:pass
                st.session_state.pop('_remember_token',None)
            st.rerun()
        else:
            st.session_state['failed_logins']=st.session_state.get('failed_logins',0)+1
            st.error('Credenciales incorrectas o usuario desactivado.')

def normalized_username(raw):
    """Identificador de acceso estable: se normaliza sin cambiar el ID del usuario."""
    username=str(raw or '').strip().lower()
    if not re.fullmatch(r'[a-z0-9][a-z0-9._-]{2,39}',username):
        raise ValueError('El usuario debe tener entre 3 y 40 caracteres: letras, números, punto, guion o guion bajo; empezar con letra o número.')
    return username


def update_user_by_admin(target_id,username,full_name,role,active,new_password='',project_ids=None):
    """Edita una cuenta por su ID; conserva permisos de cliente y las referencias."""
    require_admin()
    target_id=int(target_id)
    actor=int(st.session_state['user_id'])
    username=normalized_username(username)
    full_name=str(full_name or '').strip()
    if not full_name:raise ValueError('El nombre completo es obligatorio.')
    if len(full_name)>150:raise ValueError('El nombre completo no debe exceder 150 caracteres.')
    if role not in ('Administrador','Consulta','Cliente'):
        raise ValueError('Rol no válido.')
    if new_password and len(new_password)<6:
        raise ValueError('La contraseña nueva debe tener al menos 6 caracteres.')
    if target_id==actor and (role!='Administrador' or not active):
        raise ValueError('No puedes desactivar tu propia cuenta ni quitarte el rol Administrador.')
    with connection() as con:
        previous=con.execute('SELECT username,full_name,role,active FROM users WHERE id=?',(target_id,)).fetchone()
        if not previous:raise ValueError('El usuario seleccionado ya no existe.')
        if previous['role']=='Administrador' and previous['active'] and (role!='Administrador' or not active):
            others=con.execute("SELECT COUNT(*) FROM users WHERE role='Administrador' AND active=1 AND id<>?",(target_id,)).fetchone()[0]
            if others<1:raise ValueError('Debe quedar al menos un Administrador activo.')
        conflict=con.execute('SELECT id FROM users WHERE lower(username)=? AND id<>?',(username,target_id)).fetchone()
        if conflict:raise ValueError('El nombre de usuario ya está registrado. Elige otro.')
        con.execute('UPDATE users SET username=?,full_name=?,role=?,active=? WHERE id=?',
                    (username,full_name,role,int(bool(active)),target_id))
        if new_password:
            con.execute('UPDATE users SET password_hash=? WHERE id=?',(hash_password(new_password),target_id))
        # Sólo los cambios de credenciales, acceso o rol revocan sesiones recordadas.
        # Cambiar el nombre visible o el usuario no invalida la sesión existente.
        if new_password or role!=previous['role'] or bool(active)!=bool(previous['active']):
            con.execute('DELETE FROM remembered_sessions WHERE user_id=?',(target_id,))
        if role=='Consulta':
            ensure_consulta_person(con,target_id,full_name,username)
        if project_ids is not None or role!='Consulta':
            set_consulta_project_access(con,target_id,project_ids or [])
        con.execute('INSERT INTO audit(module,record_code,action) VALUES (?,?,?)',
                    ('Usuarios',username,f'Perfil editado por Administrador ID {actor}'))
    if target_id==actor:
        st.session_state['username']=username
        st.session_state['full_name']=full_name
        if new_password:
            st.session_state.pop('_remember_token',None)
            try:delete_remember_cookie()
            except Exception:pass
    return username


def update_own_profile(username,full_name,current_password):
    """Cada usuario puede editar solamente su propia cuenta, comprobando su clave."""
    actor=st.session_state.get('user_id')
    if actor is None:raise PermissionError('Debes iniciar sesión para modificar tus datos.')
    username=normalized_username(username)
    full_name=str(full_name or '').strip()
    if not full_name:raise ValueError('El nombre completo es obligatorio.')
    if len(full_name)>150:raise ValueError('El nombre completo no debe exceder 150 caracteres.')
    if not current_password:raise ValueError('Ingresa tu contraseña actual para confirmar los cambios.')
    with connection() as con:
        current=con.execute('SELECT password_hash,active,username,full_name FROM users WHERE id=?',(actor,)).fetchone()
        if not current or not current['active']:
            raise PermissionError('La cuenta no está activa.')
        if not verify_password(current_password,current['password_hash']):
            raise ValueError('Contraseña actual incorrecta.')
        conflict=con.execute('SELECT id FROM users WHERE lower(username)=? AND id<>?',(username,actor)).fetchone()
        if conflict:raise ValueError('El nombre de usuario ya está registrado. Elige otro.')
        con.execute('UPDATE users SET username=?,full_name=? WHERE id=?',(username,full_name,actor))
        ensure_consulta_person(con,actor,full_name,username)
        con.execute('INSERT INTO audit(module,record_code,action) VALUES (?,?,?)',
                    ('Mi cuenta',username,'Datos personales actualizados'))
    st.session_state['username']=username
    st.session_state['full_name']=full_name
    return username


def consulta_project_ids(user_id):
    """IDs de proyectos que un usuario Consulta puede ver en Proyectos."""
    with connection() as con:
        return {int(r[0]) for r in con.execute(
            'SELECT project_id FROM consulta_project_access WHERE user_id=?', (int(user_id),))}


def consulta_project_codes(user_id):
    """Proyectos autorizados expresamente al usuario Consulta."""
    with connection() as con:
        rows=con.execute('''SELECT p.code FROM projects p
            JOIN consulta_project_access pa ON pa.project_id=p.id
            JOIN users u ON u.id=pa.user_id
            WHERE pa.user_id=? AND u.role='Consulta' AND u.active=1''',(int(user_id),)).fetchall()
    return {str(row['code']) for row in rows}


def consulta_delivery_codes(user_id):
    """Entregables con tarea asignada al usuario y proyecto autorizado."""
    with connection() as con:
        rows=con.execute('''SELECT DISTINCT c.code
            FROM deliverable_catalog c
            JOIN projects p ON p.code=c.project_code
            JOIN consulta_project_access pa ON pa.project_id=p.id AND pa.user_id=?
            JOIN tasks t ON t.project_code=c.project_code AND t.delivery_code=c.code
            JOIN task_people tp ON tp.task_id=t.id
            JOIN people pers ON pers.id=tp.person_id AND pers.user_id=?
            JOIN users u ON u.id=pers.user_id AND u.role='Consulta' AND u.active=1''',
            (int(user_id),int(user_id))).fetchall()
    return {str(row['code']) for row in rows}


def consulta_document_ids(user_id):
    """Acceso al checklist SOLO si hay asignación como revisor/responsable.

    Un documento con lista editada explícitamente usa esa lista; los documentos
    sin lista propia heredan las personas de tareas del mismo entregable.
    En ambos casos se exige autorización del proyecto.
    """
    with connection() as con:
        rows=con.execute('''SELECT DISTINCT dd.id
          FROM delivery_documents dd
          JOIN deliverables d ON d.id=dd.delivery_id
          JOIN projects p ON p.code=d.project_code
          JOIN consulta_project_access pa ON pa.project_id=p.id AND pa.user_id=?
          JOIN people pe ON pe.user_id=?
          JOIN users u ON u.id=pe.user_id AND u.active=1 AND u.role='Consulta'
          WHERE (
             (EXISTS(SELECT 1 FROM document_assignment_mode m WHERE m.document_id=dd.id)
               AND EXISTS(SELECT 1 FROM document_people dp
                    WHERE dp.document_id=dd.id AND dp.person_id=pe.id))
             OR
             (NOT EXISTS(SELECT 1 FROM document_assignment_mode m WHERE m.document_id=dd.id)
               AND EXISTS(SELECT 1 FROM tasks t
                    JOIN task_people tp ON tp.task_id=t.id
                    WHERE t.project_code=d.project_code AND t.delivery_code=d.code
                          AND tp.person_id=pe.id))
          )''',(int(user_id),int(user_id))).fetchall()
    return {int(row['id']) for row in rows}


def document_assignments(document_id):
    with connection() as con:
        customized=bool(con.execute('SELECT 1 FROM document_assignment_mode WHERE document_id=?',
                                    (int(document_id),)).fetchone())
        if customized:
            records=con.execute('SELECT person_id,assignment_role FROM document_people WHERE document_id=?',
                                (int(document_id),)).fetchall()
        else:
            records=con.execute('''SELECT DISTINCT tp.person_id,tp.assignment_role
                FROM delivery_documents doc JOIN deliverables d ON d.id=doc.delivery_id
                JOIN tasks t ON t.project_code=d.project_code AND t.delivery_code=d.code
                JOIN task_people tp ON tp.task_id=t.id
                WHERE doc.id=?''',(int(document_id),)).fetchall()
    return customized, ({int(r['person_id']) for r in records if r['assignment_role']=='Responsable'},
                        {int(r['person_id']) for r in records if r['assignment_role']=='Revisor'})


def save_document_assignments(document_id,owners,reviewers):
    require_admin()
    owners=set(map(int,owners));reviewers=set(map(int,reviewers))
    with connection() as con:
        if not con.execute('SELECT 1 FROM delivery_documents WHERE id=?',(int(document_id),)).fetchone():
            raise ValueError('El documento ya no existe.')
        allowed_ids=owners|reviewers
        if allowed_ids:
            slots=','.join('?' for _ in allowed_ids)
            existing={r[0] for r in con.execute(f'SELECT id FROM people WHERE id IN ({slots})',tuple(allowed_ids))}
            if existing!=allowed_ids:raise ValueError('Un integrante ya no existe.')
        con.execute('INSERT OR IGNORE INTO document_assignment_mode(document_id) VALUES(?)',(int(document_id),))
        con.execute('DELETE FROM document_people WHERE document_id=?',(int(document_id),))
        con.executemany('INSERT INTO document_people(document_id,person_id,assignment_role) VALUES (?,?,?)',
                        [(document_id,i,'Responsable') for i in owners]+[(document_id,i,'Revisor') for i in reviewers])
        con.execute('INSERT INTO audit(module,record_code,action) VALUES(?,?,?)',
                    ('Checklist',str(document_id),'Asignaciones de documento actualizadas'))

def set_consulta_project_access(con, user_id, project_ids):
    """Actualiza solo accesos de un usuario, validando rol e IDs en el servidor."""
    require_admin()
    user_id=int(user_id)
    ids={int(i) for i in project_ids}
    account=con.execute('SELECT role FROM users WHERE id=?',(user_id,)).fetchone()
    if not account:raise ValueError('El usuario seleccionado no existe.')
    if account['role']!='Consulta':
        # Los roles Cliente y Administrador no usan este permiso.
        con.execute('DELETE FROM consulta_project_access WHERE user_id=?',(user_id,))
        return
    if ids:
        placeholders=','.join('?' for _ in ids)
        existing={int(r[0]) for r in con.execute(
            f'SELECT id FROM projects WHERE id IN ({placeholders})', tuple(ids))}
        if existing!=ids:raise ValueError('Uno o más proyectos ya no existen; actualiza la selección.')
    con.execute('DELETE FROM consulta_project_access WHERE user_id=?',(user_id,))
    con.executemany('INSERT INTO consulta_project_access(user_id,project_id) VALUES(?,?)',
                    [(user_id,i) for i in sorted(ids)])
    con.execute('INSERT INTO audit(module,record_code,action) VALUES (?,?,?)',
                ('Usuarios',str(user_id),f'Proyectos visibles Consulta actualizados: {len(ids)}'))


def project_permission_options():
    """Retorna el catálogo de proyectos para el formulario del Administrador."""
    with connection() as con:
        return [dict(r) for r in con.execute(
            'SELECT id,code,name,client FROM projects ORDER BY code')]


def visible_project_rows(frame,user_id):
    """Filtra los proyectos antes de construir la vista de Consulta."""
    if frame.empty:return frame.copy()
    allowed=consulta_project_ids(user_id)
    return frame[frame['id'].isin(allowed)].copy()


def user_management():
    require_admin()
    st.title('🔐 Administración de usuarios')
    with connection() as con:
        users=pd.read_sql_query('SELECT id,username,full_name,role,active,created_at FROM users ORDER BY id',con)
    if not users.empty:
        display=users.rename(columns={'username':'Usuario','full_name':'Nombre completo',
            'role':'Rol','active':'Activo','created_at':'Fecha de creación'})
        st.dataframe(display[['Usuario','Nombre completo','Rol','Activo','Fecha de creación']],
                     hide_index=True,use_container_width=True)
    available_projects=project_permission_options()
    project_ids=[int(p['id']) for p in available_projects]
    project_labels={int(p['id']):f"{p['code']} · {p['name']}" + (f" · {p['client']}" if p.get('client') else '') for p in available_projects}
    with st.expander('➕ Crear usuario',expanded=users.empty):
        with st.form('new_user'):
            username=st.text_input('Usuario nuevo').strip().lower()
            name=st.text_input('Nombre completo').strip()
            role=st.selectbox('Rol',['Consulta','Cliente','Administrador'])
            password=st.text_input('Contraseña inicial (mínimo 6 caracteres)',type='password')
            initial_projects=st.multiselect('Proyectos visibles (aplica solo a Consulta)',project_ids,
                format_func=lambda i:project_labels[i],key='new_consulta_project_grants',
                help='Selecciona uno o varios proyectos. Si no seleccionas ninguno, Consulta no verá proyectos en ese apartado.')
            if st.form_submit_button('Crear usuario',type='primary'):
                try:
                    username=normalized_username(username)
                    if not name:raise ValueError('El nombre completo es obligatorio.')
                    if len(password)<6:raise ValueError('La contraseña debe tener al menos 6 caracteres.')
                    with connection() as con:
                        cursor=con.execute('INSERT INTO users(username,full_name,password_hash,role) VALUES(?,?,?,?)',
                                    (username,name,hash_password(password),role))
                        if role=='Consulta':
                            ensure_consulta_person(con,cursor.lastrowid,name,username)
                        set_consulta_project_access(con,cursor.lastrowid,initial_projects)
                    st.success('Usuario creado.');st.rerun()
                except sqlite3.IntegrityError:st.error('El nombre de usuario ya existe.')
                except ValueError as exc:st.error(str(exc))
    if not users.empty:
        # Edición plegable mediante un único botón. Se vuelve a consultar la cuenta
        # seleccionada al guardar; nunca se confía en un rol recibido desde el navegador.
        if st.button('✏️ Editar usuario',use_container_width=True,key='open_user_editor'):
            st.session_state['show_user_editor']=not st.session_state.get('show_user_editor',False)
        if st.session_state.get('show_user_editor',False):
            with st.container(border=True):
                selected=st.selectbox('Seleccionar usuario',users.id.astype(int).tolist(),
                    format_func=lambda i: f"{users.loc[users.id==i,'full_name'].iloc[0]} · {users.loc[users.id==i,'username'].iloc[0]}",
                    key='user_edit_select')
                r=users.loc[users.id==selected].iloc[0]
                selected_projects=consulta_project_ids(selected)
                with st.form(f'edit_user_form_{selected}'):
                    new_username=st.text_input('Nombre de usuario',value=str(r['username']))
                    new_full_name=st.text_input('Nombre completo',value=str(r['full_name']))
                    allowed_roles=['Administrador','Consulta','Cliente']
                    new_role=st.selectbox('Rol',allowed_roles,index=allowed_roles.index(str(r['role'])))
                    active=st.checkbox('Cuenta activa',value=bool(r['active']))
                    reset=st.text_input('Nueva contraseña (opcional; dejar vacío para conservar)',type='password')
                    edited_projects=st.multiselect('Proyectos que este usuario Consulta puede visualizar',project_ids,
                        default=[i for i in project_ids if i in selected_projects],
                        format_func=lambda i:project_labels[i],key=f'edit_consulta_project_grants_{selected}',
                        help='Solo aplica si el rol es Consulta. Sin proyectos asignados, el apartado Proyectos estará vacío.')
                    save=st.form_submit_button('💾 Guardar cambios',type='primary')
                if save:
                    try:
                        updated=update_user_by_admin(selected,new_username,new_full_name,new_role,active,reset,
                                                     project_ids=edited_projects)
                        st.success(f'Usuario {updated} actualizado.');st.rerun()
                    except (ValueError,PermissionError,sqlite3.Error) as exc:st.error(str(exc))
        st.divider()
        with st.expander('🗑️ Eliminar usuario'):
            st.warning('Se eliminará la cuenta y sus autorizaciones; no se borrarán proyectos, entregables ni documentos en Google Drive.')
            candidates=[int(i) for i in users.id.tolist() if int(i)!=int(st.session_state['user_id'])]
            if not candidates:
                st.info('No hay otros usuarios que se puedan eliminar.')
            else:
                target=st.selectbox('Usuario a eliminar',candidates,key='user_delete_select',
                    format_func=lambda i: f"{users.loc[users.id==i,'full_name'].iloc[0]} · {users.loc[users.id==i,'username'].iloc[0]}")
                target_username=str(users.loc[users.id==target,'username'].iloc[0])
                with st.form(f'user_delete_confirm_{target}'):
                    typed=st.text_input(f'Escribe el usuario «{target_username}» para confirmar')
                    confirmed=st.checkbox('Comprendo que esta cuenta dejará de tener acceso inmediatamente')
                    clicked=st.form_submit_button('🗑️ Eliminar usuario definitivamente',type='primary')
                if clicked:
                    if not confirmed or typed.strip()!=target_username:
                        st.error('Escribe el nombre de usuario exacto y marca la confirmación.')
                    else:
                        try:
                            removed=delete_user_account(target)
                            st.success(f'Usuario {removed} eliminado. Sus sesiones se revocaron.');st.rerun()
                        except (ValueError,sqlite3.Error) as exc:st.error(str(exc))
    client_permissions_ui()

def client_permissions_ui():
    """ALTIVIA: Cliente -> Proyecto -> Entregables autorizados, sin borrar otros proyectos."""
    require_admin()
    st.subheader('📂 Autorización de entregables para clientes')
    st.caption('El cliente solo verá entregables Aprobados o Entregados que ALTIVIA autorice expresamente.')
    with connection() as con:
        clients=[dict(r) for r in con.execute("SELECT id,username,full_name FROM users WHERE role='Cliente' AND active=1 ORDER BY full_name")]
        eligible=[dict(r) for r in con.execute('''SELECT d.id,d.code,d.name,d.project_code,d.status,p.name AS project_name
           FROM deliverables d LEFT JOIN projects p ON p.code=d.project_code
           WHERE d.status IN ('Aprobado','Entregado') ORDER BY d.project_code,d.name''')]
    if not clients:
        st.info('Primero crea un usuario con el rol Cliente.');return
    # Todos los filtros se vuelven a verificar contra el servidor al guardar.
    project_names={r['project_code']:(r['project_name'] or r['project_code']) for r in eligible}
    with st.container(border=True):
        c1,c2=st.columns(2)
        with c1:
            st.markdown('**01 · Cliente**')
            cid=st.selectbox('Cliente',[r['id'] for r in clients],
                format_func=lambda i:next(r['full_name']+' · '+r['username'] for r in clients if r['id']==i),key='client_acl_user')
        if not project_names:
            st.info('No existen entregables aprobados o entregados para autorizar.');return
        with c2:
            st.markdown('**02 · Proyecto**')
            proj=st.selectbox('Proyecto',sorted(project_names),
                format_func=lambda c:f'{c} – {project_names[c]}',key='client_acl_project')
        with connection() as con:
            current={int(r[0]) for r in con.execute('SELECT delivery_id FROM client_access WHERE user_id=?',(cid,))}
        options={r['id']:f"{r['code']} · {r['name']} ({r['status']})" for r in eligible if r['project_code']==proj}
        st.markdown('**03 · Entregables autorizados**')
        st.caption('Solo se listan los entregables del proyecto seleccionado. Los permisos de otros proyectos no se modificarán.')
        selected=st.multiselect('Entregables autorizados',list(options),
            default=[i for i in options if i in current],format_func=lambda i:options[i],key=f'client_acl_docs_{cid}_{proj}')
        if st.button('💾 Guardar autorizaciones de este proyecto',type='primary',key='save_client_grants'):
            valid_ids=set(options)
            if not set(selected).issubset(valid_ids):
                st.error('La selección contiene entregables no autorizables.');return
            with connection() as con:
                active=con.execute("SELECT role,active FROM users WHERE id=?",(cid,)).fetchone()
                if not active or active['role']!='Cliente' or not active['active']:
                    st.error('El usuario ya no está habilitado como Cliente.');return
                con.execute('''DELETE FROM client_access WHERE user_id=? AND delivery_id IN
                    (SELECT id FROM deliverables WHERE project_code=?)''',(cid,proj))
                con.executemany('INSERT INTO client_access(user_id,delivery_id) VALUES (?,?)',[(cid,int(i)) for i in selected])
                con.execute('INSERT INTO audit(module,record_code,action) VALUES (?,?,?)',
                    ('Clientes',proj,f'Accesos actualizados para usuario {cid}'))
            st.success('Permisos del proyecto guardados; se conservaron los demás proyectos.');st.rerun()


def internal_pdf_preview(data):
    if hasattr(st,'pdf'):
        st.pdf(data,height=680)
    else:
        import streamlit.components.v1 as components
        payload=base64.b64encode(data).decode('ascii')
        components.html('<iframe title="Plano PDF" src="data:application/pdf;base64,'+payload+'" style="width:100%;height:680px;border:0" loading="lazy"></iframe>',height=700,scrolling=False)
        st.caption('Si tu navegador bloquea el visor integrado, utiliza Descargar PDF.')


MAX_UPLOAD_MB = 15
MAX_UPLOAD_BYTES = MAX_UPLOAD_MB * 1024 * 1024

def file_info(delivery_id, version):
    with connection() as con:
        rows=con.execute('SELECT kind,filename,uploaded_at,length(content) AS size FROM delivery_files WHERE delivery_id=? AND version=?',(delivery_id,version)).fetchall()
    return {r['kind']:dict(r) for r in rows}

def file_bytes(delivery_id,version,kind):
    with connection() as con:
        r=con.execute('SELECT filename,content FROM delivery_files WHERE delivery_id=? AND version=? AND kind=?',(delivery_id,version,kind)).fetchone()
    return (r['filename'],bytes(r['content'])) if r else None

def upload_delivery_files(delivery_id,code,version):
    require_admin()
    st.markdown('**📎 Archivos de esta versión**')
    existing=file_info(delivery_id,version)
    pdf=st.file_uploader('PDF para visualizar',type=['pdf'],key=f'pdf_{delivery_id}_{version}')
    editable=st.file_uploader('Archivo editable para descargar (DWG, DOC o DOCX)',type=['dwg','doc','docx'],key=f'edit_{delivery_id}_{version}')
    st.caption(f'Máximo {MAX_UPLOAD_MB} MB por archivo. Se guardan en la base de datos local; no subas documentos confidenciales a Streamlit Community Cloud.')
    for kind,title in [('pdf','PDF'),('editable','Editable')]:
        if kind in existing:st.caption(f'{title} actual: {existing[kind]["filename"]} · {existing[kind]["size"] / 1048576:.1f} MB')
    if st.button('💾 Guardar archivos de esta versión',key=f'save_files_{delivery_id}_{version}'):
        pending=[]
        for kind,f in [('pdf',pdf),('editable',editable)]:
            if f is None:continue
            raw=f.getvalue()
            if not raw or len(raw)>MAX_UPLOAD_BYTES:
                st.error(f'{f.name}: archivo vacío o superior a {MAX_UPLOAD_MB} MB.');return
            if kind=='pdf' and not raw.startswith(b'%PDF-'):
                st.error('El PDF no tiene una cabecera válida.');return
            ext=f.name.rsplit('.',1)[-1].lower() if '.' in f.name else ''
            if (kind=='pdf' and ext!='pdf') or (kind=='editable' and ext not in ('dwg','doc','docx')):
                st.error('Formato no permitido.');return
            pending.append((kind,f.name,raw))
        if not pending:st.info('Selecciona al menos un archivo para guardar.');return
        with connection() as con:
            for kind,name,raw in pending:
                con.execute("""INSERT INTO delivery_files(delivery_id,version,kind,filename,content) VALUES (?,?,?,?,?)
                    ON CONFLICT(delivery_id,version,kind) DO UPDATE SET filename=excluded.filename,content=excluded.content,uploaded_at=CURRENT_TIMESTAMP""",
                    (delivery_id,version,kind,name,raw))
            con.execute('INSERT INTO audit(module,record_code,action) VALUES (?,?,?)',('Archivos',code,'Archivos actualizados versión '+version))
        st.success('Archivos guardados.');st.rerun()

# Los documentos se guardan exclusivamente en Drive, no como BLOB local.
DRIVE_FOLDER_ID = '19sEVR8-vrm9m_8adHJVU_l1gYCYIz4Sq'
DRIVE_LIMIT = 25 * 1024 * 1024

def drive_credentials():
    try:
        cfg=st.secrets.get('google_drive',{})
        return {k:str(cfg.get(k,'')).strip() for k in ('client_id','client_secret','refresh_token')}
    except Exception:
        return {}

def drive_ready():
    cfg=drive_credentials()
    return all(cfg.values())

def drive_token():
    cfg=drive_credentials()
    if not all(cfg.values()):
        raise ValueError('Google Drive todavía no está autorizado en Secrets. No se subió ningún archivo.')
    resp=requests.post('https://oauth2.googleapis.com/token',data={
        'client_id':cfg['client_id'],'client_secret':cfg['client_secret'],
        'refresh_token':cfg['refresh_token'],'grant_type':'refresh_token'},timeout=20)
    if not resp.ok:raise ValueError(f'No se pudo autenticar Google Drive (HTTP {resp.status_code}).')
    return resp.json()['access_token']

def drive_safe_folder_name(value):
    """Nombre legible, estable y sin separadores de ruta."""
    name=re.sub(r'[\\/:*?"<>|\x00-\x1f]', '_', str(value or '').strip())
    return name[:135] or 'Sin nombre'


def drive_find_folder(token,parent_id,folder_name):
    """Busca solamente carpetas accesibles a esta credencial, sin crearlas."""
    name=drive_safe_folder_name(folder_name)
    escaped=name.replace('\\','\\\\').replace("'", "\\'")
    q=f"name = '{escaped}' and '{parent_id}' in parents and mimeType = 'application/vnd.google-apps.folder' and trashed = false"
    response=requests.get('https://www.googleapis.com/drive/v3/files',
        params={'q':q,'fields':'nextPageToken,files(id,name)','pageSize':100},
        headers={'Authorization':f'Bearer {token}'},timeout=30)
    if not response.ok:
        raise ValueError(f'No se pudo localizar la carpeta en Drive (HTTP {response.status_code}).')
    results=response.json().get('files',[])
    return results[0]['id'] if results else None


def drive_rename_item(token,file_id,new_name):
    response=requests.patch('https://www.googleapis.com/drive/v3/files/'+quote(file_id,safe=''),
        headers={'Authorization':f'Bearer {token}'},json={'name':new_name},
        params={'fields':'id,name'},timeout=45)
    if not response.ok:
        raise ValueError(f'Google Drive no permitió cambiar el nombre (HTTP {response.status_code}).')


def drive_find_or_create_folder(token,parent_id,folder_name):
    """Reutiliza carpetas por nombre y padre; nunca crea una carpeta por cada guardado."""
    name=drive_safe_folder_name(folder_name)
    headers={'Authorization':f'Bearer {token}'}
    escaped=name.replace('\\','\\\\').replace("'", "\\'")
    q=f"name = '{escaped}' and '{parent_id}' in parents and mimeType = 'application/vnd.google-apps.folder' and trashed = false"
    response=requests.get('https://www.googleapis.com/drive/v3/files',
        params={'q':q,'fields':'nextPageToken,files(id,name)','page_size':100},headers=headers,timeout=30)
    if not response.ok:
        raise ValueError(f'No se pudo buscar carpeta en Google Drive (HTTP {response.status_code}): {response.text[:200]}')
    matches=response.json().get('files',[])
    if matches:return matches[0]['id']
    created=requests.post('https://www.googleapis.com/drive/v3/files',headers=headers,
        json={'name':name,'mimeType':'application/vnd.google-apps.folder','parents':[parent_id]},
        params={'fields':'id,name'},timeout=30)
    if not created.ok:
        raise ValueError(f'No se pudo crear carpeta en Google Drive (HTTP {created.status_code}): {created.text[:200]}')
    return created.json()['id']


def drive_delivery_folder(token,delivery_id,code,version):
    """Carpeta raíz / proyecto / entregable / versión."""
    with connection() as con:
        record=con.execute("""SELECT d.project_code,p.name AS project_name,d.name AS delivery_name
                              FROM deliverables d LEFT JOIN projects p ON p.code=d.project_code
                              WHERE d.id=?""",(delivery_id,)).fetchone()
    if record is None:raise ValueError('No existe el entregable para vincular archivos.')
    project=record['project_code'] or 'PROYECTO'
    project_name=(record['project_name'] or '').strip()
    delivery_name=(record['delivery_name'] or '').strip()
    project_folder=drive_safe_folder_name(project + (' - '+project_name if project_name else ''))
    delivery_folder=drive_safe_folder_name(code + (' - '+delivery_name if delivery_name else ''))
    project_id=drive_find_or_create_folder(token,DRIVE_FOLDER_ID,project_folder)
    deliverable_id=drive_find_or_create_folder(token,project_id,delivery_folder)
    return drive_find_or_create_folder(token,deliverable_id,version)


def drive_upload(upload,delivery_id,code,version,kind,allow_new_consulta=False):
    if not can_edit():
        if not (allow_new_consulta and st.session_state.get('role')=='Consulta'):
            raise PermissionError('Solo el administrador puede modificar archivos de versiones existentes.')
    if upload is None:return
    raw=upload.getvalue()
    if not raw or len(raw)>DRIVE_LIMIT:raise ValueError('El archivo está vacío o supera los 25 MB.')
    ext=upload.name.rsplit('.',1)[-1].lower() if '.' in upload.name else ''
    if (kind=='pdf' and (ext!='pdf' or not raw.startswith(b'%PDF-'))) or (kind=='editable' and ext not in ('dwg','doc','docx')):
        raise ValueError('Formato de archivo no permitido.')
    mime= 'application/pdf' if kind=='pdf' else 'application/octet-stream'
    safe_code=re.sub(r'[^A-Za-z0-9_-]','_',code)
    with connection() as con:
        exists=con.execute('SELECT 1 FROM deliverables WHERE id=? AND code=?',(delivery_id,code)).fetchone()
    if not exists:raise ValueError('No se encontró el entregable registrado.')
    # NUNCA anteponer el proyecto ni su nombre: el ID Entregable ya es completo.
    # Ej.: 202607_ARQ-001_V01.pdf y 202607_ARQ-001_V01.dwg
    filename=f'{safe_code}_{version}.{ext}'
    token=drive_token()
    version_folder_id=drive_delivery_folder(token,delivery_id,code,version)
    metadata={'name':filename,'parents':[version_folder_id], 'description':f'ALTIVIA {code} {version} {kind}'}
    # Subida multipart; el token corresponde a la cuenta que posee la carpeta.
    resp=requests.post('https://www.googleapis.com/upload/drive/v3/files?uploadType=multipart&fields=id,name',
        headers={'Authorization':f'Bearer {token}'},
        files={'metadata':('metadata',json.dumps(metadata),'application/json; charset=UTF-8'),
               'file':(filename,raw,mime)},timeout=90)
    if not resp.ok:raise ValueError(f'Google Drive rechazó la carga ({resp.status_code}): {resp.text[:250]}')
    file_id=resp.json()['id']
    with connection() as con:
        old=con.execute('SELECT file_id FROM drive_files WHERE delivery_id=? AND version=? AND kind=?',
            (delivery_id,version,kind)).fetchone()
        con.execute('''INSERT INTO drive_files(delivery_id,version,kind,file_id,filename)
            VALUES(?,?,?,?,?) ON CONFLICT(delivery_id,version,kind)
            DO UPDATE SET file_id=excluded.file_id,filename=excluded.filename,uploaded_at=CURRENT_TIMESTAMP''',
            (delivery_id,version,kind,file_id,filename))
        con.execute('INSERT INTO audit(module,record_code,action) VALUES(?,?,?)',
            ('Drive',code,f'Archivo {kind} subido en {version}'))
    # No se elimina automáticamente el fichero anterior: protege el historial.
    return filename

def drive_file_info(delivery_id,version):
    with connection() as con:
        return {r['kind']:dict(r) for r in con.execute('SELECT * FROM drive_files WHERE delivery_id=? AND version=?',(delivery_id,version)).fetchall()}

def drive_bytes(file_id):
    token=drive_token()
    response=requests.get(f'https://www.googleapis.com/drive/v3/files/{file_id}',params={'alt':'media'},
       headers={'Authorization':f'Bearer {token}'},timeout=75,stream=True)
    if not response.ok:raise ValueError('No se pudo recuperar el documento autorizado desde Google Drive.')
    chunks=[]; size=0
    for chunk in response.iter_content(262144):
        size+=len(chunk)
        if size>DRIVE_LIMIT:raise ValueError('El documento supera los 25 MB.')
        chunks.append(chunk)
    return b''.join(chunks)


def _download_drive_bytes(file_id, credentials, *, validate_pdf=False):
    """Descarga diferida segura: no usa st.session_state ni st.secrets en el hilo del botón.

    Las credenciales se capturan del servidor al crear el botón; nunca se envían al cliente.
    Se solicita un token OAuth nuevo por descarga para evitar tokens caducados.
    """
    if not all(credentials.get(k) for k in ('client_id','client_secret','refresh_token')):
        raise ValueError('Google Drive no está conectado.')
    response=requests.post('https://oauth2.googleapis.com/token',data={
        'client_id':credentials['client_id'],
        'client_secret':credentials['client_secret'],
        'refresh_token':credentials['refresh_token'],
        'grant_type':'refresh_token'},timeout=20)
    if not response.ok:
        raise ValueError('No se pudo iniciar la descarga desde Google Drive.')
    token=response.json()['access_token']
    # Evitar URLs basadas en datos de usuario y seguir la API oficial de Drive.
    with requests.get('https://www.googleapis.com/drive/v3/files/'+quote(file_id,safe=''),
        params={'alt':'media'},headers={'Authorization':f'Bearer {token}'},
        timeout=90,stream=True) as response:
        if not response.ok:
            raise ValueError('No se pudo recuperar el archivo desde Google Drive.')
        chunks=[]
        length=0
        for block in response.iter_content(262144):
            if not block:continue
            length+=len(block)
            if length>DRIVE_LIMIT:
                raise ValueError('El archivo supera el límite de 25 MB.')
            chunks.append(block)
    data=b''.join(chunks)
    if not data:raise ValueError('El archivo descargado está vacío.')
    if validate_pdf and not data.startswith(b'%PDF-'):
        raise ValueError('El archivo recuperado no es un PDF válido.')
    return data


def direct_document_download(document, kind):
    """Proporciona la función diferida al botón de descarga, sin paso Preparar."""
    files=document.get('files',{})
    remote=files.get(kind)
    if remote:
        credentials=drive_credentials()
        if not all(credentials.values()):return None,remote['filename']
        file_id=remote['file_id']
        # Importante: callback sin comandos Streamlit; se ejecuta fuera del script.
        def on_demand():
            return _download_drive_bytes(file_id, credentials, validate_pdf=(kind=='pdf'))
        return on_demand,remote['filename']
    # Compatibilidad con documentos heredados que todavía están en SQLite.
    legacy=legacy_local_document(document['id'],kind)
    if legacy:
        filename,raw=legacy
        def on_demand_legacy():
            return raw
        return on_demand_legacy,filename
    return None,None


def docs_for_version(delivery_id, version):
    with connection() as con:
        docs=[dict(r) for r in con.execute(
            'SELECT id,doc_no,title FROM delivery_documents WHERE delivery_id=? AND version=? ORDER BY id',
            (delivery_id,version)).fetchall()]
        for doc in docs:
            doc['files']={r['kind']:dict(r) for r in con.execute(
                'SELECT kind,file_id,filename,uploaded_at FROM document_files WHERE document_id=?',(doc['id'],)).fetchall()}
    return docs


def register_document(delivery_id, version, title, actor_role=None):
    role=actor_role or st.session_state.get('role')
    if role not in ('Administrador','Consulta'):
        raise PermissionError('No tienes permiso para registrar documentos.')
    title=str(title or '').strip()
    if not title:raise ValueError('Cada documento debe tener un nombre descriptivo.')
    with connection() as con:
        if role=='Consulta':
            row=con.execute('''SELECT d.code,d.project_code FROM deliverables d WHERE d.id=? AND d.version=?''',
                            (int(delivery_id),str(version))).fetchone()
            if not row or row['code'] not in consulta_delivery_codes(st.session_state['user_id']):
                raise PermissionError('No puedes agregar documentos a una versión no autorizada.')
        count=con.execute('SELECT highest FROM document_serials WHERE delivery_id=? AND version=?',
                          (delivery_id,version)).fetchone()
        n=int(count['highest'])+1 if count else 1
        doc_no=f'D{n:02d}'
        cursor=con.execute('INSERT INTO delivery_documents(delivery_id,version,doc_no,title) VALUES(?,?,?,?)',
                           (delivery_id,version,doc_no,title))
        con.execute('''INSERT INTO document_serials(delivery_id,version,highest) VALUES(?,?,?)
            ON CONFLICT(delivery_id,version) DO UPDATE SET highest=MAX(highest,excluded.highest)''',
                    (delivery_id,version,n))
        con.execute('UPDATE deliverables SET last_modified_at=? WHERE id=?',
                    (peru_timestamp(),int(delivery_id)))
        return cursor.lastrowid,doc_no


def upload_document_file(upload, document_id, kind, allow_consulta=False):
    if not can_edit() and not (allow_consulta and st.session_state.get('role')=='Consulta'):
        raise PermissionError('Solo el Administrador puede reemplazar documentos existentes.')
    if upload is None:return None
    with connection() as con:
        row=con.execute('''SELECT doc.id,doc.doc_no,doc.version,d.code,d.id AS delivery_id
             FROM delivery_documents doc JOIN deliverables d ON doc.delivery_id=d.id
             WHERE doc.id=?''',(document_id,)).fetchone()
    if not row:raise ValueError('Documento no encontrado.')
    raw=upload.getvalue()
    ext=upload.name.rsplit('.',1)[-1].lower() if '.' in upload.name else ''
    if not raw or len(raw)>DRIVE_LIMIT:raise ValueError('Archivo vacío o mayor de 25 MB.')
    if kind=='pdf' and (ext!='pdf' or not raw.startswith(b'%PDF-')):
        raise ValueError('Selecciona un PDF válido.')
    if kind=='editable' and ext not in ('dwg','doc','docx','xlsx','xls','ifc','rvt'):
        raise ValueError('Editable no admitido. Usa DWG, DOC, DOCX, XLSX, XLS, IFC o RVT.')
    safe=re.sub(r'[^A-Za-z0-9_-]','_',row['code'])
    filename=f'{safe}_{row["version"]}_{row["doc_no"]}.{ext}'
    folder=drive_delivery_folder(drive_token(),row['delivery_id'],row['code'],row['version'])
    token=drive_token()
    meta={'name':filename,'parents':[folder], 'description':f'GP Altivia {row["code"]} {row["version"]} {row["doc_no"]}'}
    mime='application/pdf' if kind=='pdf' else 'application/octet-stream'
    response=requests.post('https://www.googleapis.com/upload/drive/v3/files?uploadType=multipart&fields=id,name',
        headers={'Authorization':f'Bearer {token}'},files={
            'metadata':('metadata',json.dumps(meta),'application/json; charset=UTF-8'),
            'file':(filename,raw,mime)},timeout=90)
    if not response.ok:
        raise ValueError(f'Google Drive no pudo subir {filename} (HTTP {response.status_code}): {response.text[:200]}')
    file_id=response.json()['id']
    with connection() as con:
        con.execute('''INSERT INTO document_files(document_id,kind,file_id,filename)
            VALUES (?,?,?,?) ON CONFLICT(document_id,kind)
            DO UPDATE SET file_id=excluded.file_id,filename=excluded.filename,uploaded_at=CURRENT_TIMESTAMP''',
            (document_id,kind,file_id,filename))
        con.execute('UPDATE deliverables SET last_modified_at=? WHERE id=?',
                    (peru_timestamp(),int(row['delivery_id'])))
        con.execute('INSERT INTO audit(module,record_code,action) VALUES(?,?,?)',
            ('Documentos',row['code'],f'{row["version"]}/{row["doc_no"]} - {kind} subido'))
    # Si se sustituyó un archivo, el anterior permanece en Drive como respaldo.
    return filename


def attach_new_documents(delivery_id,version,entries,consulta=False):
    """Cada entrada (título, PDF opcional, editable opcional) crea un DNN único."""
    result=[]
    for title,pdf,editable in entries:
        if not (str(title or '').strip() or pdf or editable):continue
        doc_id,number=register_document(delivery_id,version,title)
        result.append(number)
        for kind,item in (('pdf',pdf),('editable',editable)):
            if item:upload_document_file(item,doc_id,kind,allow_consulta=consulta)
    return result


def legacy_local_document(document_id,kind):
    """Lectura transitoria para archivos antiguos conservados como BLOB en SQLite."""
    with connection() as con:
        doc=con.execute('SELECT delivery_id,version,doc_no FROM delivery_documents WHERE id=?',(document_id,)).fetchone()
    if doc and doc['doc_no']=='D01':
        return file_bytes(doc['delivery_id'],doc['version'],kind)
    return None


def document_preview_ui(document, *, allow_download=True, allow_editable=True, prefix='doc'):
    """Vista previa y descargas de un clic: PDF en toda versión, editable solo vigente.

    Streamlit 1.52+ admite data=callable para descargar el archivo bajo demanda.
    Así se evita el botón «Preparar» y la precarga de todos los archivos de Drive.
    """
    doc_id=int(document['id'])
    files=document.get('files',{})
    legacy_pdf=legacy_local_document(doc_id,'pdf') if 'pdf' not in files else None
    legacy_edit=legacy_local_document(doc_id,'editable') if 'editable' not in files else None
    has_pdf=('pdf' in files or legacy_pdf is not None)
    has_edit=('editable' in files or legacy_edit is not None)
    cols=st.columns(3 if allow_download else 1,gap='small')
    with cols[0]:
        if has_pdf:
            if st.button('👁 Visualizar PDF',key=f'{prefix}_view_{doc_id}',use_container_width=True):
                key=f'{prefix}_open_{doc_id}'
                st.session_state[key]=not st.session_state.get(key,False)
        else:
            st.caption('Sin PDF')
    if allow_download:
        with cols[1]:
            if has_pdf:
                pdf_data,pdf_filename=direct_document_download(document,'pdf')
                st.download_button('⬇ Descargar PDF',data=pdf_data or b'',
                    file_name=pdf_filename,mime='application/pdf',
                    disabled=pdf_data is None,on_click='ignore',
                    key=f'{prefix}_download_pdf_{doc_id}',use_container_width=True)
                if pdf_data is None:st.caption('Google Drive no está conectado.')
            else:
                st.caption('Sin PDF')
        with cols[2]:
            if allow_editable and has_edit:
                edit_data,edit_filename=direct_document_download(document,'editable')
                st.download_button('⬇ Descargar editable',data=edit_data or b'',
                    file_name=edit_filename,mime='application/octet-stream',
                    disabled=edit_data is None,on_click='ignore',
                    key=f'{prefix}_download_edit_{doc_id}',use_container_width=True)
                if edit_data is None:st.caption('Google Drive no está conectado.')
            else:
                st.caption('Editable solo en versión más reciente' if has_edit and not allow_editable else 'Sin editable')
    if st.session_state.get(f'{prefix}_open_{doc_id}') and has_pdf:
        try:
            pdf=drive_bytes(files['pdf']['file_id']) if 'pdf' in files else legacy_pdf[1]
            if not pdf.startswith(b'%PDF-'):raise ValueError('El PDF no es válido.')
            internal_pdf_preview(pdf)
        except Exception as exc:
            st.error(f'No se pudo visualizar el PDF: {exc}')

def next_version(code,current=None):
    """No reutiliza números de versiones eliminadas; registra el máximo histórico."""
    versions=[str(current)] if current else []
    historic=0
    if code:
        with connection() as con:
            versions += [r[0] for r in con.execute('SELECT version FROM versions WHERE delivery_code=?',(code,))]
            saved=con.execute('SELECT highest FROM version_counters WHERE delivery_code=?',(code,)).fetchone()
            historic=int(saved['highest']) if saved else 0
    n=max([historic]+[int(m.group(1)) for v in versions if (m:=re.fullmatch(r'V(\d+)',str(v).upper()))])
    return f'V{n+1:02d}'


def client_portal():
    # Cliente: únicamente documentos de la versión vigente de entregables aprobados/autorizados.
    if st.session_state.get('role')!='Cliente':
        st.error('Acceso restringido.');st.stop()
    with connection() as con:
        records=[dict(r) for r in con.execute('''SELECT d.id,d.code,d.project_code,d.name,d.version,
            d.specialty,d.status,p.name AS project_name
            FROM deliverables d JOIN client_access a ON a.delivery_id=d.id
            LEFT JOIN projects p ON p.code=d.project_code
            WHERE a.user_id=? AND d.status IN ('Aprobado','Entregado')
            ORDER BY d.project_code,d.name''',(st.session_state['user_id'],))]
    if not records:
        st.title('📁 Mis proyectos')
        st.info('Todavía no tienes entregables aprobados autorizados.');return
    projects={r['project_code']:r['project_name'] or r['project_code'] for r in records}
    cols=st.columns([1.2,2,1.2],gap='small')
    with cols[0]:
        st.caption('01 · CLIENTE')
        st.markdown('**'+escape(str(st.session_state.get('full_name') or 'Cliente'))+'**')
    with cols[1]:
        st.caption('02 · PROYECTO')
        chosen=st.selectbox('Proyecto autorizado',sorted(projects),
            format_func=lambda c:f'{c} – {projects[c]}',label_visibility='collapsed',key='client_project_choice')
    visible=[r for r in records if r['project_code']==chosen]
    with cols[2]:
        st.caption('03 · ENTREGABLES AUTORIZADOS')
        st.metric('Disponibles',len(visible),label_visibility='collapsed')
    st.title('📁 '+projects[chosen])
    st.caption('Planos y documentos aprobados · Acceso de solo lectura')
    query=st.text_input('🔎 Buscar entregable o documento',key='client_search').strip().casefold()
    visible_docs=0
    for item in visible:
        docs=docs_for_version(item['id'],item['version'])
        if query:
            matches_delivery=query in ' '.join(str(item.get(k) or '') for k in ('code','name','specialty','version')).casefold()
            if not matches_delivery:
                docs=[d for d in docs if query in (d['doc_no']+' '+d['title']).casefold()]
            if not matches_delivery and not docs:continue
        with st.container(border=True):
            st.markdown(f'### {escape(item["code"])} – {escape(item["name"])}')
            st.caption(f'{item["specialty"] or ""} · {item["version"]} · {len(docs)} documento(s)')
            st.success(item['status'],icon='✅')
            if not docs:
                st.caption('Este entregable aún no tiene documentos disponibles.')
            for document in docs:
                visible_docs+=1
                with st.container(border=True):
                    st.markdown(f'**📄 {escape(document["doc_no"])} — {escape(document["title"])}**')
                    document_preview_ui(document,allow_download=True,prefix='client')
    if not visible_docs:st.info('No hay documentos que coincidan con la búsqueda.')

def my_account():
    st.title('👤 Mi cuenta')
    uid=st.session_state.get('user_id')
    with connection() as con:
        current=con.execute('SELECT username,full_name,role,active FROM users WHERE id=?',(uid,)).fetchone()
    if not current or not current['active']:
        st.error('No se pudo verificar una cuenta activa. Inicia sesión de nuevo.');return
    st.caption(f"Rol: {current['role']} · Usuario: {current['username']}")

    st.subheader('Editar mis datos')
    with st.form('edit_own_profile'):
        new_username=st.text_input('Nombre de usuario',value=current['username'])
        new_full_name=st.text_input('Nombre completo',value=current['full_name'])
        st.caption('Para proteger tu cuenta, confirma los cambios con tu contraseña actual. Tu rol y permisos no cambiarán.')
        profile_password=st.text_input('Contraseña actual para confirmar',type='password')
        save_profile=st.form_submit_button('💾 Guardar mis datos',type='primary')
    if save_profile:
        try:
            update_own_profile(new_username,new_full_name,profile_password)
            st.success('Tus datos se actualizaron. Puedes seguir usando GP Altivia sin cerrar sesión.');st.rerun()
        except (ValueError,PermissionError,sqlite3.Error) as exc:st.error(str(exc))

    st.divider()
    st.subheader('🔒 Cambiar contraseña')
    with st.form('change_password'):
        current_password=st.text_input('Contraseña actual',type='password')
        new=st.text_input('Nueva contraseña (mínimo 6 caracteres)',type='password')
        confirm=st.text_input('Confirmar contraseña',type='password')
        if st.form_submit_button('Cambiar contraseña'):
            with connection() as con:
                stored=con.execute('SELECT password_hash FROM users WHERE id=?',(uid,)).fetchone()
                if not stored or not verify_password(current_password,stored[0]):
                    st.error('Contraseña actual incorrecta.')
                elif len(new)<6 or new!=confirm:
                    st.error('La nueva contraseña no cumple los requisitos o no coincide.')
                else:
                    con.execute('UPDATE users SET password_hash=? WHERE id=?',(hash_password(new),uid))
                    con.execute('DELETE FROM remembered_sessions WHERE user_id=?',(uid,))
                    st.session_state.pop('_remember_token',None)
                    try:delete_remember_cookie()
                    except Exception:pass
                    st.success('Contraseña actualizada. Los accesos recordados quedaron revocados; al recargar deberás ingresar nuevamente.')

def df(table):
    with connection() as con: return pd.read_sql_query(f'SELECT * FROM {table} ORDER BY id DESC',con)

def datespan(value):
    try: return (date.fromisoformat(str(value)[:10])-date.today()).days
    except (ValueError,TypeError): return None

def _person_task_ids(person_id):
    with connection() as con:
        return con.execute('SELECT DISTINCT task_id FROM task_people WHERE person_id=?',
                           (int(person_id),)).fetchall()


def decorate():
    data={x:df(t) for x,(t,_) in SPECS.items()}
    p,t,e,c,pe=[data[k] for k in ['Proyectos','Plan de trabajo','Entregables','Control de cambios','Personal']]
    if not t.empty:
        t['Días restantes']=t.due_date.map(datespan)
        t['Días atraso']=t.apply(lambda r:max(0,-r['Días restantes']) if pd.notna(r['Días restantes']) and r.status not in FINISHED_TASK and int(r.progress or 0)<100 else 0,axis=1)
        t['Semáforo']=t.apply(lambda r:'🟢 Terminada' if r.status in FINISHED_TASK or r.progress==100 else '🔴 ATRASADA' if r['Días atraso']>0 else '🟡 Vence pronto' if pd.notna(r['Días restantes']) and r['Días restantes']<=2 else '🔵 Revisión' if r.status=='En revisión' else '⚪ No iniciada' if r.status=='No iniciado' else '🟢 En plazo',axis=1)
    if not e.empty:
        e['Días restantes']=e.due_date.map(datespan)
        e['Alerta']=e.apply(lambda r:'🔴 ATRASADO' if pd.notna(r['Días restantes']) and r['Días restantes']<0 and r.status not in FINISHED_DELIVERY else '🟡 Próximo' if pd.notna(r['Días restantes']) and 0<=r['Días restantes']<=3 and r.status not in FINISHED_DELIVERY else '🔵 En revisión' if 'revisión' in r.status.lower() else '🟢 Cerrado' if r.status in FINISHED_DELIVERY else '',axis=1)
    if not p.empty:
        p['Días restantes']=p.due_date.map(datespan)
        progresses=[]; risks=[]; sem=[]
        for _,r in p.iterrows():
            pt=t[t.project_code==r.code] if not t.empty else pd.DataFrame()
            progress=round(float(pt.progress.fillna(0).mean()),1) if len(pt) else 0
            progresses.append(progress)
            blocked=False
            late=(not pt.empty and bool((pt['Días atraso']>0).any())) or (datespan(r.due_date) is not None and datespan(r.due_date)<0 and r.status not in ('Terminado',))
            risk='Alto' if late or blocked else 'Medio' if (datespan(r.due_date) is not None and datespan(r.due_date)<=7 and progress<100 and r.status!='Terminado') else 'Bajo'
            risks.append(risk)
            sem.append('🔴 Atrasado' if late else '🟠 Bloqueado' if blocked else '🟡 En riesgo' if risk in ('Medio','Alto','Crítico') else '🟢 En plazo')
        p['% Avance']=progresses;p['Riesgo calculado']=risks;p['Situación']=sem
    if not pe.empty:
        work=[];projects=[];load=[]
        for _,r in pe.iterrows():
            assigned=(t[t['id'].isin([x[0] for x in _person_task_ids(int(r['id']))]) &
                (~t.status.isin(FINISHED_TASK))] if not t.empty else pd.DataFrame())
            n=len(assigned)
            work.append(n)
            projects.append(assigned.project_code.nunique() if not assigned.empty else 0)
            load.append('🟡 Con tareas' if n else '🟢 Sin tareas')
        pe['Tareas activas']=work;pe['Proyectos asignados']=projects;pe['Carga de trabajo']=load
    return data

def programmed_deliveries(project_code):
    # Catálogo previo a V01, filtrado siempre por ID Proyecto.
    if not project_code:
        return {}
    with connection() as con:
        rows=con.execute('''SELECT code,name FROM deliverable_catalog
            WHERE project_code=? ORDER BY code''',(str(project_code),)).fetchall()
    return {r['code']:r['name'] for r in rows}


def programmed_delivery_specialty(project_code, delivery_code):
    """La especialidad de una tarea vinculada SIEMPRE viene del catálogo admin.

    Acepta entregables aún sin V01. Un valor antiguo (p. ej. Geotecnia) se
    conserva para registros históricos, aunque ya no aparezca en la lista
    desplegable de especialidades nuevas.
    """
    if not project_code or not delivery_code:
        return ''
    with connection() as con:
        row=con.execute("SELECT specialty FROM deliverable_catalog WHERE project_code=? AND code=?",
                        (str(project_code), str(delivery_code))).fetchone()
    return (row['specialty'] or '') if row else ''


def selectors(kind,current=None,project_code=None):
    if kind=='project': return df('projects').code.dropna().tolist()
    if kind=='delivery': return ['']+list(programmed_deliveries(project_code))
    if kind=='person': return df('people').name.dropna().tolist()
    return OPTIONS.get(kind,[])

def form_input(key,label,kind,required,choice,value,form_key,project_code=None):
    tag=form_key+'_'+key
    if kind in ('select','person','project','delivery'):
        vals=selectors(choice if kind=='select' else kind,project_code=project_code)
        if kind=='select' and key=='specialty' and 'form_people_' in form_key:
            vals=[v for v in vals if v not in ('Seguridad','Relaves')]
        if not required and kind!='delivery': vals=['']+vals
        if value and value not in vals and kind not in ('select','delivery'):vals=[value]+vals
        if not vals: st.warning(f'Primero registre {"personal" if kind=="person" else "proyectos" if kind=="project" else "entregables"}.') ;vals=['']
        if kind=='delivery':
            names=programmed_deliveries(project_code)
            if len(vals)==1 and vals[0]=='':
                st.caption('Este proyecto aún no tiene entregables programados por el Administrador.')
            if value and value not in names:
                st.warning('El entregable vinculado anteriormente no corresponde al proyecto seleccionado o fue retirado. Selecciona otro para guardar.')
            return st.selectbox(label,vals,index=vals.index(value) if value in vals else 0,
                format_func=lambda c:(f'{c} — {names[c]}' if c else '— Sin entregable relacionado —'),
                key=tag+'_'+str(project_code)+'_'+str(value))
        return st.selectbox(label,vals,index=vals.index(value) if value in vals else 0,key=tag+'_'+str(value))
    if kind=='version': return st.text_input(label,value=value or 'V01',key=tag,help='Formato recomendado: V01, V02, V03…')
    if kind=='date':
        val=date.fromisoformat(str(value)[:10]) if value and str(value)!='nan' else None
        return st.date_input(label,value=val,key=tag,format='DD/MM/YYYY')
    if kind=='bool':return st.checkbox(label,value=bool(value),key=tag)
    if kind=='long':return st.text_area(label,value=value or '',key=tag,height=90)
    if kind=='int':return st.number_input(label,min_value=0,max_value=100 if key=='progress' else 100000,value=int(value or (4 if key=='capacity' else 0)),step=1,key=tag)
    if kind=='float':return st.number_input(label,min_value=0.0,value=float(value or 0),step=0.5,key=tag)
    return st.text_input(label,value=str(value or ''),key=tag)

def update_project_sync(values,record_id):
    """Actualiza referencias internas y metadatos Drive. Nunca recrea ni borra archivos."""
    require_admin()
    new_code=str(values.get('code') or '').strip()
    new_name=str(values.get('name') or '').strip()
    if not new_code or not new_name:raise ValueError('Código y nombre de proyecto son obligatorios.')
    with connection() as con:
        old=con.execute('SELECT code,name FROM projects WHERE id=?',(record_id,)).fetchone()
        if not old:raise ValueError('No se encontró el proyecto.')
        old_code,old_name=old['code'],old['name'] or ''
        clash=con.execute('SELECT id FROM projects WHERE code=? AND id<>?',(new_code,record_id)).fetchone()
        if clash:raise ValueError('El código nuevo pertenece a otro proyecto.')
        remote=[dict(x) for x in con.execute('''SELECT f.file_id,f.filename,doc.version,doc.doc_no,d.code AS delivery_code
            FROM document_files f JOIN delivery_documents doc ON f.document_id=doc.id
            JOIN deliverables d ON doc.delivery_id=d.id WHERE d.project_code=?''',(old_code,)).fetchall()]
        known_remote={r['file_id'] for r in remote}
        remote += [dict(x) for x in con.execute('''SELECT f.file_id,f.filename,f.version,'D01' AS doc_no,d.code AS delivery_code
            FROM drive_files f JOIN deliverables d ON f.delivery_id=d.id
            WHERE d.project_code=?''',(old_code,)).fetchall() if x['file_id'] not in known_remote]
    to_rename=[];token=None
    project_old_folder=drive_safe_folder_name(old_code + (' - '+old_name.strip() if old_name.strip() else ''))
    project_new_folder=drive_safe_folder_name(new_code + (' - '+new_name if new_name else ''))
    if old_code!=new_code or project_old_folder!=project_new_folder:
        if remote and not drive_ready():
            raise ValueError('Debes autorizar Google Drive antes de modificar el código o nombre de un proyecto con archivos asociados.')
        if drive_ready():
            token=drive_token()
            old_folder_id=drive_find_folder(token,DRIVE_FOLDER_ID,project_old_folder)
            if remote and not old_folder_id:
                raise ValueError('No se localizó la carpeta antigua del proyecto. No se modificó la base de datos; verifica el nombre en Drive.')
            if old_folder_id and project_old_folder!=project_new_folder:
                target_id=drive_find_folder(token,DRIVE_FOLDER_ID,project_new_folder)
                if target_id and target_id!=old_folder_id:
                    raise ValueError('Ya existe una carpeta con el nuevo nombre. No se modificó el proyecto para evitar mezclar archivos.')
                to_rename.append((old_folder_id,project_old_folder,project_new_folder))
            # Los archivos se nombran solo con ID Entregable + versión, nunca con
            # el nombre del proyecto. Mantener inmutables los IDs de entregables
            # existentes protege sus vínculos y los contadores históricos.
            if old_code!=new_code or old_name!=new_name:
                for f in remote:
                    ext=f['filename'].rsplit('.',1)[-1].lower() if '.' in f['filename'] else 'bin'
                    safe_delivery=re.sub(r'[^A-Za-z0-9_-]','_',f['delivery_code'])
                    target=f'{safe_delivery}_{f["version"]}_{f["doc_no"]}.{ext}'
                    if target!=f['filename']:
                        to_rename.append((f['file_id'],f['filename'],target))
    changes=[]
    try:
        for file_id,old_filename,new_filename in to_rename:
            drive_rename_item(token,file_id,new_filename)
            changes.append((file_id,old_filename))
        with connection() as con:
            vals=dict(values)
            sets=', '.join(f'"{k}"=?' for k in vals)
            con.execute(f'UPDATE projects SET {sets} WHERE id=?',list(vals.values())+[record_id])
            if old_code!=new_code:
                # Las tareas existentes conservan su ID (clave histórica). La serie
                # para tareas nuevas sí sigue el código nuevo del proyecto.
                con.execute('UPDATE task_counters SET project_code=? WHERE project_code=?',(new_code,old_code))
                for table in ('tasks','deliverables','changes','meetings','deliverable_catalog'):
                    con.execute(f'UPDATE {table} SET project_code=? WHERE project_code=?',(new_code,old_code))
            for file_id,old_filename,new_filename in to_rename:
                con.execute('UPDATE drive_files SET filename=? WHERE file_id=?',(new_filename,file_id))
                con.execute('UPDATE document_files SET filename=? WHERE file_id=?',(new_filename,file_id))
            con.execute('INSERT INTO audit(module,record_code,action) VALUES(?,?,?)',
                ('Proyectos',new_code,'Proyecto modificado (código anterior: '+old_code+')'))
    except Exception as exc:
        rollback_errors=[]
        for file_id,old_filename in reversed(changes):
            try:drive_rename_item(token,file_id,old_filename)
            except Exception:rollback_errors.append(file_id)
        if rollback_errors:
            raise ValueError('Falló la actualización y parte de Google Drive no pudo restaurarse. Revisa la carpeta antes de intentarlo de nuevo.') from exc
        raise


def task_code_for_project(project_code,con):
    """Siguiente código disponible sin avanzar contador (para previsualización)."""
    project_code=str(project_code or '').strip()
    if not project_code:
        return ''
    row=con.execute('SELECT highest FROM task_counters WHERE project_code=?',(project_code,)).fetchone()
    highest=int(row['highest']) if row else 0
    while True:
        highest+=1
        code=f'{project_code}_T{highest:03d}'
        if con.execute('SELECT 1 FROM tasks WHERE code=?',(code,)).fetchone() is None:
            return code


def preview_task_code(project_code):
    with connection() as con:
        return task_code_for_project(project_code,con)


def linked_delivery_tasks(project_code,delivery_code,user_id=None):
    """Tareas del mismo entregable; Consulta ve solo sus asignaciones."""
    with connection() as con:
        params=[str(project_code),str(delivery_code)]
        where='t.project_code=? AND t.delivery_code=?'
        if user_id is not None:
            where+=""" AND EXISTS (
                SELECT 1 FROM task_people tp JOIN people p ON p.id=tp.person_id
                JOIN users u ON u.id=p.user_id
                WHERE tp.task_id=t.id AND u.id=? AND u.role='Consulta' AND u.active=1)"""
            params.append(int(user_id))
        rows=[dict(r) for r in con.execute(
            f"""SELECT t.id,t.code,t.activity,t.owner,t.reviewer FROM tasks t
                WHERE {where} ORDER BY t.code""",params)]
        for row in rows:
            assigned=con.execute("""SELECT tp.assignment_role,p.name
                FROM task_people tp JOIN people p ON p.id=tp.person_id
                WHERE tp.task_id=? ORDER BY p.name COLLATE NOCASE""",(row['id'],)).fetchall()
            owners=[r['name'] for r in assigned if r['assignment_role']=='Responsable']
            reviewers=[r['name'] for r in assigned if r['assignment_role']=='Revisor']
            # Las tareas anteriores siguen funcionando aun sin tabla task_people poblada.
            row['owners']=', '.join(owners) if owners else str(row['owner'] or '')
            row['reviewers']=', '.join(reviewers) if reviewers else str(row['reviewer'] or '')
    return rows


def save_record(module,values,record_id=None,task_assignees=None):
    '''Guarda datos; en entregables el catálogo controla IDs y fecha final.

    El servidor aplica las mismas reglas aunque un usuario altere widgets locales.
    '''
    values=dict(values)
    was_update=(record_id is not None)
    is_consulta=(st.session_state.get('role')=='Consulta')
    consulta_delivery=(module=='Entregables' and is_consulta)
    if not consulta_delivery:
        require_admin()
    if consulta_delivery:
        if str(values.get('code') or '') not in consulta_delivery_codes(st.session_state['user_id']):
            raise PermissionError('No tienes acceso a este entregable: falta proyecto autorizado o tarea asignada.')
    if module=='Proyectos' and record_id is not None:
        return update_project_sync(values,record_id)
    table,fields=SPECS[module]
    if module=='Plan de trabajo':
        # Se valida en el servidor: el ID nunca se acepta desde el navegador.
        # Se conserva el código antiguo en ediciones para proteger las referencias.
        if record_id is not None:
            with connection() as con:
                original=con.execute('SELECT code,project_code FROM tasks WHERE id=?',(int(record_id),)).fetchone()
            if not original:
                raise ValueError('La tarea ya no existe. Actualiza la pantalla.')
            if str(values.get('project_code') or '')!=str(original['project_code'] or ''):
                raise ValueError('No es posible cambiar el proyecto de una tarea existente. Crea otra tarea en el proyecto correspondiente.')
            values['code']=original['code']
        else:
            values['code']='POR_ASIGNAR'  # Marcador interno; se sustituye en transacción.
        if task_assignees is not None:
            owners,reviewers=task_assignees
            owners=set(map(int,owners));reviewers=set(map(int,reviewers))
            if not owners:raise ValueError('Selecciona por lo menos un responsable.')
            with connection() as con:
                all_ids=owners | reviewers
                marks=','.join('?' for _ in all_ids)
                valid={int(r[0]) for r in con.execute(f'SELECT id FROM people WHERE id IN ({marks})',list(all_ids))}
                if all_ids!=valid:raise ValueError('Un responsable o revisor ya no existe en Personal.')
                all_people={int(r['id']):r['name'] for r in con.execute('SELECT id,name FROM people')}
            values['owner']=', '.join(sorted(all_people[i] for i in owners))
            values['reviewer']=', '.join(sorted(all_people[i] for i in reviewers))
        project=str(values.get('project_code') or '').strip()
        selected=str(values.get('delivery_code') or '').strip()
        with connection() as con:
            if not con.execute('SELECT 1 FROM projects WHERE code=?',(project,)).fetchone():
                raise ValueError('Selecciona un proyecto registrado para esta tarea.')
            if selected:
                catalog=con.execute("SELECT specialty FROM deliverable_catalog WHERE project_code=? AND code=?",
                                    (project,selected)).fetchone()
                if catalog is None:
                    raise ValueError('El entregable relacionado no está programado para este proyecto. '
                                     'Selecciona uno del catálogo del proyecto correspondiente.')
                # La BD prevalece frente al campo del navegador.
                values['specialty']=catalog['specialty'] or ''
            elif values.get('specialty') not in ('',None,*SPECIALTIES):
                raise ValueError('Selecciona una especialidad vigente del listado.')
    if module=='Entregables':
        code=str(values.get('code') or '').strip()
        project=str(values.get('project_code') or '').strip()
        if not code or not project:raise ValueError('Selecciona un proyecto y un entregable del catálogo.')
        with connection() as con:
            catalog=con.execute('SELECT * FROM deliverable_catalog WHERE code=?',(code,)).fetchone()
            # Carga ficticia/operación histórica por Admin: alta en catálogo.
            if catalog is None and can_edit() and record_id is None:
                con.execute('''INSERT INTO deliverable_catalog(project_code,code,name,specialty,final_due_date)
                    VALUES(?,?,?,?,?)''',(project,code,str(values.get('name') or '').strip(),
                    values.get('specialty') or '',values.get('due_date') or date.today().isoformat()))
                catalog=con.execute('SELECT * FROM deliverable_catalog WHERE code=?',(code,)).fetchone()
            if not catalog or catalog['project_code']!=project:
                raise ValueError('El ID de entregable no está autorizado en el catálogo de este proyecto.')
            registered=con.execute('SELECT * FROM deliverables WHERE code=?',(code,)).fetchone()
        if record_id is None and registered:
            raise ValueError('Este entregable ya tiene una versión. Actualiza la página antes de registrar otra.')
        if consulta_delivery and record_id is not None:
            if not registered:raise ValueError('El entregable aún no existe.')
            expected=next_version(code,registered['version'])
            if str(values.get('version') or '')!=expected:
                raise ValueError(f'Consulta solo puede registrar una nueva versión ({expected}); no editar versiones existentes.')
        if record_id is not None and (not registered or registered['id']!=record_id or registered['project_code']!=project):
            raise ValueError('No está permitido reasignar un entregable a otro proyecto.')
        values['code']=code
        values['project_code']=project
        values['name']=catalog['name']
        values['specialty']=catalog['specialty'] or ''
        values['due_date']=catalog['final_due_date']
        values.pop('drawing_code',None)
        values.pop('file_path',None)
        if record_id is None:
            values['version']=next_version(code)
            values['actual_date']=today_peru().isoformat()
            if is_consulta:
                values['status']='Pendiente'
                for f in ('review_date','correction_date','approval_date'):
                    values[f]=None
        else:
            original_version=registered['version'] or ''
            if values.get('version')!=original_version:
                expected=next_version(code,original_version)
                if values['version']!=expected:
                    raise ValueError('La nueva versión debe ser '+expected)
                values['actual_date']=today_peru().isoformat()
            else:
                values['actual_date']=registered['actual_date']
            if consulta_delivery:
                values['status']='Pendiente'
                for f in ('review_date','correction_date','approval_date'):
                    values[f]=None
        # task_id NO es un campo libre. Se valida nuevamente contra la BD.
        chosen_id=values.get('task_id')
        try:chosen_id=int(chosen_id) if chosen_id not in ('',None) else None
        except (TypeError,ValueError):raise ValueError('Selecciona un ID Tarea válido.')
        if consulta_delivery and chosen_id is None:
            raise ValueError('Selecciona una tarea vinculada a este entregable.')
        if chosen_id is not None:
            permitted=linked_delivery_tasks(project,code,
                st.session_state['user_id'] if consulta_delivery else None)
            selected_task=next((t for t in permitted if int(t['id'])==chosen_id),None)
            if selected_task is None:
                raise PermissionError('La tarea no pertenece al entregable o no estás asignado a ella.')
            values['task_id']=chosen_id
            values['owner']=selected_task['owners']
            values['reviewer']=selected_task['reviewers']
        else:
            selected_task=None
            values['task_id']=None
        # Cada guardado/versión reinicia la antigüedad; la edición rápida solo el estado.
        values['last_modified_at']=peru_timestamp()
        registering_version=(record_id is None or
                              str(values['version'])!=str(registered['version'] if registered else ''))
        if values.get('status') not in DELIVERY_STATES:
            raise ValueError('Estado de entregable no válido.')
    mandatory=[label for k,label,_,required,_ in fields if required and (values.get(k) in ('',None))]
    if mandatory:raise ValueError('Campos obligatorios: '+', '.join(mandatory))
    if module=='Plan de trabajo' and values['start_date']>values['due_date']:raise ValueError('Fecha término anterior al inicio.')
    if module=='Entregables' and not str(values['version']).upper().startswith('V'):
        raise ValueError('Use versiones V01, V02, etc.')
    if module=='Entregables' and values['status'] in ('Enviado al cliente','Aprobado','Entregado') and not values.get('review_date'):
        raise ValueError('Debe registrar fecha de revisión antes de enviar/aprobar/entregar.')
    if module=='Personal':
        values['code']=str(values.get('code') or '').strip()
        if not values['code']:raise ValueError('El nombre es obligatorio.')
        values['name']=values['code']
    with connection() as con:
        if consulta_delivery:
            con.execute('BEGIN IMMEDIATE')
            live=con.execute('SELECT id,version,project_code FROM deliverables WHERE code=?',
                             (values['code'],)).fetchone()
            if record_id is None and live:
                raise ValueError('Este entregable acaba de registrarse. Recarga la pantalla.')
            if record_id is not None:
                if not live or live['id']!=int(record_id):
                    raise ValueError('El entregable fue modificado. Actualiza la página.')
                expected=next_version(values['code'],live['version'])
                if values['version']!=expected:
                    raise ValueError('Ya se registró otra versión. Actualiza la página para obtener la siguiente.')
        if module=='Personal':
            other=con.execute('SELECT id FROM people WHERE lower(name)=lower(?) AND id!=?',(values['name'],record_id or -1)).fetchone()
            if other:raise ValueError('Ya existe una persona con ese nombre.')
            if record_id:
                old=con.execute('SELECT name FROM people WHERE id=?',(record_id,)).fetchone()
                if old and old['name']!=values['name']:
                    for linked_table,column in [('projects','manager'),('deliverables','owner'),('deliverables','reviewer'),('changes','owner'),('meetings','owner')]:
                        con.execute(f'UPDATE {linked_table} SET {column}=? WHERE {column}=?',(values['name'],old['name']))
                linked=con.execute('SELECT user_id FROM people WHERE id=?',(record_id,)).fetchone()
                if linked and linked['user_id']:
                    con.execute('UPDATE users SET full_name=? WHERE id=?',
                                (values['name'],linked['user_id']))
        if record_id:
            sets=', '.join(f'{k}=?' for k in values)
            con.execute(f'UPDATE {table} SET {sets} WHERE id=?',list(values.values())+[record_id])
            if module=='Personal':
                _refresh_task_display(con,[r[0] for r in con.execute(
                    'SELECT DISTINCT task_id FROM task_people WHERE person_id=?',(record_id,))])
        else:
            if module=='Plan de trabajo':
                # Bloqueo de escritura ANTES de leer el contador: dos administradores
                # no pueden obtener el mismo código si guardan al mismo tiempo.
                con.execute('BEGIN IMMEDIATE')
                project_code=str(values['project_code']).strip()
                generated=task_code_for_project(project_code,con)
                values['code']=generated
                serial=int(generated.rsplit('_T',1)[1])
                con.execute('''INSERT INTO task_counters(project_code,highest) VALUES(?,?)
                    ON CONFLICT(project_code) DO UPDATE SET highest=MAX(highest,excluded.highest)''',
                    (project_code,serial))
            keys=','.join(values.keys());placeholders=','.join('?' for _ in values)
            inserted=con.execute(f'INSERT INTO {table} ({keys}) VALUES ({placeholders})',list(values.values()))
            if module=='Plan de trabajo':record_id=inserted.lastrowid
        if module=='Plan de trabajo' and task_assignees is not None:
            owners,reviewers=task_assignees
            con.execute('DELETE FROM task_people WHERE task_id=?',(record_id,))
            con.executemany('INSERT INTO task_people(task_id,person_id,assignment_role) VALUES(?,?,?)',
                [(record_id,int(pid),'Responsable') for pid in set(owners)] +
                [(record_id,int(pid),'Revisor') for pid in set(reviewers)])
            con.execute('INSERT OR IGNORE INTO task_people_migration(task_id) VALUES(?)',(record_id,))
        if module=='Entregables':
            if registering_version or (was_update and can_edit() and selected_task is not None):
                stored=con.execute('SELECT id FROM deliverables WHERE code=?',(values['code'],)).fetchone()
                if stored:
                    con.execute('''INSERT INTO delivery_version_tasks
                        (delivery_id,version,task_id,task_code,activity,owners,reviewers)
                        VALUES(?,?,?,?,?,?,?) ON CONFLICT(delivery_id,version) DO UPDATE SET
                        task_id=excluded.task_id,task_code=excluded.task_code,
                        activity=excluded.activity,owners=excluded.owners,reviewers=excluded.reviewers''',
                        (int(stored['id']),values['version'],values['task_id'],
                         selected_task['code'] if selected_task else '',
                         selected_task['activity'] if selected_task else '',
                         values.get('owner') or '',values.get('reviewer') or ''))
            con.execute('INSERT OR IGNORE INTO versions(delivery_code,version,registered_at,notes) VALUES (?,?,?,?)',
                (values['code'],values['version'],date.today().isoformat(),'Versión registrada desde formulario'))
            m=re.fullmatch(r'V(\d+)',str(values['version']).upper())
            if m:
                con.execute('''INSERT INTO version_counters(delivery_code,highest) VALUES(?,?)
                    ON CONFLICT(delivery_code) DO UPDATE SET highest=MAX(highest,excluded.highest)''',
                    (values['code'],int(m.group(1))))
        con.execute('INSERT INTO audit(module,record_code,action) VALUES (?,?,?)',
            (module,values['code'],'Actualización' if was_update else 'Alta'))


def backup_database():
    """Respaldo SQLite coherente; no copiar el archivo mientras está escribiéndose."""
    folder=os.path.join(ROOT,'respaldos')
    os.makedirs(folder,exist_ok=True)
    stamp=datetime.now().strftime('%Y%m%d_%H%M%S_%f')
    filename=os.path.join(folder,f'ALTIVIA_automatico_{stamp}.db')
    with sqlite3.connect(DB,timeout=30) as origin, sqlite3.connect(filename) as destination:
        origin.backup(destination)
    return filename


def delete_selected(module,ids):
    require_admin()
    table,_=SPECS[module]
    ids=list(dict.fromkeys(int(i) for i in ids))
    if not ids:return 0
    with connection() as con:
        marks=','.join('?' for _ in ids)
        rows=con.execute(f'SELECT code FROM {table} WHERE id IN ({marks})',ids).fetchall()
        codes=[r['code'] for r in rows]
        if module=='Personal':
            protected=[dict(x) for x in con.execute(f'''SELECT p.name,p.user_id,
                EXISTS(SELECT 1 FROM task_people tp WHERE tp.person_id=p.id) AS assigned
                FROM people p WHERE p.id IN ({marks})''',ids)]
            if any(r['user_id'] or r['assigned'] for r in protected):
                raise ValueError('No se puede eliminar Personal vinculado a un usuario o a tareas. '
                                 'Primero reasigna las tareas o elimina la cuenta y desvincula las tareas.')
        if module=='Proyectos' and codes:
            pm=','.join('?' for _ in codes)
            con.execute(f'DELETE FROM task_counters WHERE project_code IN ({pm})',codes)
            dc=[r[0] for r in con.execute(f'SELECT code FROM deliverables WHERE project_code IN ({pm})',codes)]
            if dc:
                dm=','.join('?' for _ in dc)
                con.execute(f'DELETE FROM versions WHERE delivery_code IN ({dm})',dc)
                con.execute(f'DELETE FROM checklist WHERE delivery_code IN ({dm})',dc)
            con.execute(f'DELETE FROM delivery_files WHERE delivery_id IN (SELECT id FROM deliverables WHERE project_code IN ({pm}))',codes)
            # La eliminación de un PROYECTO elimina también sus catálogos y sus
            # contadores (no solo los que tenían una entrega registrada).
            all_catalog_codes=[r[0] for r in con.execute(f'SELECT code FROM deliverable_catalog WHERE project_code IN ({pm})',codes)]
            all_counter_codes=list(dict.fromkeys(all_catalog_codes+dc))
            if all_counter_codes:
                cm=','.join('?' for _ in all_counter_codes)
                con.execute(f'DELETE FROM version_counters WHERE delivery_code IN ({cm})',all_counter_codes)
            con.execute(f'DELETE FROM deliverable_catalog WHERE project_code IN ({pm})',codes)
            for linked in ('tasks','deliverables','changes','meetings'):
                con.execute(f'DELETE FROM {linked} WHERE project_code IN ({pm})',codes)
        if module=='Plan de trabajo':
            con.execute(f'DELETE FROM task_people_migration WHERE task_id IN ({marks})',ids)
        if module=='Entregables' and codes:
            con.execute(f'DELETE FROM delivery_files WHERE delivery_id IN ({marks})',ids)
            dm=','.join('?' for _ in codes)
            con.execute(f'DELETE FROM versions WHERE delivery_code IN ({dm})',codes)
            con.execute(f'DELETE FROM checklist WHERE delivery_code IN ({dm})',codes)
            # Las tareas no se borran: solo se elimina la relación al entregable.
            con.execute(f"UPDATE tasks SET delivery_code='' WHERE delivery_code IN ({dm})",codes)
        con.execute(f'DELETE FROM {table} WHERE id IN ({marks})',ids)
        for code in codes:
            con.execute('INSERT INTO audit(module,record_code,action) VALUES (?,?,?)',
                        (module,code,'Eliminación múltiple'))
    return len(codes)


def bulk_delete_ui(module,data):
    require_admin()
    table,_=SPECS[module]
    view=data[module]
    if view.empty:return
    with st.expander('🗑️ Eliminar varios registros a la vez'):
        st.warning('Esta acción es irreversible. Los proyectos eliminados también eliminarán sus tareas, entregables, revisiones, cambios y reuniones asociados. Se creará un respaldo automático.')
        options=view['id'].astype(int).tolist()
        names={int(r['id']):f"{r.get('code','')} — {r.get('name',r.get('activity',r.get('topic','')))}" for _,r in view.iterrows()}
        chosen=st.multiselect('Seleccione uno o varios registros',options,format_func=lambda rid:names.get(rid,str(rid)),key='multidel_'+table)
        confirm=st.checkbox(f'Confirmo eliminar {len(chosen)} registro(s) seleccionado(s)',key='multiconfirm_'+table)
        typed=st.text_input('Escriba ELIMINAR para confirmar',key='multitype_'+table)
        if st.button('Eliminar registros seleccionados',type='primary',disabled=not chosen or not confirm or typed!='ELIMINAR',key='multibutton_'+table):
            try:
                path=backup_database()
                n=delete_selected(module,chosen)
                st.success(f'Se eliminaron {n} registros. Respaldo automático: {os.path.basename(path)}')
                st.rerun()
            except Exception as exc:st.error(f'No se pudo eliminar: {exc}')


def reset_database_ui():
    require_admin()
    st.divider()
    st.subheader('⚠️ Restablecer información de ALTIVIA')
    st.warning('Esta herramienta BORRA datos definitivamente del sistema activo. Antes se genera una copia .db en la carpeta respaldos. Resguarde también una copia fuera de esta computadora.')
    mode=st.radio('Seleccione el alcance del restablecimiento',
        ['Limpiar datos operativos (conservar todos los usuarios)',
         'Restablecimiento general (borrar datos y otros usuarios; conservar mi cuenta administradora)'],
         key='reset_mode')
    st.caption('Se borrarán proyectos, sus autorizaciones de visualización, tareas, entregables, versiones, checklists, cambios, reuniones, personal y bitácora. La segunda opción también borra todas las cuentas excepto el administrador que ejecuta el procedimiento.')
    with st.form('reset_form'):
        typed=st.text_input('Escriba RESTABLECER ALTIVIA')
        password=st.text_input('Su contraseña de administrador',type='password')
        accept=st.checkbox('Entiendo que todos los datos seleccionados serán eliminados')
        go=st.form_submit_button('Restablecer base de datos',type='primary')
    if go:
        if typed!='RESTABLECER ALTIVIA' or not accept:
            st.error('Complete la confirmación de forma exacta.');return
        with connection() as con:
            user=con.execute('SELECT password_hash FROM users WHERE id=? AND role=? AND active=1',
                             (st.session_state['user_id'],'Administrador')).fetchone()
        if not user or not verify_password(password,user['password_hash']):
            st.error('Contraseña de administrador incorrecta.');return
        try:
            path=backup_database()
            with connection() as con:
                for table in ('task_counters','document_files','document_checklist','delivery_documents','document_serials','legacy_document_imports','drive_files','delivery_files','checklist','versions','version_counters','meetings','changes','task_people','task_people_migration','tasks','deliverables','deliverable_catalog','projects','people','audit'):
                    con.execute(f'DELETE FROM {table}')
                if mode.startswith('Restablecimiento general'):
                    con.execute('DELETE FROM users WHERE id<>?',(st.session_state['user_id'],))
                con.execute('INSERT INTO audit(module,record_code,action) VALUES (?,?,?)',
                            ('Sistema','ALTIVIA','Restablecimiento general' if mode.startswith('Restablecimiento general') else 'Limpieza de datos'))
            st.success(f'Información restablecida. Respaldo guardado: {path}')
            st.rerun()
        except Exception as exc:st.error(f'El restablecimiento falló: {exc}')


def open_task_editor(task_id=None):
    """Solo el administrador puede abrir la edición (también desde callbacks)."""
    require_admin()
    st.session_state['task_editor_id']=int(task_id) if task_id is not None else None
    st.session_state['task_editor_open']=True


def close_task_editor():
    st.session_state['task_editor_id']=None
    st.session_state['task_editor_open']=False


def task_cards(frame):
    """Fichas específicas de Plan de trabajo; filtradas por permisos antes de entrar."""
    st.markdown('### 📂 Registros')
    if frame.empty:
        st.info('Todavía no hay actividades disponibles.');return
    query=st.text_input('🔎 Buscar tarea, actividad o integrante',key='cards_search_tasks')
    a,b=st.columns(2)
    with a:
        states=sorted(str(x) for x in frame['status'].dropna().unique())
        state=st.selectbox('Estado',['Todos']+states,key='cards_state_tasks')
    with b:
        projects=sorted(str(x) for x in frame['project_code'].dropna().unique())
        project=st.selectbox('Proyecto',['Todos']+projects,key='cards_proj_tasks') if projects else 'Todos'
    if query:
        frame=frame[frame.astype(str).apply(lambda ser:ser.str.contains(query,case=False,regex=False)).any(axis=1)]
    if state!='Todos':frame=frame[frame['status'].astype(str)==state]
    if project!='Todos':frame=frame[frame['project_code'].astype(str)==project]
    st.caption(f'{len(frame)} tarea(s) encontradas')
    # Descripción del proyecto / cliente se recupera de Proyectos, no del ID.
    with connection() as con:
        project_names={str(p['code']):(str(p['name'] or ''),str(p['client'] or ''))
            for p in con.execute('SELECT code,name,client FROM projects')}
    page_size=10
    max_page=max(1,(len(frame)+page_size-1)//page_size)
    page=st.number_input('Página',min_value=1,max_value=max_page,value=1,step=1,key='cards_page_tasks')
    for _,r in frame.iloc[(page-1)*page_size:page*page_size].iterrows():
        code=escape(str(r.get('code') or ''))
        activity=escape(str(r.get('activity') or 'Actividad sin nombre'))
        status=escape(str(r.get('status') or 'Sin estado'))
        owners=escape(str(r.get('owner') or 'Sin asignar'))
        reviewers=escape(str(r.get('reviewer') or 'Sin asignar'))
        start=escape(str(r.get('start_date') or '—'))
        due=escape(str(r.get('due_date') or '—'))
        with st.container(border=True):
            # 1. Actividad destacada: 1.2rem = 20% más que el texto base.
            st.markdown(f'<div style="font-size:1.2rem;font-weight:750;line-height:1.35;overflow-wrap:anywhere">{activity}</div>',
                        unsafe_allow_html=True)
            # 2. Código de tarea: segunda línea, sin especialidad.
            st.caption(f'ID Tarea: {code}')
            project_name,client_name=project_names.get(str(r.get('project_code') or ''),('',''))
            project_info=' · '.join(s for s in (project_name, f'Cliente: {client_name}' if client_name else '') if s)
            if project_info:
                st.caption('🏗️ '+project_info)
            st.markdown(f'**Estado:** {status}')
            cols=st.columns(2)
            with cols[0]:
                st.markdown('**Responsables**')
                st.write(owners)
            with cols[1]:
                st.markdown('**Revisores**')
                st.write(reviewers)
            dates=st.columns(2)
            with dates[0]:st.caption(f'Fecha de inicio: {start}')
            with dates[1]:st.caption(f'Fecha de término: {due}')
            with st.expander('Ver detalles'):
                details=[('ID Proyecto','project_code'),('Entregable relacionado','delivery_code'),
                         ('Especialidad','specialty'),('Avance (%)','progress'),
                         ('Prioridad','priority'),('Fecha actualización','updated_at'),
                         ('Observaciones','notes')]
                for label,key in details:
                    value=r.get(key)
                    if value is not None and pd.notna(value) and str(value).strip():
                        st.markdown(f'**{label}:** {escape(str(value))}')
            if can_edit():
                st.button('✏️ Editar',key=f'edit_task_card_{int(r["id"])}',
                          on_click=open_task_editor,args=(int(r['id']),),
                          use_container_width=True)



def update_delivery_status(delivery_id, new_status):
    """Cambio rápido de estado por el Administrador, sin modificar versiones ni archivos."""
    require_admin()
    if new_status not in DELIVERY_STATES:
        raise ValueError('El estado seleccionado no es válido.')
    with connection() as con:
        con.execute('BEGIN IMMEDIATE')
        record=con.execute("SELECT id,code,status,review_date,approval_date,correction_date FROM deliverables WHERE id=?",(int(delivery_id),)).fetchone()
        if not record:raise ValueError('El entregable ya no existe.')
        if record['status']==new_status:return False
        fields={'status':new_status,'last_modified_at':peru_timestamp()}
        stamp=today_peru().isoformat()
        # Respeta la validación del módulo: estos estados requieren revisión registrada.
        if new_status in ('Enviado al cliente','Aprobado','Entregado') and not record['review_date']:
            fields['review_date']=stamp
        if new_status in ('Aprobado','Entregado') and not record['approval_date']:
            fields['approval_date']=stamp
        if new_status=='Corregido' and not record['correction_date']:
            fields['correction_date']=stamp
        assignments=', '.join(f'{k}=?' for k in fields)
        con.execute(f'UPDATE deliverables SET {assignments} WHERE id=?',
                    (*fields.values(),int(delivery_id)))
        con.execute('INSERT INTO audit(module,record_code,action) VALUES (?,?,?)',
                    ('Entregables',record['code'],f'Cambio rápido de estado: {record["status"]} → {new_status}'))
    return True


def card_browser(module, frame, allow_version_edit=True):
    """Visualizador responsive sin exponer campos técnicos."""
    st.markdown('### 📂 Registros')
    if frame.empty:
        st.info('Todavía no hay registros en este módulo.')
        return
    table,fields=SPECS[module]
    labels={k:label for k,label,_,_,_ in fields}
    q=st.text_input('🔎 Buscar por código, nombre o responsable',key='cards_search_'+table)
    a,b=st.columns(2)
    with a:
        states=sorted(str(x) for x in frame['status'].dropna().unique()) if 'status' in frame else []
        status=('Todos' if module=='Entregables' and not can_edit() else
                st.selectbox('Estado',['Todos']+states,key='cards_state_'+table))
    with b:
        projects=sorted(str(x) for x in frame['project_code'].dropna().unique()) if 'project_code' in frame else []
        project=st.selectbox('Proyecto',['Todos']+projects,key='cards_proj_'+table) if projects else 'Todos'
    if q:frame=frame[frame.astype(str).apply(lambda series:series.str.contains(q,case=False,regex=False)).any(axis=1)]
    if status!='Todos':frame=frame[frame['status'].astype(str)==status]
    if project!='Todos':frame=frame[frame['project_code'].astype(str)==project]
    st.caption(f'{len(frame)} registro(s) encontrados')
    page_size=10
    max_page=max(1,(len(frame)+page_size-1)//page_size)
    page=st.number_input('Página',min_value=1,max_value=max_page,value=1,step=1,key='cards_page_'+table)
    frame=frame.iloc[(page-1)*page_size:page*page_size]
    for _,r in frame.iterrows():
        title=str(r.get('name') or r.get('activity') or r.get('description') or r.get('code') or 'Registro')
        status_text=str(r.get('status') or 'Sin estado')
        code=str(r.get('code') or '')
        subtitle=' · '.join(str(r.get(k)) for k in ('project_code','specialty','version') if k in r and pd.notna(r.get(k)) and str(r.get(k)).strip())
        if module=='Proyectos':
            card_key=f'gp_project_card_{int(r["id"])}'
        elif module=='Entregables':
            signal,elapsed=delivery_signal(status_text,r.get('last_modified_at'))
            card_key=f'gp_delivery_{signal}_card_{int(r["id"])}'
        else:
            card_key=f'gp_generic_card_{table}_{int(r["id"])}'
        with st.container(border=True,key=card_key):
            st.markdown(f'**{escape(title)}**')
            if module=='Entregables':
                label={'yellow':'En seguimiento','red':'Sin cambios por más de 3 días',
                       'green':'Estado actualizado por Administrador'}[signal]
                st.caption('● '+label)
            st.caption(f'{escape(code)}  ·  {escape(subtitle)}')
            if module!='Entregables' or can_edit():
                st.markdown(f'**Estado:** {escape(status_text)}')
            if module=='Entregables' and can_edit():
                # Edición rápida directamente en la ficha, sin abrir el formulario completo.
                delivery_id=int(r['id'])
                with st.form(f'quick_delivery_state_{delivery_id}'):
                    left,right=st.columns([3,1],gap='small',vertical_alignment='bottom')
                    with left:
                        selected_state=st.selectbox('Cambiar estado',DELIVERY_STATES,
                            index=DELIVERY_STATES.index(status_text) if status_text in DELIVERY_STATES else 0,
                            key=f'quick_delivery_state_value_{delivery_id}')
                    with right:
                        saved=st.form_submit_button('💾 Guardar estado',use_container_width=True)
                if saved:
                    try:
                        changed=update_delivery_status(delivery_id,selected_state)
                        if changed:
                            st.toast('Estado actualizado',icon='✅')
                            st.rerun()
                        else:st.info('El entregable ya tiene ese estado.')
                    except (ValueError,PermissionError,sqlite3.Error) as exc:
                        st.error(f'No se pudo actualizar el estado: {exc}')
            cols=st.columns(2)
            details=[k for k in ('owner','reviewer','manager','due_date','actual_date','priority','progress','% Avance','Riesgo calculado','client','role','email') if k in r and pd.notna(r.get(k)) and str(r.get(k)).strip()]
            for idx,k in enumerate(details):
                with cols[idx%2]:st.caption(f'{labels.get(k,k)}: {r[k]}')
            with st.expander('Ver detalles'+(' y versiones' if module=='Entregables' else '')):
                for k,label,_,_,_ in fields:
                    v=r.get(k)
                    if k not in details and k not in ('code','name','activity','description','status','file_path','drawing_code') and v is not None and pd.notna(v) and str(v).strip():
                        st.markdown(f'**{label}:** {escape(str(v))}')
                if module=='Entregables':
                    versions_for_delivery(code,allow_edit=allow_version_edit and can_edit())
            if module=='Proyectos' and can_edit():
                rid=int(r['id'])
                with st.expander('✏️ Editar información del proyecto',expanded=False):
                    st.caption('Si cambias el código, ALTIVIA también actualizará las tareas, entregables y cambios vinculados. '
                               'Los archivos vinculados se renombrarán en Google Drive si la conexión está autorizada.')
                    with st.form(f'edit_project_form_{rid}'):
                        values={}
                        for key,label,kind,required,opt in SPECS['Proyectos'][1]:
                            raw=r.get(key)
                            if raw is not None and pd.isna(raw):raw=None
                            values[key]=form_input(key,label+(' *' if required else ''),kind,required,opt,
                                raw,f'project_edit_{rid}')
                        if st.form_submit_button('💾 Guardar información del proyecto',type='primary'):
                            values={k:(v.isoformat() if isinstance(v,date) else v) for k,v in values.items()}
                            try:
                                save_record('Proyectos',values,rid)
                                st.success('Proyecto y referencias actualizados.');st.rerun()
                            except Exception as exc:st.error(f'No se completó el cambio: {exc}')


def versions_for_delivery(code,allow_edit=False):
    """Historial visual y descarga; Consulta ve versiones de proyectos autorizados."""
    if st.session_state.get('role')=='Consulta':
        with connection() as con:
            authorized=con.execute('''SELECT 1 FROM deliverables d
                JOIN projects p ON p.code=d.project_code
                JOIN consulta_project_access pa ON pa.project_id=p.id
                WHERE d.code=? AND pa.user_id=?''',(code,st.session_state['user_id'])).fetchone()
        if not authorized:
            st.error('No tienes acceso a las versiones de este proyecto.');return
    with connection() as con:
        delivery=con.execute('SELECT id,version FROM deliverables WHERE code=?',(code,)).fetchone()
        versions=[dict(v) for v in con.execute(
            'SELECT id,version,registered_at,notes FROM versions WHERE delivery_code=? ORDER BY id DESC',(code,))]
    if not delivery:st.warning('Entregable no disponible.');return
    did=int(delivery['id'])
    is_admin=bool(allow_edit and can_edit())
    latest_version=delivery['version'] or (versions[0]['version'] if versions else '')
    st.markdown('**📚 Versiones y documentos**')
    for ver in versions:
        version=ver['version'];vid=ver['id']
        docs=docs_for_version(did,version)
        with connection() as con:
            version_task=con.execute('''SELECT task_code,activity,owners,reviewers
                FROM delivery_version_tasks WHERE delivery_id=? AND version=?''',
                (did,version)).fetchone()
        with st.container(border=True):
            st.markdown(f'**{escape(version)}** · {escape(str(ver["registered_at"]))}' +
                (' · **Vigente**' if version==delivery['version'] else ''))
            st.caption((ver['notes'] or 'Sin descripción') if is_admin else f'{len(docs)} documento(s)')
            if version_task and version_task['task_code']:
                with st.expander('Tarea vinculada a esta versión'):
                    st.markdown('**ID Tarea:** '+escape(str(version_task['task_code'])))
                    st.markdown('**Actividad:** '+escape(str(version_task['activity'] or '—')))
                    st.markdown('**Responsables:** '+escape(str(version_task['owners'] or '—')))
                    st.markdown('**Revisores:** '+escape(str(version_task['reviewers'] or '—')))
            for d in docs:
                with st.container(border=True):
                    st.markdown(f'**📄 {escape(d["doc_no"])} — {escape(d["title"])}**')
                    document_preview_ui(d,allow_download=True,
                        allow_editable=(is_admin or version==latest_version),
                        prefix=f'version_{vid}')
                    if is_admin:
                        with st.expander(f'👥 Asignaciones del documento {d["doc_no"]} (Checklist)'):
                            customized,(owners,reviewers)=document_assignments(d['id'])
                            st.caption('Por defecto se heredan responsables y revisores de las tareas del entregable. '
                                       'Al guardar una lista personalizada, el Checklist solo será visible para estas personas, '
                                       'si además tienen acceso al proyecto.')
                            roster=people_for_assignment()
                            labels={int(p['id']):p['name'] for p in roster}
                            with st.form(f'doc_assign_{d["id"]}'):
                                selected_owners=st.multiselect('Responsables del documento',list(labels),
                                    default=[i for i in labels if i in owners],
                                    format_func=lambda i:labels[i],key=f'doc_assign_o_{d["id"]}')
                                selected_reviewers=st.multiselect('Revisores del documento',list(labels),
                                    default=[i for i in labels if i in reviewers],
                                    format_func=lambda i:labels[i],key=f'doc_assign_r_{d["id"]}')
                                if st.form_submit_button('💾 Guardar asignaciones'):
                                    try:
                                        save_document_assignments(d['id'],selected_owners,selected_reviewers)
                                        st.success('Asignaciones de documento actualizadas');st.rerun()
                                    except Exception as exc:st.error(str(exc))
                            if customized and st.button('↩ Volver a heredar asignaciones de las tareas',
                                                        key=f'doc_assign_reset_{d["id"]}'):
                                with connection() as con:
                                    con.execute('DELETE FROM document_assignment_mode WHERE document_id=?',(d['id'],))
                                    con.execute('DELETE FROM document_people WHERE document_id=?',(d['id'],))
                                st.rerun()
                        with st.expander(f'✏️ Administrar {d["doc_no"]}',expanded=False):
                            with st.form(f'doc_edit_{d["id"]}'):
                                title=st.text_input('Nombre descriptivo',value=d['title'])
                                pdf=st.file_uploader('Reemplazar PDF',type=['pdf'],key=f'doc_pdf_{d["id"]}')
                                editable=st.file_uploader('Reemplazar editable',type=['dwg','doc','docx','xlsx','xls','ifc','rvt'],key=f'doc_edit_{d["id"]}')
                                if st.form_submit_button('💾 Guardar documento'):
                                    try:
                                        if not title.strip():raise ValueError('Indica el nombre del documento.')
                                        if (pdf or editable) and not drive_ready():raise ValueError('Conecta Google Drive primero.')
                                        if pdf:upload_document_file(pdf,d['id'],'pdf')
                                        if editable:upload_document_file(editable,d['id'],'editable')
                                        with connection() as con:
                                            con.execute('UPDATE delivery_documents SET title=? WHERE id=?',(title.strip(),d['id']))
                                            con.execute('INSERT INTO audit(module,record_code,action) VALUES(?,?,?)',
                                                ('Documentos',code,f'{version}/{d["doc_no"]} editado'))
                                        st.success('Documento actualizado');st.rerun()
                                    except Exception as exc:st.error(str(exc))
                            if st.button('🗑 Eliminar documento '+d['doc_no'],key=f'doc_rm_{d["id"]}'):
                                st.session_state[f'doc_confirm_{d["id"]}']=True
                            if st.session_state.get(f'doc_confirm_{d["id"]}'):
                                st.warning('Se retirará el documento y su checklist de ALTIVIA. Los archivos físicos permanecerán en Google Drive como respaldo.')
                                if st.button('Confirmar eliminación de '+d['doc_no'],key=f'doc_delete_yes_{d["id"]}'):
                                    with connection() as con:
                                        con.execute('DELETE FROM delivery_documents WHERE id=?',(d['id'],))
                                        con.execute('INSERT INTO audit(module,record_code,action) VALUES(?,?,?)',
                                            ('Documentos',code,'Retirado '+version+'/'+d['doc_no']+' (archivos en Drive conservados)'))
                                    st.session_state.pop(f'doc_confirm_{d["id"]}',None)
                                    st.rerun()
            if not docs:st.caption('Esta versión aún no tiene documentos.')
            if not is_admin:continue
            with st.expander('➕ Agregar documentos a '+version):
                how_many=st.number_input('Cantidad de documentos nuevos',min_value=1,max_value=10,value=1,
                    key=f'add_doc_count_{vid}')
                with st.form(f'add_doc_form_{vid}'):
                    entries=[]
                    for i in range(how_many):
                        st.markdown(f'**Documento adicional {i+1}**')
                        title=st.text_input('Nombre',key=f'add_doc_name_{vid}_{i}')
                        pdf=st.file_uploader('PDF',type=['pdf'],key=f'add_doc_pdf_{vid}_{i}')
                        ed=st.file_uploader('Editable',type=['dwg','doc','docx','xlsx','xls','ifc','rvt'],key=f'add_doc_edit_{vid}_{i}')
                        entries.append((title,pdf,ed))
                    submitted=st.form_submit_button('💾 Añadir documentos')
                if submitted:
                    try:
                        if any(not title.strip() for title,_,_ in entries):raise ValueError('Asigna un nombre a cada documento.')
                        if any((p or e) for _,p,e in entries) and not drive_ready():raise ValueError('Conecta Google Drive.')
                        numbers=attach_new_documents(did,version,entries)
                        st.success('Añadidos: '+', '.join(numbers));st.rerun()
                    except Exception as exc:st.error('No se completó la carga: '+str(exc))
            with st.expander('✏️ Editar descripción de la versión'):
                with st.form(f'version_note_{vid}'):
                    note=st.text_area('Descripción de revisión',value=ver['notes'] or '')
                    if st.form_submit_button('Guardar descripción'):
                        with connection() as con:
                            con.execute('UPDATE versions SET notes=? WHERE id=?',(note,vid))
                        st.rerun()
            if st.button('🗑 Eliminar versión '+version,key=f'version_delete_{vid}'):
                st.session_state[f'version_confirm_{vid}']=True
            if st.session_state.get(f'version_confirm_{vid}'):
                st.error('Se eliminará toda la versión y sus documentos. Los archivos vinculados se eliminarán también de Google Drive.')
                if st.button('Confirmar eliminación de versión '+version,key=f'version_yes_{vid}'):
                    try:
                        ids={f['file_id'] for d in docs for f in d['files'].values()}
                        with connection() as con:
                            ids.update(r['file_id'] for r in con.execute('SELECT file_id FROM drive_files WHERE delivery_id=? AND version=?',(did,version)))
                        for fileid in ids:drive_delete_document(fileid)
                        with connection() as con:
                            con.execute('DELETE FROM delivery_documents WHERE delivery_id=? AND version=?',(did,version))
                            con.execute('DELETE FROM drive_files WHERE delivery_id=? AND version=?',(did,version))
                            con.execute('DELETE FROM delivery_files WHERE delivery_id=? AND version=?',(did,version))
                            con.execute('DELETE FROM versions WHERE id=? AND delivery_code=?',(vid,code))
                            remaining=con.execute('SELECT version FROM versions WHERE delivery_code=? ORDER BY id DESC LIMIT 1',(code,)).fetchone()
                            if delivery['version']==version:
                                con.execute('UPDATE deliverables SET version=? WHERE id=?',(remaining['version'] if remaining else '',did))
                        st.session_state.pop(f'version_confirm_{vid}',None)
                        st.rerun()
                    except Exception as exc:st.error('No se pudo eliminar completamente: '+str(exc))

def drive_delete_document(file_id):
    """Borrado remoto solo para archivos previamente vinculados al entregable."""
    require_admin()
    response=requests.delete('https://www.googleapis.com/drive/v3/files/'+quote(file_id,safe=''),
        headers={'Authorization':'Bearer '+drive_token()},timeout=45)
    if response.status_code not in (200,204,404):
        raise ValueError(f'Google Drive rechazó la eliminación (HTTP {response.status_code}).')


def project_editor(data):
    """Proyecto: creación en la cabecera; edición junto a su tarjeta."""
    st.title('📁 Proyectos')
    st.caption('Seguimiento general, avance y riesgo calculado desde el plan de trabajo.')
    if can_edit():
        if st.button('➕ Crear proyecto',type='primary',key='project_create_toggle'):
            st.session_state['show_project_create']=not st.session_state.get('show_project_create',False)
        if st.session_state.get('show_project_create'):
            with st.container(border=True):
                st.subheader('Nuevo proyecto')
                with st.form('new_project_form',clear_on_submit=False):
                    values={}
                    for key,label,kind,required,opt in SPECS['Proyectos'][1]:
                        values[key]=form_input(key,label+(' *' if required else ''),kind,required,opt,None,'new_project')
                    if st.form_submit_button('💾 Guardar proyecto',type='primary'):
                        values={k:(v.isoformat() if isinstance(v,date) else v) for k,v in values.items()}
                        try:
                            save_record('Proyectos',values)
                            st.session_state['show_project_create']=False
                            st.success('Proyecto creado correctamente.');st.rerun()
                        except (ValueError,sqlite3.IntegrityError) as exc:st.error(str(exc))
    projects=data['Proyectos'] if can_edit() else visible_project_rows(
        data['Proyectos'],st.session_state['user_id'])
    if not can_edit():
        st.caption('Solo se muestran los proyectos autorizados por el Administrador.')
    card_browser('Proyectos',projects)
    if can_edit():bulk_delete_ui('Proyectos',data)


def build_deliverable_id(project_code,admin_code):
    """ID estable: prefijo obligatorio del proyecto + código dado por el Admin.

    La convención se aplica a nuevas altas, no reescribe entregables anteriores.
    """
    project=str(project_code or '').strip()
    suffix=str(admin_code or '').strip().upper()
    if not project:
        raise ValueError('Selecciona un ID Proyecto antes de registrar el entregable.')
    if not suffix:
        raise ValueError('Escribe el código del entregable, por ejemplo ARQ-001.')
    if not re.fullmatch(r'[A-Z0-9][A-Z0-9_-]*',suffix):
        raise ValueError('El código adicional solo permite letras, números, guiones y guiones bajos.')
    if suffix.startswith(project.upper()+'_'):
        raise ValueError('Escribe solo el código adicional; el ID Proyecto se agrega automáticamente.')
    full=f'{project}_{suffix}'
    if len(full)>120:
        raise ValueError('El ID Entregable supera los 120 caracteres.')
    return full


def catalog_rows():
    with connection() as con:
        return [dict(r) for r in con.execute('''SELECT c.*,p.name AS project_name,
            (SELECT COUNT(*) FROM deliverables d WHERE d.code=c.code) AS registered
            FROM deliverable_catalog c LEFT JOIN projects p ON p.code=c.project_code
            ORDER BY c.project_code,c.code''')]


def catalog_update(catalog_id,project_code,code,name,specialty,final_date):
    '''Actualiza el catálogo y entregas vinculadas; preserva referencias Drive.'''
    require_admin()
    code=code.strip();name=name.strip()
    if not code or not name:raise ValueError('Se requiere ID y nombre del entregable.')
    with connection() as con:
        previous=con.execute('SELECT * FROM deliverable_catalog WHERE id=?',(catalog_id,)).fetchone()
        if previous is None:raise ValueError('El entregable de catálogo ya no existe.')
        # Los IDs creados antes de esta convención se mantienen intactos.
        # Si el Admin cambia un ID, el nuevo debe incluir el prefijo del proyecto.
        if code!=previous['code'] or project_code!=previous['project_code']:
            prefix=project_code+'_'
            if not code.startswith(prefix):
                raise ValueError('El ID Entregable debe iniciar con el ID Proyecto y un guion bajo: '+prefix)
            if build_deliverable_id(project_code,code[len(prefix):])!=code:
                raise ValueError('El ID Entregable debe seguir el formato del proyecto y código adicional.')
        current=con.execute('SELECT * FROM deliverables WHERE code=?',(previous['code'],)).fetchone()
        linked_tasks=con.execute('SELECT COUNT(*) FROM tasks WHERE delivery_code=?',
                                 (previous['code'],)).fetchone()[0]
        if linked_tasks and project_code!=previous['project_code']:
            raise ValueError('Este entregable tiene tareas vinculadas. Antes de moverlo a otro proyecto, '
                             'desvincula esas tareas en Plan de trabajo.')
        history=con.execute('SELECT highest FROM version_counters WHERE delivery_code=?',
                            (previous['code'],)).fetchone()
        if (current or (history and history['highest']>0)) and (code!=previous['code'] or project_code!=previous['project_code']):
            raise ValueError('El ID y el proyecto no pueden cambiarse si existen entregas o versiones históricas.')
        if con.execute('SELECT id FROM deliverable_catalog WHERE code=? AND id<>?',(code,catalog_id)).fetchone():
            raise ValueError('El código ya está reservado para otro entregable.')
        if con.execute('SELECT id FROM deliverables WHERE code=? AND code<>?',(code,previous['code'])).fetchone():
            raise ValueError('El código coincide con un entregable registrado.')
        remote=con.execute('''SELECT COUNT(*) FROM document_files f JOIN delivery_documents d ON d.id=f.document_id WHERE d.delivery_id=?''',(current['id'],)).fetchone()[0] if current else 0
        project=con.execute('SELECT name FROM projects WHERE code=?',(previous['project_code'],)).fetchone()
    change_folder=None
    if current and remote and previous['name']!=name:
        if not drive_ready():raise ValueError('Primero autoriza Google Drive para renombrar un entregable con archivos.')
        token=drive_token()
        folder=drive_safe_folder_name(previous['project_code']+' - '+(project['name'] or ''))
        project_id=drive_find_folder(token,DRIVE_FOLDER_ID,folder)
        original_name=drive_safe_folder_name(previous['code']+' - '+previous['name'])
        new_name=drive_safe_folder_name(code+' - '+name)
        found=drive_find_folder(token,project_id,original_name) if project_id else None
        if not found:raise ValueError('No se encontró la carpeta anterior en Google Drive; no se modificaron datos.')
        collision=drive_find_folder(token,project_id,new_name)
        if collision and collision!=found:raise ValueError('Ya existe una carpeta con el nuevo nombre.')
        drive_rename_item(token,found,new_name)
        change_folder=(token,found,original_name)
    try:
        with connection() as con:
            con.execute('''UPDATE deliverable_catalog SET project_code=?,code=?,name=?,specialty=?,final_due_date=?
                WHERE id=?''',(project_code,code,name,specialty,final_date,catalog_id))
            if code!=previous['code']:
                # El ID es referencia de tareas, incluidas las que aún no tienen V01.
                con.execute('UPDATE tasks SET delivery_code=? WHERE delivery_code=?',
                            (code,previous['code']))
            # Los vínculos existen incluso si el entregable aún no tiene V01.
            con.execute('''UPDATE tasks SET specialty=? WHERE project_code=? AND delivery_code=?''',
                (specialty or '',project_code,code))
            if current:
                con.execute('''UPDATE deliverables SET name=?,specialty=?,due_date=? WHERE id=?''',
                    (name,specialty,final_date,current['id']))
            con.execute('INSERT INTO audit(module,record_code,action) VALUES (?,?,?)',
                ('Catálogo entregables',code,'Edición por Administrador'))
    except Exception:
        if change_folder:
            try:drive_rename_item(*change_folder)
            except Exception:pass
        raise



def catalog_delete(catalog_id):
    """Elimina la programación sin bloquearla por versiones antiguas eliminadas.

    No borra entregas activas, tareas ni archivos de Drive. Borra el contador
    histórico ÚNICAMENTE de este entregable completo, permitiendo que un nuevo
    entregable con el mismo ID comience en V01. Al borrar solo una versión, el
    contador se sigue conservando (next_version y version_counters no cambian).
    También limpia versiones/checklists locales huérfanos de ese código.
    """
    require_admin()
    with connection() as con:
        con.execute('BEGIN IMMEDIATE')
        item=con.execute('SELECT id,code,project_code FROM deliverable_catalog WHERE id=?',
                         (int(catalog_id),)).fetchone()
        if item is None:
            raise ValueError('Este entregable programado ya no existe. Actualiza la pantalla.')
        code=item['code']
        # Este registro puede tener un contador de versiones o notas antiguas, aun
        # cuando su última versión ya se eliminó. Eso NO impide borrar el catálogo.
        if con.execute('SELECT 1 FROM deliverables WHERE code=? LIMIT 1',(code,)).fetchone():
            raise ValueError('Este entregable todavía tiene una entrega registrada. '
                             'Elimínala primero desde los registros del módulo Entregables.')
        if con.execute('SELECT 1 FROM tasks WHERE delivery_code=? LIMIT 1',(code,)).fetchone():
            raise ValueError('Existen tareas vinculadas a este ID. Desvincúlalas en Plan de trabajo '
                             'antes de borrar la programación; las versiones históricas no bloquean.')
        # No hay entrega activa: se limpian referencias locales y el contador
        # solo de ESTE ID; no afecta a versiones ni a contadores de otros IDs.
        versions_deleted=con.execute('DELETE FROM versions WHERE delivery_code=?',(code,)).rowcount
        checks_deleted=con.execute('DELETE FROM checklist WHERE delivery_code=?',(code,)).rowcount
        counter_deleted=con.execute('DELETE FROM version_counters WHERE delivery_code=?',(code,)).rowcount
        con.execute('DELETE FROM deliverable_catalog WHERE id=?',(int(catalog_id),))
        con.execute('INSERT INTO audit(module,record_code,action) VALUES (?,?,?)',
                    ('Catálogo entregables',code,
                     'Eliminado del proyecto '+item['project_code']+
                     f'; referencias históricas locales limpiadas: {versions_deleted} versiones, {checks_deleted} checks; '
                     f'contador histórico del entregable retirado: {counter_deleted}; sin cambios en Google Drive'))
    return code


def delivery_catalog_admin_ui():
    '''Catálogo por proyecto, solo Admin; permite seleccionar desde Consulta sin escribir IDs.'''
    require_admin()
    st.markdown('### 📋 Entregables programados por proyecto')
    st.caption('Aquí defines el ID, el nombre, la especialidad y la Fecha de Presentación Final. '
               'Esta programación NO equivale a registrar una entrega o subir archivos.')
    projects=[dict(r) for r in connection_project_rows()]
    if not projects:
        st.warning('Primero crea un proyecto en el módulo Proyectos.');return
    project_names={p['code']:p['name'] for p in projects}
    # Fuera del formulario para actualizar en vivo el ID cuando cambia el proyecto
    # o el sufijo escrito por el administrador.
    pr=st.selectbox('ID Proyecto *',list(project_names),
        format_func=lambda c:f'{c} — {project_names[c]}',key='catalog_new_project')
    suffix=st.text_input('Código adicional del entregable *',placeholder='Ej.: ARQ-001',
        help='ALTIVIA antepone automáticamente el ID Proyecto. Escribe solo el código adicional.',
        key='catalog_new_suffix').strip()
    code=f'{pr}_{suffix.upper()}' if suffix else f'{pr}_'
    st.text_input('ID Entregable (automático)',value=code,disabled=True,
        help='Este es el identificador que se guardará. Ej.: 202607_ARQ-001.')
    with st.form('catalog_create_form'):
        name=st.text_input('Nombre del entregable *',placeholder='Ej.: Plantas arquitectónicas').strip()
        specialty=st.selectbox('Especialidad',['']+SPECIALTIES)
        due=st.date_input('Fecha de Presentación Final *',value=today_peru(),format='DD/MM/YYYY')
        if st.form_submit_button('➕ Registrar entregable',type='primary'):
            try:
                code=build_deliverable_id(pr,suffix)
                if not name:raise ValueError('Indica el nombre del entregable.')
                with connection() as con:
                    if not con.execute('SELECT 1 FROM projects WHERE code=?',(pr,)).fetchone():
                        raise ValueError('El proyecto seleccionado ya no existe.')
                    con.execute('''INSERT INTO deliverable_catalog(project_code,code,name,specialty,final_due_date)
                        VALUES(?,?,?,?,?)''',(pr,code,name,specialty,due.isoformat()))
                    con.execute('INSERT INTO audit(module,record_code,action) VALUES (?,?,?)',
                        ('Catálogo entregables',code,'Alta en proyecto '+pr))
                st.success('Entregable añadido al catálogo. Ya puede seleccionarse desde Consulta.');st.rerun()
            except sqlite3.IntegrityError:
                st.error('Ese ID de entregable ya existe. Utiliza otro código adicional.')
            except ValueError as exc:st.error(str(exc))
    catalog=catalog_rows()
    st.caption(f'{len(catalog)} entregable(s) programado(s)')
    term=st.text_input('🔎 Buscar entregable del catálogo',key='catalog_search').strip().casefold()
    pfilter=st.selectbox('Filtrar proyecto',['Todos']+list(project_names),key='catalog_project_filter',
                         format_func=lambda x:x if x=='Todos' else f'{x} — {project_names[x]}')
    for item in catalog:
        if pfilter!='Todos' and item['project_code']!=pfilter:continue
        if term and term not in (item['code']+' '+item['name']+' '+(item['project_name'] or '')).casefold():continue
        with st.container(border=True):
            st.markdown(f"**{escape(item['code'])} — {escape(item['name'])}**")
            st.caption(f"{escape(item['project_code'])} · Fecha final: {escape(item['final_due_date'])} · "
                       + ('Entrega registrada' if item['registered'] else 'Pendiente de registrar'))
            with st.expander('✏️ Editar entregable programado'):
                with st.form(f'catalog_edit_{item["id"]}'):
                    pr2=st.selectbox('ID Proyecto',list(project_names),
                        index=list(project_names).index(item['project_code']),
                        disabled=bool(item['registered']),key=f'cat_pr_{item["id"]}')
                    code2=st.text_input('ID Entregable',value=item['code'],
                        disabled=bool(item['registered']),key=f'cat_code_{item["id"]}').strip()
                    name2=st.text_input('Nombre del entregable',value=item['name'],key=f'cat_name_{item["id"]}').strip()
                    sp_list=['']+SPECIALTIES
                    old_sp=item['specialty'] or ''
                    if old_sp and old_sp not in sp_list:
                        st.caption(f'Especialidad anterior: {old_sp}. Ya no está en el listado; '
                                   'selecciona una especialidad vigente al guardar.')
                    sp=st.selectbox('Especialidad',sp_list,
                                    index=sp_list.index(old_sp) if old_sp in sp_list else 0,
                                    key=f'cat_spec_{item["id"]}')
                    due2=st.date_input('Fecha de Presentación Final',
                        value=date.fromisoformat(item['final_due_date']),format='DD/MM/YYYY',key=f'cat_due_{item["id"]}')
                    if item['registered']:
                        st.caption('El ID y el proyecto permanecen bloqueados porque ya existe una entrega. '
                                   'Sí puedes cambiar nombre, especialidad y fecha final.')
                    if st.form_submit_button('💾 Guardar cambios'):
                        try:
                            catalog_update(item['id'],pr2,code2,name2,sp,due2.isoformat())
                            st.success('Catálogo actualizado.');st.rerun()
                        except (ValueError,sqlite3.IntegrityError,requests.RequestException) as exc:
                            st.error(str(exc))
            # Botón visible dentro de cada tarjeta. Confirmación en dos pasos.
            # Los documentos/archivos de Google Drive nunca se eliminan desde aquí.
            if item['registered']:
                st.button('🗑️ Eliminar entregable',key=f'catalog_del_{item["id"]}',disabled=True)
                st.caption('Elimina primero la entrega activa en Registros. Las versiones históricas ya eliminadas no bloquearán el catálogo.')
            else:
                if st.button('🗑️ Eliminar entregable',key=f'catalog_del_{item["id"]}'):
                    st.session_state[f'catalog_delete_open_{item["id"]}']=True
                if st.session_state.get(f'catalog_delete_open_{item["id"]}'):
                    with st.container(border=True):
                        st.warning(f'¿Eliminar la programación {item["code"]} del proyecto {item["project_code"]}? '
                                   'No se borrarán archivos de Google Drive ni otros entregables. '
                                   'Se reiniciará el contador histórico de este ID, pero no el de los demás. '
                                   'Si luego reutilizas el mismo ID, revisa o archiva sus archivos antiguos en Drive '
                                   'para evitar documentos con nombres duplicados.')
                        confirm=st.checkbox(
                            f'Confirmo eliminar el entregable {item["code"]}',
                            key=f'catalog_del_confirm_{item["id"]}')
                        accept,cancel=st.columns(2)
                        with accept:
                            if st.button('Confirmar eliminación',type='primary',
                                         key=f'catalog_del_confirm_btn_{item["id"]}',
                                         disabled=not confirm,use_container_width=True):
                                try:
                                    backup=backup_database()
                                    deleted=catalog_delete(item['id'])
                                    st.session_state.pop(f'catalog_delete_open_{item["id"]}',None)
                                    st.success(f'Entregable {deleted} eliminado del catálogo. '
                                               f'Respaldo: {os.path.basename(backup)}')
                                    st.rerun()
                                except (ValueError,sqlite3.Error,OSError) as exc:
                                    st.error(f'No se pudo eliminar: {exc}')
                        with cancel:
                            if st.button('Cancelar',key=f'catalog_del_cancel_{item["id"]}',
                                         use_container_width=True):
                                st.session_state.pop(f'catalog_delete_open_{item["id"]}',None)
                                st.rerun()



def connection_project_rows():
    with connection() as con:
        return [dict(r) for r in con.execute('SELECT code,name FROM projects ORDER BY code')]


def delivery_editor(data):
    st.title('📑 Entregables')
    admin=can_edit()
    role=st.session_state.get('role')
    if admin:
        if st.button('📋 Registrar entregables',type='primary',key='toggle_delivery_catalog'):
            st.session_state['catalog_visible']=not st.session_state.get('catalog_visible',False)
        if st.session_state.get('catalog_visible'):
            delivery_catalog_admin_ui()
            st.divider()
    if not (admin or role=='Consulta'):
        card_browser('Entregables',data['Entregables'],allow_version_edit=False);return
    st.subheader('Registrar entrega' if role=='Consulta' else 'Crear o editar una entrega')
    if role=='Consulta':
        st.info('Puedes registrar V01 y después V02, V03… para los entregables asignados. Cada versión anterior se conserva y no puede editarse desde Consulta.')
    projects=connection_project_rows()
    if role=='Consulta':
        allowed_projects=consulta_project_codes(st.session_state['user_id'])
        projects=[pr for pr in projects if pr['code'] in allowed_projects]
    if not projects:
        st.warning('No hay proyectos registrados.');return
    pn={p['code']:p['name'] for p in projects}
    project=st.selectbox('01 · ID Proyecto',list(pn),
        format_func=lambda code:f'{code} — {pn[code]}',key='new_delivery_project')
    cat=[r for r in catalog_rows() if r['project_code']==project]
    if role=='Consulta':
        allowed_deliveries=consulta_delivery_codes(st.session_state['user_id'])
        cat=[r for r in cat if r['code'] in allowed_deliveries]
    with connection() as con:
        used={r['code']:dict(r) for r in con.execute('SELECT * FROM deliverables WHERE project_code=?',(project,))}
    if not cat:
        st.info('No hay entregables asignados mediante tareas en este proyecto, o no tienes permiso de acceso.')
        visible=data['Entregables']
        if role=='Consulta':visible=visible[visible['code'].isin(consulta_delivery_codes(st.session_state['user_id']))]
        card_browser('Entregables',visible,allow_version_edit=admin)
        return
    by_code={r['code']:r for r in cat}
    code=st.selectbox('02 · ID Entregable',list(by_code),
        format_func=lambda c:f'{c} — {by_code[c]["name"]}',key=f'new_delivery_code_{project}')
    catalog=by_code[code]
    existing=used.get(code)
    rid=existing['id'] if existing else None
    action='Registrar nueva versión' if (role=='Consulta' and rid) else 'Nuevo registro'
    if admin and rid:
        action=st.radio('Acción',['Editar registro actual','Registrar nueva versión'],horizontal=True,
            key=f'delivery_action_{code}')
    current=existing or {}
    version=(current.get('version') or 'V01') if existing and action=='Editar registro actual' else next_version(code,current.get('version'))
    st.caption('ID, especialidad y fecha final definidos por Administrador.')
    form_key=f'delivery_editor_{code}_{action}_{version}'
    if action=='Editar registro actual' and rid:
        st.info('Para reemplazar documentos existentes usa «Ver detalles y versiones». Aquí puedes agregar más documentos.')
    optional_docs=bool(rid and action=='Editar registro actual')
    how_many=st.number_input('Cantidad de documentos nuevos',min_value=0 if optional_docs else 1,
        max_value=10,value=0 if optional_docs else 1,key=f'document_count_{form_key}')
    # Fuera de st.form para que las asignaciones se actualicen al elegir tarea.
    st.text_input('Nombre del entregable',value=catalog['name'],disabled=True,
                  key=f'{form_key}_delivery_name')
    st.text_input('Especialidad',value=catalog['specialty'] or 'No asignada',disabled=True,
                  key=f'{form_key}_delivery_specialty')
    st.text_input('Versión',value=version,disabled=True,key=f'{form_key}_delivery_version')
    due=date.fromisoformat(catalog['final_due_date'])
    st.date_input('Fecha de Presentación Final',value=due,disabled=True,
                  format='DD/MM/YYYY',key=f'{form_key}_delivery_due')
    delivery_date=(current.get('actual_date') if (existing and action=='Editar registro actual') else None)
    if delivery_date:
        st.date_input('Fecha de entregable',value=date.fromisoformat(delivery_date),
                      disabled=True,format='DD/MM/YYYY',key=f'{form_key}_delivery_date')
    elif existing and action=='Editar registro actual':
        st.text_input('Fecha de entregable',value='No registrada (dato histórico)',
                      disabled=True,key=f'{form_key}_delivery_date_old')
    else:
        delivery_date=today_peru().isoformat()
        st.date_input('Fecha de entregable',value=date.fromisoformat(delivery_date),
                      disabled=True,format='DD/MM/YYYY',key=f'{form_key}_delivery_date')
    # Solo tareas del proyecto y entregable; Consulta solo las suyas.
    linked=linked_delivery_tasks(project,code,
        st.session_state['user_id'] if role=='Consulta' else None)
    choices={int(t['id']):t for t in linked}
    saved_task=current.get('task_id') if existing else None
    current_task_id=int(saved_task) if saved_task and not pd.isna(saved_task) else None
    default_task=current_task_id if current_task_id in choices else None
    options=[None]+list(choices)
    task_id=st.selectbox('ID Tarea vinculada al entregable',options,
        index=options.index(default_task) if default_task in options else 0,
        format_func=lambda tid:choices[tid]['code'] if tid is not None else '— Seleccionar tarea —',
        key=f'{form_key}_linked_task',
        help='Solo muestra tareas asignadas a este entregable y al proyecto seleccionado.')
    chosen=choices.get(task_id)
    if chosen:
        st.markdown('**Actividad:** '+escape(str(chosen['activity'] or 'Sin actividad')))
        cols=st.columns(2)
        with cols[0]:st.text_area('Responsables (automático)',value=chosen['owners'] or 'Sin asignar',
                                 disabled=True,key=f'{form_key}_owners_auto_{task_id}',height=75)
        with cols[1]:st.text_area('Revisores (automático)',value=chosen['reviewers'] or 'Sin asignar',
                                 disabled=True,key=f'{form_key}_reviewers_auto_{task_id}',height=75)
    else:
        st.info('Selecciona un ID Tarea para completar automáticamente la actividad, responsables y revisores.')
        if not linked:
            st.warning('No hay tareas vinculadas disponibles. Vincula una tarea a este entregable desde Plan de trabajo.')
    with st.form(f'{form_key}_form'):
        vals={'project_code':project,'code':code,'name':catalog['name'],
              'specialty':catalog['specialty'] or '','version':version,'due_date':due.isoformat(),
              'actual_date':delivery_date,'file_path':current.get('file_path') or '',
              'task_id':task_id,
              'owner':chosen['owners'] if chosen else current.get('owner') or '',
              'reviewer':chosen['reviewers'] if chosen else current.get('reviewer') or ''}
        for key,label,kind,required,opt in SPECS['Entregables'][1]:
            if key in ('project_code','code','name','specialty','version','due_date','actual_date','file_path','owner','reviewer'):continue
            if key=='status' and role=='Consulta':
                vals[key]='Pendiente';continue
            old=current.get(key)
            if old is not None and pd.isna(old):old=None
            if key=='status' and action=='Registrar nueva versión':old='Pendiente'
            if key in ('review_date','correction_date','approval_date') and action=='Registrar nueva versión':old=None
            vals[key]=form_input(key,label+(' *' if required else ''),kind,required,opt,old,form_key)
        st.markdown('#### 📎 Documentos de '+version)
        st.caption('Cada documento se guarda como D01, D02… con su PDF y editable opcional en Google Drive.')
        entries=[]
        for i in range(how_many):
            with st.container(border=True):
                st.markdown(f'**Documento {i+1}**')
                title=st.text_input('Nombre del documento *',key=f'{form_key}_title_{i}',
                    placeholder='Ej.: Planta del primer piso')
                pdf=st.file_uploader('PDF para visualizar',type=['pdf'],key=f'{form_key}_pdf_{i}')
                editable=st.file_uploader('Editable (DWG, Word, Excel, IFC o RVT)',
                    type=['dwg','doc','docx','xls','xlsx','ifc','rvt'],key=f'{form_key}_editable_{i}')
                entries.append((title,pdf,editable))
        if not drive_ready():st.warning('Google Drive no está autorizado; no podrás cargar archivos.')
        submitted=st.form_submit_button(
            '💾 Registrar '+version if role=='Consulta' else ('💾 Registrar entrega' if not rid else '💾 Guardar cambios'),
            type='primary')
    if submitted:
        try:
            if any(not name.strip() for name,_,_ in entries):raise ValueError('Escribe el nombre de cada documento.')
            if any((pdf or editable) for _,pdf,editable in entries) and not drive_ready():
                raise ValueError('Google Drive no está conectado. No se guardó el formulario.')
            # Validación anterior al guardado para reducir operaciones parciales.
            for _,pdf,editable in entries:
                for kind,f in (('pdf',pdf),('editable',editable)):
                    if f is None:continue
                    raw=f.getvalue();ext=f.name.rsplit('.',1)[-1].lower() if '.' in f.name else ''
                    if not raw or len(raw)>DRIVE_LIMIT:raise ValueError(f'{f.name}: vacío o mayor de 25 MB.')
                    if kind=='pdf' and (ext!='pdf' or not raw.startswith(b'%PDF-')):
                        raise ValueError(f'{f.name}: PDF no válido.')
                    if kind=='editable' and ext not in ('dwg','doc','docx','xls','xlsx','ifc','rvt'):
                        raise ValueError(f'{f.name}: formato editable no admitido.')
            if role=='Consulta':vals['status']='Pendiente'
            # Revalidación de permisos y versión vigente dentro de save_record.
            if role=='Consulta' and code not in consulta_delivery_codes(st.session_state['user_id']):
                raise PermissionError('Ya no estás asignado a este entregable o proyecto.')
            save_record('Entregables',vals,rid)
            with connection() as con:
                stored=con.execute('SELECT id FROM deliverables WHERE code=?',(code,)).fetchone()
            numbers=attach_new_documents(int(stored['id']),version,entries,
                                         consulta=(role=='Consulta'))
            st.success('Entrega registrada. Documentos: '+', '.join(numbers));st.rerun()
        except (ValueError,PermissionError,sqlite3.IntegrityError,requests.RequestException) as exc:
            st.error(f'No se completó la operación: {exc}')
            st.caption('Si Google Drive falló durante la subida, el registro puede haberse creado parcialmente. '
                       'No repitas la carga sin revisar el historial; pide al Administrador completar los archivos faltantes.')
    st.divider()
    visible=data['Entregables']
    if role=='Consulta':
        visible=visible[visible['code'].isin(consulta_delivery_codes(st.session_state['user_id']))]
    card_browser('Entregables',visible,allow_version_edit=admin)
    if admin:bulk_delete_ui('Entregables',data)

def task_editor(data):
    """Crea tareas con código automático y edita desde cada ficha de Registros."""
    st.title('Plan de trabajo')
    admin=can_edit()
    _,fields=SPECS['Plan de trabajo']
    if admin:
        st.button('➕ Crear tarea',type='primary',on_click=open_task_editor,args=(None,),key='create_task_button')
        if st.session_state.get('task_editor_open'):
            rid=st.session_state.get('task_editor_id')
            existing=data['Plan de trabajo']
            row={}
            if rid is not None:
                matches=existing[existing['id']==rid]
                if matches.empty:
                    st.warning('La tarea seleccionada ya no existe.');close_task_editor();st.rerun()
                row=matches.iloc[0].to_dict()
            heading=('✏️ Editar tarea · '+str(row.get('code'))) if rid is not None else '➕ Crear tarea'
            with st.container(border=True):
                st.markdown('#### '+heading)
                st.caption('El ID de la tarea se asigna automáticamente al guardar y es independiente por proyecto.')
                project_codes=selectors('project')
                if not project_codes:
                    st.warning('Registra primero un proyecto.')
                else:
                    if rid is not None:
                        project=str(row.get('project_code') or '')
                        st.text_input('ID Proyecto *',value=project,disabled=True,key=f'task_edit_project_{rid}')
                        code=str(row['code'])
                    else:
                        project=st.selectbox('ID Proyecto *',project_codes,key='task_project_create')
                        code=preview_task_code(project)
                    st.text_input('ID Tarea (automático)',value=code,disabled=True,
                                  key=f'task_code_readonly_{rid}_{project}_{code}')
                    form_key=f'task_fields_{rid if rid is not None else "nuevo"}_{project}'
                    vals={'project_code':project,'code':code}
                    for key,label,kind,required,opt in fields:
                        if key in ('code','project_code','owner','reviewer'):
                            continue
                        raw=row.get(key)
                        if raw is not None and not isinstance(raw,(list,dict)) and pd.isna(raw):raw=None
                        if key=='delivery_code':
                            vals[key]=form_input(key,label,kind,required,opt,raw,form_key,project_code=project)
                            continue
                        if key=='specialty' and vals.get('delivery_code'):
                            auto=programmed_delivery_specialty(project,vals['delivery_code'])
                            vals[key]=auto
                            st.text_input('Especialidad (automática)',
                                value=auto or 'Sin especialidad en el catálogo',disabled=True,
                                key=f'task_auto_spec_{rid}_{project}_{vals["delivery_code"]}_{auto}')
                            st.caption('La especialidad se obtiene del entregable programado.')
                            continue
                        if key=='specialty' and raw and raw not in SPECIALTIES:
                            st.caption(f'Especialidad anterior: {raw}. Selecciona una vigente si no hay entregable relacionado.')
                        vals[key]=form_input(key,label+(' *' if required else ''),kind,required,opt,
                                             raw,form_key,project_code=project)
                    st.markdown('#### Asignación de equipo')
                    roster=people_for_assignment()
                    labels={int(person['id']):person for person in roster}
                    all_person_ids=list(labels)
                    if rid is not None:
                        default_owners,default_reviewers=task_assigned_people(rid)
                    else:
                        default_owners,default_reviewers=set(),set()
                    def person_label(pid):
                        person=labels[pid]
                        account=' · Usuario Consulta' if person.get('account_role')=='Consulta' else ''
                        return str(person['name'])+account
                    owner_ids=st.multiselect('Responsables * (uno o varios)',all_person_ids,
                        default=[i for i in all_person_ids if i in default_owners],
                        format_func=person_label,key=f'task_owners_{rid}_{project}')
                    reviewer_ids=st.multiselect('Revisores (uno o varios)',all_person_ids,
                        default=[i for i in all_person_ids if i in default_reviewers],
                        format_func=person_label,key=f'task_reviewers_{rid}_{project}')
                    if not all_person_ids:
                        st.warning('Agrega Personal o crea un usuario Consulta para asignar tareas.')
                    st.caption('Los usuarios Consulta solo verán actividades donde son responsables o revisores.')
                    save,cancel=st.columns([2,1])
                    with save:
                        saving=st.button('💾 Guardar tarea' if rid is None else '💾 Guardar cambios',
                            type='primary',key=f'save_task_{rid}_{project}',use_container_width=True)
                    with cancel:
                        st.button('Cancelar',key=f'cancel_task_{rid}_{project}',
                                  on_click=close_task_editor,use_container_width=True)
                    if saving:
                        converted={k:(v.isoformat() if isinstance(v,date) else int(v) if isinstance(v,bool) else v)
                                   for k,v in vals.items()}
                        try:
                            save_record('Plan de trabajo',converted,rid,task_assignees=(owner_ids,reviewer_ids))
                            close_task_editor()
                            st.success('Tarea guardada correctamente.')
                            st.rerun()
                        except (ValueError,sqlite3.IntegrityError,requests.RequestException,PermissionError) as exc:
                            st.error(str(exc))
        bulk_delete_ui('Plan de trabajo',data)
    else:
        st.info('Mis actividades: únicamente se muestran las tareas donde figuras como responsable o revisor.')
    visible_tasks=(data['Plan de trabajo'] if admin else
        visible_consulta_tasks(data['Plan de trabajo'],st.session_state['user_id']))
    task_cards(visible_tasks)


def edit_module(module,data):
    if module=='Proyectos':
        return project_editor(data)
    if module=='Plan de trabajo':
        return task_editor(data)
    if module=='Entregables':
        return delivery_editor(data)
    table,fields=SPECS[module]
    st.title(module)
    admin=can_edit()
    consulta_creates=(module=='Entregables' and st.session_state.get('role')=='Consulta')
    if admin or consulta_creates:
        header='➕ Crear entregable' if consulta_creates else '✏️ Crear o editar registro'
        with st.expander(header,expanded=consulta_creates):
            existing=data[module]
            rid=None;row={}
            if admin:
                items=['➕ Nuevo registro']+[f'{r["code"]} — {r.get("name",r.get("activity",r.get("description","")))}' for _,r in existing.iterrows()]
                choice=st.selectbox('Registro a editar',items,key='pick_'+table)
                if choice!='➕ Nuevo registro':
                    idx=items.index(choice)-1;row=existing.iloc[idx].to_dict();rid=int(row['id'])
            else:
                st.info('Puedes crear un entregable nuevo y adjuntar sus archivos iniciales. Después de guardarlo, solo el administrador podrá modificarlo.')
            task_project=None
            if module=='Plan de trabajo':
                # Fuera de st.form: al cambiar el proyecto, Streamlit actualiza
                # inmediatamente los entregables disponibles dentro del formulario.
                project_list=selectors('project')
                if not project_list:
                    st.warning('Primero registra un proyecto para poder crear una tarea.')
                    return
                preferred=row.get('project_code')
                task_project=st.selectbox('ID Proyecto *',project_list,
                    index=project_list.index(preferred) if preferred in project_list else 0,
                    key='task_project_'+str(rid if rid is not None else 'nuevo'))
                st.caption('Los entregables relacionados proceden del catálogo del Administrador; '
                           'no necesitan tener una primera versión (V01).')
            with st.form('form_'+table+'_'+str(rid),clear_on_submit=False):
                vals={}
                if module=='Plan de trabajo':
                    vals['project_code']=task_project
                version_choice=None
                if module=='Entregables' and admin and rid:
                    st.caption('Una revisión nueva aumenta la versión. Los datos existentes se pueden modificar conservando la versión actual.')
                    version_choice=st.radio('Acción de versión',['Conservar versión actual','Registrar nueva versión'],
                        horizontal=True,index=1,key=f'version_action_{rid}')
                for key,label,kind,required,opt in fields:
                    if module=='Plan de trabajo' and key=='project_code':
                        continue  # Ya elegido antes del formulario para filtrar entregables.
                    if module=='Entregables' and key=='file_path':
                        vals[key]=row.get('file_path') or ''
                        continue
                    raw=row.get('name') if module=='Personal' and key=='code' and row else row.get(key)
                    if raw is not None and not isinstance(raw,(list,dict)) and pd.isna(raw):raw=None
                    if module=='Entregables' and key=='version':
                        code=row.get('code') if row else ''
                        proposed=next_version(code,row.get('version')) if (version_choice=='Registrar nueva versión' or not rid) else (raw or 'V01')
                        vals['version']=st.text_input('Versión *',value=proposed,disabled=True,
                            key=f'computed_version_{rid}_{version_choice}_{proposed}')
                    elif module=='Entregables' and consulta_creates and key=='status':
                        vals[key]='Pendiente';st.text_input('Estado *',value='Pendiente',disabled=True)
                    else:
                        vals[key]=form_input(key,label+(' *' if required else ''),kind,required,opt,raw,
                            'form_'+table+'_'+str(rid if rid is not None else 'nuevo'),
                            project_code=task_project if module=='Plan de trabajo' else None)
                pdf_file=None;edit_file=None
                if module=='Entregables':
                    st.markdown('#### 📎 Documentos del entregable')
                    st.caption('Los archivos se cargarán a Google Drive; no se almacenan como BLOB en SQLite.')
                    pdf_file=st.file_uploader('PDF para visualizar',type=['pdf'],key=f'main_pdf_{rid}_{version_choice}')
                    edit_file=st.file_uploader('Editable para descargar (DWG, DOC o DOCX)',type=['dwg','doc','docx'],
                        key=f'main_edit_{rid}_{version_choice}')
                    if not drive_ready():st.warning('Google Drive no está autorizado. Puedes guardar los datos, pero no subir archivos.')
                if st.form_submit_button('💾 Guardar entregable' if consulta_creates else '💾 Guardar cambios',type='primary'):
                    vals={k:(v.isoformat() if isinstance(v,date) else int(v) if isinstance(v,bool) else v) for k,v in vals.items()}
                    try:
                        if module=='Entregables' and (pdf_file or edit_file) and not drive_ready():
                            raise ValueError('Falta autorizar Google Drive. No se guardó el formulario ni los archivos.')
                        if consulta_creates:
                            vals['status']='Pendiente';vals['version']='V01'
                            if df('deliverables').code.eq(vals['code']).any():
                                raise ValueError('El código del entregable ya existe. Utiliza otro código.')
                        save_record(module,vals,rid)
                        if module=='Entregables' and (pdf_file or edit_file):
                            with connection() as con:
                                rowid=con.execute('SELECT id FROM deliverables WHERE code=?',(vals['code'],)).fetchone()['id']
                            for kind,upload in [('pdf',pdf_file),('editable',edit_file)]:
                                if upload:drive_upload(upload,rowid,vals['code'],vals['version'],kind,
                                    allow_new_consulta=consulta_creates and rid is None)
                        st.success('Registro guardado');st.rerun()
                    except (ValueError,sqlite3.IntegrityError,requests.RequestException,PermissionError) as ex:st.error(str(ex))
        if admin:bulk_delete_ui(module,data)
    elif st.session_state.get('role')=='Consulta':
        st.info('Modo consulta: puedes buscar y visualizar los registros sin modificarlos.')
    card_browser(module,data[module],allow_version_edit=admin)


def add_examples():
    require_admin()
    with connection() as con:
        if con.execute('SELECT COUNT(*) FROM projects').fetchone()[0]:raise ValueError('Solo se pueden cargar los ejemplos si aún no existen proyectos.')
    d=lambda n:(date.today()+timedelta(days=n)).isoformat()
    for item in [dict(code='Andrea Torres (EJEMPLO)',name='Andrea Torres (EJEMPLO)',role='Jefe de Proyecto',specialty='Coordinación',capacity=4,status='Disponible'),dict(code='Luis Vega (EJEMPLO)',name='Luis Vega (EJEMPLO)',role='Ingeniero',specialty='Estructuras',capacity=3,status='Disponible'),dict(code='María Rojas (EJEMPLO)',name='María Rojas (EJEMPLO)',role='Modelador BIM',specialty='BIM',capacity=2,status='Disponible')]:
        save_record('Personal',{k:item.get(k) for k,_,_,_,_ in SPECS['Personal'][1]})
    save_record('Proyectos',dict(code='PR-EX-001',name='Proyecto Ejemplo - Edificio Multifamiliar (EJEMPLO)',client='Cliente ficticio',type='Expediente técnico',location='Lima - Perú',manager='Andrea Torres (EJEMPLO)',start_date=d(-20),due_date=d(12),status='En desarrollo',priority='Alta',notes='REGISTRO FICTICIO - reemplazar para producción'))
    tasks=[('T-EX-001','Arquitectura','Plantas y elevaciones', 'Andrea Torres (EJEMPLO)',-4,65,'En desarrollo',0),('T-EX-002','Estructuras','Planos de cimentación','Luis Vega (EJEMPLO)',3,80,'En revisión',0),('T-EX-003','BIM','Compatibilización de interferencias','María Rojas (EJEMPLO)',6,25,'En desarrollo',0),('T-EX-004','Eléctricas','Circuitos y tableros','Andrea Torres (EJEMPLO)',10,0,'No iniciado',0),('T-EX-005','Sanitarias','Redes de agua y desagüe','Luis Vega (EJEMPLO)',-2,100,'Aprobado',0)]
    for code,sp,activity,owner,due,progress,status,blocked in tasks:
        save_record('Plan de trabajo',dict(code=code,project_code='PR-EX-001',specialty=sp,activity=activity,delivery_code='',owner=owner,reviewer='Andrea Torres (EJEMPLO)',start_date=d(-18),due_date=d(due),progress=progress,status=status,priority='Alta' if due<0 else 'Media',updated_at=d(0),notes='EJEMPLO'))
    for code,sp,name,owner,due,status,obs in [('E-EX-001','Arquitectura','Planta arquitectónica','Andrea Torres (EJEMPLO)',-2,'Con observaciones',3),('E-EX-002','Estructuras','Cimentaciones','Luis Vega (EJEMPLO)',4,'En revisión interna',1),('E-EX-003','BIM','Modelo federado','María Rojas (EJEMPLO)',8,'En desarrollo',0)]:
        save_record('Entregables',dict(code=code,project_code='PR-EX-001',drawing_code=code.replace('E-EX','PL-EX'),name=name,specialty=sp,owner=owner,reviewer='Andrea Torres (EJEMPLO)',version='V01',due_date=d(due),actual_date=None,status=status,review_date=d(-1) if obs else None,correction_date=None,approval_date=None,notes='EJEMPLO',file_path=''))
    save_record('Control de cambios',dict(code='C-EX-001',project_code='PR-EX-001',request_date=d(-2),requester='Cliente ficticio',description='Cambio de distribución de ambientes',reason='Nueva necesidad del cliente',specialty='Arquitectura',affected_drawings='PL-EX-001',owner='Andrea Torres (EJEMPLO)',schedule_impact='Alto',new_due_date=d(16),approved_by='',approval_date=None,status='Solicitado',notes='EJEMPLO'))

def priorities(data):
    t=data['Plan de trabajo'];e=data['Entregables'];p=data['Proyectos'];c=data['Control de cambios']
    out=[]
    def add(priority,kind,project,desc,owner,due,code):out.append(dict(Prioridad=priority,Tipo=kind,Proyecto=project,Detalle=desc,Responsable=owner,Vencimiento=str(due or ''),ID=code))
    for _,r in t.iterrows():
        if r.status in FINISHED_TASK or int(r.progress or 0)>=100:continue
        days=r['Días restantes']
        if pd.notna(days) and days<0:add('🔴 Alta','Tarea atrasada',r.project_code,r.activity,r.owner,r.due_date,r.code)
        elif pd.notna(days) and days==0:add('🔴 Alta','Vence hoy',r.project_code,r.activity,r.owner,r.due_date,r.code)
        elif pd.notna(days) and days==1:add('🟡 Media','Vence mañana',r.project_code,r.activity,r.owner,r.due_date,r.code)
    for _,r in e.iterrows():
        if r.status in FINISHED_DELIVERY:continue
        days=r['Días restantes']
        if pd.notna(days) and days<0:add('🔴 Alta','Entregable atrasado',r.project_code,r.name,r.owner,r.due_date,r.code)
        elif pd.notna(days) and days<=3:add('🟡 Media','Entregable próximo',r.project_code,r.name,r.owner,r.due_date,r.code)
        if r.status in ('Con observaciones','Observado por cliente'):add('🟡 Media','Entregable observado',r.project_code,r.name,r.owner,r.due_date,r.code)
    for _,r in p.iterrows():
        if r.status not in ('Terminado',) and r['Riesgo calculado'] in ('Alto','Crítico'):add('🔴 Alta','Proyecto en riesgo',r.code,r['name'],r.manager,r.due_date,r.code)
    for _,r in c.iterrows():
        if r.status in ('Solicitado',):add('🟡 Media','Cambio por aprobar',r.project_code,r.description,r.owner,r.new_due_date,r.code)
    return pd.DataFrame(out,columns=['Prioridad','Tipo','Proyecto','Detalle','Responsable','Vencimiento','ID'])

def dashboard(data):
    p,t,e,c,pe=[data[k] for k in ['Proyectos','Plan de trabajo','Entregables','Control de cambios','Personal']]
    active=p[~p.status.isin(['Terminado'])] if not p.empty else p
    metrics=[('Proyectos activos',len(active)),('Proyectos atrasados',int((active.Situación=='🔴 Atrasado').sum()) if not active.empty else 0),('Proyectos en riesgo',int(active['Riesgo calculado'].isin(['Alto','Crítico']).sum()) if not active.empty else 0),('Tareas pendientes',int((~t.status.isin(FINISHED_TASK)&(t.progress.fillna(0)<100)).sum()) if not t.empty else 0),('Tareas atrasadas',int((t['Días atraso']>0).sum()) if not t.empty else 0),('Entregables pendientes',int((~e.status.isin(FINISHED_DELIVERY)).sum()) if not e.empty else 0),('Entregables en revisión',int(e.status.str.contains('revisión',case=False,na=False).sum()) if not e.empty else 0),('Con observaciones',int(e.status.isin(['Con observaciones','Observado por cliente']).sum()) if not e.empty else 0),('Cambios pendientes',int(c.status.isin(['Solicitado']).sum()) if not c.empty else 0),]
    st.title('ALTIVIA  |  Panel de control')
    st.caption(f'Consultoría e ingeniería • Actualizado al {TODAY.strftime("%d/%m/%Y")} • Datos almacenados localmente')
    # Tarjetas fluidas: 5 columnas en escritorio, 2 en móvil; contraste independiente del tema.
    cards=''.join(
        '<div class="altivia-kpi-card"><div class="altivia-kpi-label">'
        +escape(str(label))+'</div><div class="altivia-kpi-value">'
        +escape(str(val))+'</div></div>' for label,val in metrics
    )
    st.markdown('<div class="altivia-kpi-grid">'+cards+'</div>',unsafe_allow_html=True)
    st.divider()
    st.subheader('🚨 PRIORIDADES DE HOY')
    urgent=priorities(data)
    if urgent.empty:st.success('No hay alertas con los registros actuales.')
    else:
        selected=st.multiselect('Filtrar prioridades',urgent['Tipo'].unique(),default=list(urgent['Tipo'].unique()))
        st.dataframe(urgent[urgent['Tipo'].isin(selected)],use_container_width=True,hide_index=True,height=340)
    st.subheader('Resumen de proyectos')
    if not p.empty:
        st.dataframe(p[['code','name','manager','% Avance','due_date','Días restantes','status','Riesgo calculado','Situación']].rename(columns={'code':'ID','name':'Proyecto','manager':'Responsable','due_date':'Fecha entrega','status':'Estado'}),hide_index=True,use_container_width=True)
    else:st.info('Registra un proyecto para comenzar.')
    st.subheader('Análisis visual')
    plots=[('Avance por proyecto',p,'name','% Avance','bar'),('Proyectos por estado',p,'status',None,'pie'),('Tareas por estado',t,'status',None,'pie'),('Entregables por estado',e,'status',None,'pie'),('Tareas por persona',pe,'name','Tareas activas','bar'),('Proyectos por nivel de riesgo',p,'Riesgo calculado',None,'pie')]
    for i in range(0,len(plots),2):
        cols=st.columns(2)
        for col,(title,frame,x,y,kind) in zip(cols,plots[i:i+2]):
            with col:
                st.markdown('**'+title+'**')
                if frame.empty:st.info('Sin datos');continue
                if kind=='pie':
                    counts=frame[x].fillna('Sin definir').value_counts().reset_index();counts.columns=['Categoría','Cantidad']
                    fig=px.pie(counts,names='Categoría',values='Cantidad',hole=0.45)
                else:fig=px.bar(frame,x=x,y=y,labels={x:'',y:y or ''},text=y)
                fig.update_layout(height=315,margin=dict(l=5,r=5,t=10,b=5),showlegend=(kind=='pie'))
                st.plotly_chart(fig,use_container_width=True)

CHECK_GROUPS={
    'Datos generales': CHECKS[:4],
    'Información y geometría': CHECKS[4:14],
    'Presentación del plano': CHECKS[14:17]+[CHECKS[18]],
    'Coordinación y revisión': [CHECKS[17]]+CHECKS[19:]
}


def checklist_page():
    st.title('✅ Checklist de calidad por documento')
    st.caption('Cada documento tiene su propio checklist; los resultados no se mezclan entre D01, D02 y otras versiones.')
    deliveries=df('deliverables')
    allowed_docs=None
    if st.session_state.get('role')=='Consulta':
        allowed_docs=consulta_document_ids(st.session_state['user_id'])
        if not allowed_docs:
            st.info('Aún no tienes documentos con asignación de responsable o revisor en proyectos autorizados.');return
        with connection() as con:
            slots=','.join('?' for _ in allowed_docs)
            delivery_ids={r[0] for r in con.execute(
                f'SELECT DISTINCT delivery_id FROM delivery_documents WHERE id IN ({slots})',
                tuple(allowed_docs))}
        deliveries=deliveries[deliveries['id'].isin(delivery_ids)].copy()
    if deliveries.empty:st.info('No hay entregables disponibles para tu cuenta.');return
    query=st.text_input('🔎 Buscar entregable',key='ck_search_doc')
    if query:
        deliveries=deliveries[deliveries.astype(str).apply(lambda x:x.str.contains(query,case=False,regex=False)).any(axis=1)]
    if deliveries.empty:st.info('No hay coincidencias.');return
    code=st.selectbox('Entregable',deliveries.code.tolist(),key='ck_delivery',
        format_func=lambda c:f'{c} — {deliveries.loc[deliveries.code==c,"name"].iloc[0]}')
    with connection() as con:
        record=con.execute('SELECT id,version FROM deliverables WHERE code=?',(code,)).fetchone()
        versions=[v[0] for v in con.execute('SELECT version FROM versions WHERE delivery_code=? ORDER BY id DESC',(code,))]
    if allowed_docs is not None:
        with connection() as con:
            marks=','.join('?' for _ in allowed_docs)
            visible_versions={row[0] for row in con.execute(
                f'SELECT DISTINCT version FROM delivery_documents WHERE delivery_id=? AND id IN ({marks})',
                (record['id'],*tuple(allowed_docs)))}
        versions=[ver for ver in versions if ver in visible_versions]
    if not versions:st.info('No hay versiones con documentos asignados a tu cuenta.');return
    selected_version=st.selectbox('Versión',versions,index=0,key=f'ck_ver_{code}')
    docs=docs_for_version(record['id'],selected_version)
    if allowed_docs is not None:
        docs=[doc for doc in docs if doc['id'] in allowed_docs]
    if not docs:
        st.info('Esta versión todavía no tiene documentos. Agrega D01 desde Entregables.');return
    doc_ids=[d['id'] for d in docs]
    doc_id=st.selectbox('Documento',doc_ids,format_func=lambda i:next(f'{d["doc_no"]} — {d["title"]}' for d in docs if d['id']==i),
        key=f'ck_doc_{code}_{selected_version}')
    doc=next(d for d in docs if d['id']==doc_id)
    with connection() as con:
        existing={r['criterion']:dict(r) for r in con.execute('SELECT criterion,result,notes FROM document_checklist WHERE document_id=?',(doc_id,))}
    results={c:existing.get(c,{}).get('result','PENDIENTE') for c in CHECKS}
    counts={v:sum(1 for result in results.values() if result==v) for v in ('OK','PENDIENTE','NO APLICA')}
    applicable=counts['OK']+counts['PENDIENTE']
    pct=100*counts['OK']/applicable if applicable else 100
    with st.container(border=True):
        st.markdown(f'### 📄 {escape(doc["doc_no"])} – {escape(doc["title"])}')
        cols=st.columns(4)
        for col,label,num in zip(cols,['Cumplimiento','Conformes','Pendientes','No aplica'],
                                  [f'{pct:.0f}%',counts['OK'],counts['PENDIENTE'],counts['NO APLICA']]):
            with col:st.metric(label,num)
        st.progress(pct/100)
    filter_option=st.segmented_control('Filtrar criterios',options=['Todos','Pendientes','Conformes','No aplica'],
        default='Todos',key=f'ck_filter_{doc_id}')
    wanted={'Todos':None,'Pendientes':'PENDIENTE','Conformes':'OK','No aplica':'NO APLICA'}[filter_option]
    if can_edit():
        with st.form(f'document_checklist_{doc_id}'):
            entered=[]
            for group,criteria in CHECK_GROUPS.items():
                subset=[c for c in criteria if wanted is None or results[c]==wanted]
                if not subset:continue
                with st.expander(group,expanded=True):
                    for c in subset:
                        with st.container(border=True):
                            st.markdown('**'+escape(c)+'**')
                            left,right=st.columns([1,2])
                            opts=['PENDIENTE','OK','NO APLICA']
                            with left:status=st.selectbox('Resultado',opts,index=opts.index(results[c]),key=f'dck_stat_{doc_id}_{c}')
                            with right:note=st.text_input('Observación',value=existing.get(c,{}).get('notes') or '',key=f'dck_note_{doc_id}_{c}')
                            entered.append((c,status,note))
            if st.form_submit_button('💾 Guardar checklist'):
                with connection() as con:
                    for criterion,status,note in entered:
                        con.execute('''INSERT INTO document_checklist(document_id,criterion,result,notes) VALUES(?,?,?,?)
                            ON CONFLICT(document_id,criterion) DO UPDATE SET result=excluded.result,notes=excluded.notes''',
                            (doc_id,criterion,status,note))
                    con.execute('INSERT INTO audit(module,record_code,action) VALUES(?,?,?)',
                        ('Checklist',code,f'{selected_version}/{doc["doc_no"]}: {len(entered)} criterios'))
                st.success('Checklist guardado');st.rerun()
    else:
        st.info('Consulta: acceso de solo lectura al checklist.')
        for group,criteria in CHECK_GROUPS.items():
            subset=[c for c in criteria if wanted is None or results[c]==wanted]
            if not subset:continue
            with st.expander(group,expanded=True):
                for criterion in subset:
                    state=results[criterion]
                    icon={'OK':'🟢','PENDIENTE':'🟡','NO APLICA':'⚪'}[state]
                    st.markdown(f'{icon} **{escape(criterion)}** · {state}')
                    note=existing.get(criterion,{}).get('notes')
                    if note:st.caption(note)

def versions_page():
    st.title('📚 Historial de versiones')
    st.caption('También puedes administrar las versiones desde la ficha de cada entregable.')
    deliveries=df('deliverables')
    if st.session_state.get('role')=='Consulta':
        allowed_projects=consulta_project_codes(st.session_state['user_id'])
        deliveries=deliveries[deliveries['project_code'].isin(allowed_projects)].copy()
    if deliveries.empty:st.info('No hay versiones disponibles en tus proyectos autorizados.');return
    q=st.text_input('🔎 Buscar entregable por nombre o código',key='versions_search')
    if q:deliveries=deliveries[deliveries.astype(str).apply(lambda x:x.str.contains(q,case=False,regex=False)).any(axis=1)]
    st.caption(f'{len(deliveries)} entregables encontrados')
    for _,r in deliveries.head(30).iterrows():
        with st.container(border=True):
            st.markdown(f'**{escape(str(r["name"]))}** · {escape(str(r["code"]))}')
            st.caption((f'{r["project_code"]} · {r["status"]} · ' if can_edit() else f'{r["project_code"]} · ') + f'Versión vigente: {r["version"]}')
            with st.expander('Consultar y gestionar versiones'):
                versions_for_delivery(r['code'],allow_edit=can_edit())

def export_xlsx(data):
    dest=io.BytesIO()
    with pd.ExcelWriter(dest,engine='openpyxl',datetime_format='DD/MM/YYYY',date_format='DD/MM/YYYY') as writer:
        for module,frame in data.items():
            frame=frame.drop(columns=['id','created_at'],errors='ignore')
            frame.to_excel(writer,sheet_name=module[:31],index=False)
            sh=writer.sheets[module[:31]];sh.freeze_panes='A2';sh.auto_filter.ref=sh.dimensions
            from openpyxl.styles import PatternFill,Font
            from openpyxl.utils import get_column_letter
            for cell in sh[1]:cell.fill=PatternFill('solid',fgColor='17365D');cell.font=Font(color='FFFFFF',bold=True)
            for i,col in enumerate(frame.columns,1):sh.column_dimensions[get_column_letter(i)].width=min(42,max(14,len(str(col))+3))
        for table,label in [('versions','VERSIONES'),('delivery_documents','DOCUMENTOS'),('document_files','ARCHIVOS_DRIVE'),('document_checklist','CHECKLIST_DOC'),('checklist','CHECKLIST_ANT'),('audit','AUDITORIA')]:
            frame=df(table).drop(columns=['id'],errors='ignore');frame.to_excel(writer,sheet_name=label,index=False)
    dest.seek(0);return dest

def setup_style():
    st.markdown('''<style>
    [data-testid="stSidebar"]{background:#102b48}
    [data-testid="stSidebar"] :is(p,span,label,h1,h2,h3){color:#f2f6fb!important}
    /* Sidebar: botón de salida con contraste alto en todos los temas. */
    [data-testid="stSidebar"] button[kind="secondary"],
    [data-testid="stSidebar"] .stButton > button{
      background:#eaf2fb!important;color:#102b48!important;
      border:1px solid #acc7e3!important;border-radius:9px!important;
    }
    [data-testid="stSidebar"] .stButton > button :is(p,span,div){color:#102b48!important}
    [data-testid="stSidebar"] .stButton > button:hover{background:#cde3f8!important;color:#102b48!important}
    /* Los KPI usan colores explícitos, sin depender del tema claro/oscuro del móvil. */
    .altivia-kpi-grid{display:grid;grid-template-columns:repeat(5,minmax(0,1fr));gap:12px;margin:14px 0 20px}
    .altivia-kpi-card{background:#eff4fa;color:#142c49!important;border:1px solid #d4e0eb;
      border-radius:11px;padding:12px 13px;min-width:0;min-height:97px;box-sizing:border-box}
    .altivia-kpi-label{font-size:0.84rem;line-height:1.3;font-weight:600;color:#334d67!important;overflow-wrap:anywhere}
    .altivia-kpi-value{font-size:1.75rem;line-height:1.2;font-weight:700;margin-top:9px;color:#102b48!important}
    /* Tarjetas de proyectos: azul pizarra suave, contraste adecuado en tema claro/oscuro. */
    [class*="st-key-gp_project_card_"]{
      background:rgba(62,126,191,.105)!important;
      border:1px solid rgba(93,153,212,.35)!important;
      border-left:4px solid #3b80b9!important;
      border-radius:13px!important;
      box-shadow:0 3px 12px rgba(18,53,83,.075);
      padding:10px 12px!important;
    }
    [class*="st-key-gp_project_card_"]:hover{
      background:rgba(62,126,191,.16)!important;
      border-color:rgba(93,153,212,.52)!important;
    }
    /* Semáforo de ENTREGABLES: tonos pastel discretos, compatibles con el azul ALTIVIA. */
    [class*="st-key-gp_delivery_yellow_card_"]{
      background:rgba(226,178,44,.11)!important;
      border:1px solid rgba(194,155,59,.30)!important;
      border-left:4px solid #c7a24b!important;border-radius:12px!important;
      padding:10px 12px!important;
    }
    [class*="st-key-gp_delivery_red_card_"]{
      background:rgba(208,88,91,.095)!important;
      border:1px solid rgba(199,93,98,.32)!important;
      border-left:4px solid #c06a72!important;border-radius:12px!important;
      padding:10px 12px!important;
    }
    [class*="st-key-gp_delivery_green_card_"]{
      background:rgba(62,157,112,.10)!important;
      border:1px solid rgba(77,156,111,.30)!important;
      border-left:4px solid #54a57e!important;border-radius:12px!important;
      padding:10px 12px!important;
    }
    .block-container{padding-top:1.5rem}
    @media(max-width:900px){
      .altivia-kpi-grid{grid-template-columns:repeat(2,minmax(0,1fr));gap:9px}
      .altivia-kpi-card{padding:11px;min-height:84px}
      .altivia-kpi-label{font-size:0.78rem}
      .altivia-kpi-value{font-size:1.5rem;margin-top:7px}
      .block-container{padding-left:0.9rem;padding-right:0.9rem}
    }
    @media(max-width:360px){.altivia-kpi-grid{gap:7px}.altivia-kpi-card{padding:9px}.altivia-kpi-value{font-size:1.35rem}}
    
    </style>''',unsafe_allow_html=True)

def main():
    initialize();setup_style()
    if not initialize_auth():
        st.error('Configuración inicial: establezca ALTIVIA_ADMIN_PASSWORD (mínimo 6 caracteres) antes de iniciar la aplicación. Consulte README_SEGURIDAD.txt.')
        st.stop()
    restore_remembered_login()
    if 'user_id' in st.session_state:
        with connection() as con:
            current=con.execute('SELECT username,full_name,role,active FROM users WHERE id=?',(st.session_state['user_id'],)).fetchone()
        if not current or not current['active']:
            for key in ('user_id','role','username','full_name','client_preview'):st.session_state.pop(key,None)
        else:
            st.session_state['role']=current['role'];st.session_state['username']=current['username'];st.session_state['full_name']=current['full_name']
    if not st.session_state.get('user_id'):
        login_page();return
    with st.sidebar:
        st.markdown('# ALTIVIA')
        st.markdown('**GESTOR DE PROYECTOS**')
        st.caption('Ingeniería · Consultoría · Planos')
        st.caption(f"{st.session_state['full_name']} · {st.session_state['role']}")
        if st.button('Cerrar sesión'):
            end_login_session()
            st.rerun()
        pages=(['Mis documentos','Mi cuenta'] if st.session_state.get('role')=='Cliente' else
               ['Proyectos','Plan de trabajo','Entregables','Checklist','Versiones','Mi cuenta']
               if st.session_state.get('role')=='Consulta' else
               ['Dashboard','Proyectos','Plan de trabajo','Entregables','Control de cambios','Personal','Checklist','Versiones','Mi cuenta'])
        if can_edit():pages+=['Administrar usuarios','Exportación y respaldo']
        page=st.radio('Navegación',pages)
        st.divider();st.caption('🔒 Datos en SQLite local (altivia.db)')
    data=decorate()
    if page=='Mis documentos':client_portal()
    elif page=='Dashboard' and can_edit():dashboard(data)
    elif page in SPECS and st.session_state.get('role')!='Cliente' and (page not in ('Personal','Control de cambios') or can_edit()):
        edit_module(page,data)
    elif page=='Checklist' and st.session_state.get('role')!='Cliente':checklist_page()
    elif page=='Versiones' and st.session_state.get('role')!='Cliente':versions_page()
    elif page=='Mi cuenta':my_account()
    elif page=='Administrar usuarios':user_management()
    else:
        require_admin()
        st.title('Exportación, respaldo e instalación')
        st.download_button('📥 Exportar registros a Excel (.xlsx)',data=export_xlsx(data),file_name='ALTIVIA_Registro_Proyectos.xlsx',mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',type='primary')
        backup_path=backup_database()
        with open(backup_path,'rb') as f:backup=f.read()
        st.download_button('💾 Descargar respaldo íntegro (.db)',data=backup,file_name=f'ALTIVIA_respaldo_{TODAY.isoformat()}.db',mime='application/octet-stream')
        st.warning('El Excel exportado sirve para reportes. La base SQLite es la fuente principal: conserva versiones y checklists. No edite la base mientras la aplicación está abierta.')
        if not len(data['Proyectos']):
            if st.button('Cargar datos ficticios de demostración'):
                try:add_examples();st.success('Ejemplos creados');st.rerun()
                except ValueError as exc:st.error(str(exc))
        st.subheader('Bitácora de actividad');st.dataframe(df('audit').head(100),hide_index=True,use_container_width=True)
        reset_database_ui()
        st.caption('Esta instalación usa SQLite local: no la publique en Streamlit Community Cloud para datos reales. Para varios usuarios, migre a PostgreSQL y alojamiento privado.')

if __name__=='__main__':main()
